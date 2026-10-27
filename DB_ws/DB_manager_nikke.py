import time
import sqlite3
import os
from pathlib import Path
import re
import sys

if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
import requests
import openpyxl
from openpyxl.styles import PatternFill

class DBManager:
    def __init__(self, path_db):
        self.driver = None

        # 현재 파일(__file__)의 절대 경로 확인
        self.current_file = Path(__file__).resolve()
        # 현재 파일이 위치한 폴더 경로
        self.current_dir = self.current_file.parent
        # 프로젝트 폴더(DB_ws의 상위). db의 프로필 이미지 경로('resource/...')는 이 폴더 기준 상대 경로다
        self.path_project = self.current_dir.parent

        # ==== db init
        self.db_connection = sqlite3.connect(path_db)
        self.db_cursor = self.db_connection.cursor()

        # ==== db nikke_info_table element
        # 니케 이름
        self.db_name = "Name"
        # 니케 프로필 이미지 저장 경로
        self.db_path_img_profile = "Path_img_profile"
        # 니케 등급
        self.db_tier = "Tier"
        # 니케 소속 스쿼드
        self.db_squad = "Squad"
        # 니케 코드(속성)
        self.db_code = "Code"
        # 니케 무기
        self.db_weapon_type = "Weapon_type"
        # 니케 무기 장탄수
        self.db_weapon_magazine_capacity = "Weapon_magazine_capacity"
        # 니케 무기 재장전 시간
        self.db_weapon_reload_time = "Weapon_reload_time"
        # 니케 조작타입
        self.db_weapon_operation_type = "Weapon_operation_type"
        # 무기 설명
        self.db_weapon_description = "Weapon_description"
        # 니케 클래스
        self.db_class = "Class"
        # 니케 소속 기업
        self.db_company = "Company"
        # 스킬 1 이름
        self.db_skill_1_name = "Skill_1_name"
        # 스킬 1 설명
        self.db_skill_1_description = "Skill_1_description"
        # 스킬 2 이름
        self.db_skill_2_name = "Skill_2_name"
        # 스킬 2 설명
        self.db_skill_2_description = "Skill_2_description"
        # 버스트 스킬 이름
        self.db_skill_bust_name = "Skill_bust_name"
        # 버스트 스킬 단계
        self.db_skill_bust_step = "Bust_step"
        # 버스트 스킬 쿨타임
        self.db_skill_bust_time = "Skill_bust_time"
        # 버스트 스킬 설명
        self.db_skill_bust_description = "Skill_bust_description"

    def __del__(self):
        # selenium 정리
        if self.driver:
            self.driver.quit()

        # db 연결 끊기
        self.db_connection.close()

    def createDBTable(self):
        '''
        db 파일에 table이 없다면 table 생성
        '''

        # 니케 이름이 중복일 수 있음. ex) 사쿠라 ssr, sr
        self.db_cursor.execute(f'''CREATE TABLE IF NOT EXISTS nikke_info_table(
                               id INTEGER PRIMARY KEY AUTOINCREMENT,
                               {self.db_name} TEXT,
                               {self.db_path_img_profile} TEXT,
                               {self.db_tier} TEXT,
                               {self.db_squad} TEXT,
                               {self.db_code} TEXT,
                               {self.db_weapon_type} TEXT,
                               {self.db_weapon_magazine_capacity} INTEGER,
                               {self.db_weapon_reload_time} REAL,
                               {self.db_weapon_operation_type} TEXT,
                               {self.db_weapon_description} TEXT,
                               {self.db_class} TEXT,
                               {self.db_company} TEXT,
                               {self.db_skill_1_name} TEXT,
                               {self.db_skill_1_description} TEXT,
                               {self.db_skill_2_name} TEXT,
                               {self.db_skill_2_description} TEXT,
                               {self.db_skill_bust_name} TEXT,
                               {self.db_skill_bust_step} TEXT,
                               {self.db_skill_bust_time} REAL,
                               {self.db_skill_bust_description} TEXT
                               )''')

        # 스킬 키워드 사전 테이블 (N:M 연결용 단일 키워드 딕셔너리)
        # 키워드는 열(category)마다 독립된 어휘라 같은 문자열이 다른 열에 있을 수 있다 → (category, keyword) 쌍이 유일
        # (앱도 keyword와 category를 함께 조회한다). 예전 정의(keyword 단독 UNIQUE)인 db는 migrateKeywordMasterUnique가 옮긴다.
        self.db_cursor.execute('''CREATE TABLE IF NOT EXISTS skill_keyword_master_table (
                               id INTEGER PRIMARY KEY AUTOINCREMENT,
                               category TEXT NOT NULL,
                               keyword TEXT NOT NULL,
                               UNIQUE (category, keyword)
                               )''')
        self.migrateKeywordMasterUnique()

        # [니케-스킬 키워드] 관계 테이블 (N:M 관계)
        self.db_cursor.execute('''CREATE TABLE IF NOT EXISTS nikke_keyword_link_table (
                               nikke_id INTEGER,
                               keyword_id INTEGER,
                               PRIMARY KEY (nikke_id, keyword_id),
                               FOREIGN KEY (nikke_id) REFERENCES nikke_info_table (id) ON DELETE CASCADE,
                               FOREIGN KEY (keyword_id) REFERENCES skill_keyword_master_table (id) ON DELETE CASCADE
                               )''')
        self.repairKeywordLinkForeignKey()
        
        # 스킬 키워드 통합 테이블 (유사 표현 통합 — 승환님이 엑셀 '키워드 통합' 시트에서 수작업으로 확정한 결과)
        # category는 마스터 시트 열 이름('발동 시점'/'발동 조건'/'적용 대상'/'효과 - 공격류' 등), 원래 표현 하나는 대표 하나에만 연결됨
        self.db_cursor.execute('''CREATE TABLE IF NOT EXISTS skill_keyword_unify_table (
                               id INTEGER PRIMARY KEY AUTOINCREMENT,
                               category TEXT NOT NULL,
                               original_keyword TEXT NOT NULL,
                               unified_keyword TEXT NOT NULL,
                               UNIQUE (category, original_keyword)
                               )''')

        self.db_connection.commit()

    def migrateKeywordMasterUnique(self):
        '''
        skill_keyword_master_table이 예전 정의(keyword 단독 UNIQUE)라면 (category, keyword) UNIQUE 정의로 옮긴다.
        키워드 id를 그대로 유지하므로 nikke_keyword_link_table의 링크는 바뀌지 않는다. 이미 새 정의면 아무것도 하지 않는다.
        '''
        row = self.db_cursor.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='skill_keyword_master_table'").fetchone()
        if not row or 'UNIQUE (category, keyword)' in row[0]:
            return
        try:
            self.db_cursor.execute('''CREATE TABLE skill_keyword_master_table_new (
                                   id INTEGER PRIMARY KEY AUTOINCREMENT,
                                   category TEXT NOT NULL,
                                   keyword TEXT NOT NULL,
                                   UNIQUE (category, keyword)
                                   )''')
            self.db_cursor.execute("INSERT INTO skill_keyword_master_table_new (id, category, keyword) "
                                   "SELECT id, category, keyword FROM skill_keyword_master_table")
            self.db_cursor.execute("DROP TABLE skill_keyword_master_table")
            self.db_cursor.execute("PRAGMA legacy_alter_table = ON")   # 이름 변경이 다른 테이블의 참조를 건드리지 않게
            self.db_cursor.execute("ALTER TABLE skill_keyword_master_table_new RENAME TO skill_keyword_master_table")
            self.db_cursor.execute("PRAGMA legacy_alter_table = OFF")
            self.db_connection.commit()
            print("[INFO] skill_keyword_master_table을 (category, keyword) UNIQUE 정의로 옮겼습니다(키워드 id 유지).")
        except Exception as error_msg:
            self.db_connection.rollback()
            print(f"[ERROR] migrateKeywordMasterUnique 실행 중 오류 발생(롤백함): {error_msg}")

    def repairKeywordLinkForeignKey(self):
        '''
        nikke_keyword_link_table의 외래키가 'nikke_info_table_legacy'를 가리키면(예전 크롤링이 남긴 흔적) nikke_info_table을
        가리키도록 테이블을 다시 만든다. 링크 데이터는 그대로 옮긴다. 정상이면 아무것도 하지 않는다.
        '''
        row = self.db_cursor.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='nikke_keyword_link_table'").fetchone()
        if not row or 'nikke_info_table_legacy' not in row[0]:
            return
        try:
            self.db_cursor.execute('''CREATE TABLE nikke_keyword_link_table_new (
                                   nikke_id INTEGER,
                                   keyword_id INTEGER,
                                   PRIMARY KEY (nikke_id, keyword_id),
                                   FOREIGN KEY (nikke_id) REFERENCES nikke_info_table (id) ON DELETE CASCADE,
                                   FOREIGN KEY (keyword_id) REFERENCES skill_keyword_master_table (id) ON DELETE CASCADE
                                   )''')
            self.db_cursor.execute("INSERT INTO nikke_keyword_link_table_new SELECT nikke_id, keyword_id FROM nikke_keyword_link_table")
            self.db_cursor.execute("DROP TABLE nikke_keyword_link_table")
            self.db_cursor.execute("PRAGMA legacy_alter_table = ON")
            self.db_cursor.execute("ALTER TABLE nikke_keyword_link_table_new RENAME TO nikke_keyword_link_table")
            self.db_cursor.execute("PRAGMA legacy_alter_table = OFF")
            self.db_connection.commit()
            print("[INFO] nikke_keyword_link_table의 외래키를 nikke_info_table로 복구했습니다(링크 데이터 유지).")
        except Exception as error_msg:
            self.db_connection.rollback()
            print(f"[ERROR] repairKeywordLinkForeignKey 실행 중 오류 발생(롤백함): {error_msg}")

    def backupDB(self):
        '''
        크롤링 또는 업데이트 실행 전 현재 DB의 백업본을 생성
        '''
        import shutil
        db_path = self.current_dir / "nikke_info.db"
        self.backup_db_path = self.current_dir / "nikke_info_backup.db"
        
        if db_path.exists():
            self.db_connection.commit()
            shutil.copy2(str(db_path), str(self.backup_db_path))
            print(f"[백업 생성 완료] {self.backup_db_path}")
        else:
            self.backup_db_path = None
            print("알림: 기존 DB 파일이 존재하지 않아 백업을 생성하지 않았습니다.")

    def generateChangeReport(self):
        '''
        백업 DB와 현재 DB를 비교하여 변경 사항 보고서를 생성
        '''
        import os
        from datetime import datetime
        
        if not hasattr(self, 'backup_db_path') or not self.backup_db_path or not os.path.exists(self.backup_db_path):
            print("알림: 대조할 백업 DB가 존재하지 않습니다.")
            return
        
        # 백업 DB 연결
        backup_conn = sqlite3.connect(str(self.backup_db_path))
        backup_curr = backup_conn.cursor()
        
        # 대조 대상 컬럼 (Name, Path_img_profile, Tier, Squad 제외)
        compare_cols = [
            "Code", "Weapon_type", "Weapon_magazine_capacity", "Weapon_reload_time",
            "Weapon_operation_type", "Weapon_description", "Class", "Company",
            "Skill_1_name", "Skill_1_description", "Skill_2_name", "Skill_2_description",
            "Skill_bust_name", "Bust_step", "Skill_bust_time", "Skill_bust_description"
        ]
        
        select_cols_str = ", ".join(compare_cols)
        
        # 백업 데이터 가져오기 (대조용)
        try:
            backup_curr.execute(f"SELECT Name, {select_cols_str} FROM nikke_info_table")
            backup_data = {r[0]: r[1:] for r in backup_curr.fetchall() if r[0]}
        except sqlite3.OperationalError:
            print("알림: 백업 DB에 nikke_info_table이 없어 대조하지 않습니다.")
            backup_conn.close()
            return
        backup_conn.close()
        
        # 현재(최신화된) 데이터 가져오기 (신규 니케의 Tier, Squad 출력을 위해 추가 선택)
        self.db_cursor.execute(f"SELECT Name, Tier, Squad, {select_cols_str} FROM nikke_info_table")
        current_data = {r[0]: r[1:] for r in self.db_cursor.fetchall() if r[0]}
        
        new_nikkes = []
        changes = []
        
        for name, current_row in current_data.items():
            current_tier = current_row[0]
            current_squad = current_row[1]
            current_vals = current_row[2:]
            
            # 1. 신규 니케 감지 (이름, 티어, 스쿼드만 전용 리스트에 수집)
            if name not in backup_data:
                new_nikkes.append({
                    "name": name,
                    "tier": current_tier if current_tier else "None",
                    "squad": current_squad if current_squad else "None"
                })
                continue
                
            # 2. 기존 니케 정보 변경 검사
            backup_vals = backup_data[name]
            nikke_changes = []
            
            for idx, col in enumerate(compare_cols):
                b_val = backup_vals[idx]
                c_val = current_vals[idx]
                
                b_val_str = str(b_val) if b_val is not None else "None"
                c_val_str = str(c_val) if c_val is not None else "None"
                
                if b_val_str != c_val_str:
                    nikke_changes.append({
                        "column": col,
                        "before": b_val_str,
                        "after": c_val_str
                    })
            
            if nikke_changes:
                changes.append({
                    "nikke_name": name,
                    "diffs": nikke_changes
                })
                
        # 변경 또는 신규 정보가 하나라도 있을 때 보고서 파일 작성
        if new_nikkes or changes:
            date_format = "%Y-%m-%d_%H-%M-%S"
            timestamp = datetime.now().strftime(date_format)
            report_file = self.current_dir / f"db_change_report_{timestamp}.txt"
            
            with open(report_file, "w", encoding="utf-8") as f:
                f.write(f"============================================================\n")
                f.write(f" DB 변경 감지 보고서 (생성 시점: {datetime.now().strftime(date_format)})\n")
                f.write(f"============================================================\n\n")
                
                # [1] 신규 추가된 니케 목록 기재
                if new_nikkes:
                    f.write(f"[1] 신규 추가된 니케 (총 {len(new_nikkes)}명)\n")
                    f.write(f"------------------------------------------------------------\n")
                    for nn in new_nikkes:
                        f.write(f"● 니케 이름: {nn['name']}\n")
                        f.write(f"  - 등급 (Tier): {nn['tier']}\n")
                        f.write(f"  - 소속 (Squad): {nn['squad']}\n")
                        f.write(f"------------------------------------------------------------\n")
                    f.write("\n\n")
                else:
                    f.write("[1] 신규 추가된 니케 없음\n\n\n")
                    
                # [2] 기존 니케 정보 변경 목록 기재
                if changes:
                    f.write(f"[2] 기존 니케 정보 변경 (총 {len(changes)}명)\n")
                    f.write(f"------------------------------------------------------------\n")
                    for change in changes:
                        f.write(f"● 니케 이름: {change['nikke_name']}\n")
                        for diff in change['diffs']:
                            f.write(f"  [수정 컬럼] {diff['column']}\n")
                            before_raw = diff['before'].replace('\r\n', '\\n').replace('\n', '\\n')
                            after_raw = diff['after'].replace('\r\n', '\\n').replace('\n', '\\n')
                            f.write(f"    - 변경 전 : '{before_raw}'\n")
                            f.write(f"    - 변경 후 : '{after_raw}'\n")
                        f.write("-"*60 + "\n\n")
                else:
                    f.write("[2] 기존 니케 정보 변경 없음\n\n")
                    
            print(f"\n[성공] DB 변경 감지 보고서가 생성되었습니다:\n       -> {report_file}")
        else:
            print("\n[알림] 대조 완료: 변경된 기존 니케 및 신규 니케가 존재하지 않습니다.")
            
        # 임시 백업 파일 정리
        try:
            if os.path.exists(self.backup_db_path):
                os.remove(self.backup_db_path)
                print("임시 백업 DB 파일을 제거했습니다.")
        except Exception as e:
            print(f"임시 백업 DB 제거 중 에러: {e}")

    def accessPage(self):
        '''
        크롤링 시작할 페이지로 접속하는 로직
        '''
        # ==== selenium init
        chrome_options = Options()
        # 파이썬이 끝나도 크롬창이 종료되지 않게 하는 옵션
        # vsocde에서도 적용되게 하기 위해서 마지막 코드에 breakpoint 걸면 됨.
        # 그냥 py 프로그램을 실행할 때는 바로 적용됨
        chrome_options.add_experimental_option("detach", True)

        self.driver = webdriver.Chrome(options=chrome_options)

        # 블라블라 접속
        self.driver.get('https://www.blablalink.com/')
        # 동적 로딩 때문에 로딩에 시간이 필요함
        time.sleep(2)

        # 모바일 레이아웃으로 전환
        self.driver.find_element(By.XPATH, '//*[@id="layout-content"]/div/header/div/div[1]/button').click()
        time.sleep(1)

        # 쿠키 창 닫기. 쿠키 창 닫아야 이후 제대로 클릭할 수 있음
        el_cookie_btn = self.driver.find_element(By.XPATH,'//*[@id="onetrust-close-btn-container"]/button')
        el_cookie_btn.click()

        self.wait(By.XPATH,'//*[@id="layout-content"]/div/div[1]/div[2]/div[1]')
        #time.sleep(1)

        # 한국어로 변경 - 링크를 변경하는 방법이 아니어서 필수임
        # 언어 변경 버튼 클릭
        el_lang_change = self.driver.find_element(By.XPATH, '//*[@id="layout-content"]/div/div[1]/div[2]/div[1]')
        el_lang_change.click()

        #self.wait(By.XPATH, '/html/body/div[5]/div[2]/div[3]/div[1]/div[2]/div[3]')
        time.sleep(1)

        # 한국어 선택
        el_lang_korean = self.driver.find_elements(By.XPATH, "//*[text()='한국어']")[0]
        el_lang_korean.click()

        #self.wait(By.XPATH, '/html/body/div[5]/div[2]/div[5]/div[1]')
        time.sleep(1)

        # 확인 누르기
        el_lang_confirm = self.driver.find_elements(By.XPATH, "//*[text()='Confirm']")[0]
        el_lang_confirm.click()

        #self.wait(By.XPATH, '//*[@id="layout-content"]/div/div[3]/div[3]/div/div[2]/div/div[7]')
        time.sleep(3)

        # 니케 도감 페이지 접속
        el_nikkepedia = self.driver.find_element(By.XPATH, '//*[@id="layout-content"]/div/div[3]/div[3]/div/div[1]/div[2]/div/div[7]')
        el_nikkepedia.click()
    
    def crawlingAllUpdateDB(self, only_new=False):
        '''
        필수 - accessPage로 크롤링을 시작할 페이지에 접속
        크롤링을 해서 db에 넣을 데이터로 만들기
        only_new = True 일 때 기존 db에서 없는 니케만 가져옴
        '''

        # DB에서 기존 (이름, 티어) 정보 가져오기
        self.db_cursor.execute(f"SELECT {self.db_name}, {self.db_tier} FROM nikke_info_table")
        existing_nikkes = set(self.db_cursor.fetchall())

        # 니케 리스트 가져오기
        self.wait(By.XPATH, '//*[@data-cname="all-item"]')
        item_list = self.driver.find_elements(By.XPATH, '//*[@data-cname="all-item"]')
        print(f"도감에 등록된 니케 총 {len(item_list)}개")

        time.sleep(2)
        # 니케 프로필 이미지가 모두 로딩되도록 스크롤 제어
        for i in range(33, len(item_list), 24):
            temp = self.driver.find_element(By.XPATH, f'//*[@id="layout-content"]/div/div[2]/div[2]/div/div/div[3]/div[{i}]')
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", temp)
            time.sleep(2)

        # 1단계: 전체 리스트에서 이름과 티어 파악 및 새로운 니케 체크
        nikke_list_info = []
        new_nikke_found = False

        for item in item_list:
            el_span = item.find_element(By.XPATH, ".//div[contains(@class, 'nikkes-all-item-btm')]/span")
            text_name = el_span.text

            text_tier = self.parseListItemTier(item)

            nikke_list_info.append((text_name, text_tier))
            if (text_name, text_tier) not in existing_nikkes:
                new_nikke_found = True

        if only_new and not new_nikke_found:
            print("새로운 니케가 없습니다.")
            return

        if new_nikke_found:
            print("새로운 니케 발견! DB 순서 정렬을 위해 재구축을 시작합니다.")
        else:
            print("기존 데이터 전체 갱신 크롤링을 시작합니다.")

        # 2단계: 기존 테이블 백업 및 새 테이블 생성
        # legacy_alter_table을 켜야 이름을 바꿀 때 nikke_keyword_link_table의 외래키가 'nikke_info_table_legacy'로
        # 따라 바뀌지 않는다(꺼져 있으면 SQLite가 다른 테이블의 참조까지 고쳐 써서, legacy 삭제 후 없는 테이블을 가리키게 됨)
        self.db_cursor.execute("PRAGMA legacy_alter_table = ON")
        self.db_cursor.execute("ALTER TABLE nikke_info_table RENAME TO nikke_info_table_legacy")
        self.createDBTable() # 새 nikke_info_table 생성
        self.db_connection.commit()

        # 3단계: 순서대로 데이터 채우기
        for i, (name, tier) in enumerate(nikke_list_info):
            # 기존에 있던 니케인지 확인
            self.db_cursor.execute("SELECT * FROM nikke_info_table_legacy WHERE Name=? AND Tier=?", (name, tier))
            legacy_data = self.db_cursor.fetchone()

            if only_new and legacy_data:
                # 기존 데이터가 있으면 복사 (id 제외)
                columns = [
                    self.db_name, self.db_path_img_profile, self.db_tier, self.db_squad, self.db_code,
                    self.db_weapon_type, self.db_weapon_magazine_capacity, self.db_weapon_reload_time,
                    self.db_weapon_operation_type, self.db_weapon_description, self.db_class, self.db_company,
                    self.db_skill_1_name, self.db_skill_1_description, self.db_skill_2_name, 
                    self.db_skill_2_description, self.db_skill_bust_name, self.db_skill_bust_step,
                    self.db_skill_bust_time, self.db_skill_bust_description
                ]
                col_str = ", ".join(columns)
                placeholders = ", ".join(["?" for _ in columns])
                
                # legacy_data[0]은 id이므로 제외. 이미지 경로는 예전 규칙(등급 없음)일 수 있어 현재 규칙으로 다시 만든다
                legacy_values = list(legacy_data[1:])
                legacy_values[1] = self.makeProfileImagePath(name, tier)
                self.db_cursor.execute(f"INSERT INTO nikke_info_table ({col_str}) VALUES ({placeholders})", legacy_values)
                print(f"[{i+1}/{len(item_list)}] {name} ({tier}) 기존 데이터 복사 완료 (crawling skip)")
            else:
                # only_new가 False이거나 새로운 니케이면 웹 상세 페이지 크롤링
                print(f"[{i+1}/{len(item_list)}] {name} ({tier}) 상세 정보 웹 크롤링 시작")
                
                # 상세 페이지 접속
                current_item_list = self.driver.find_elements(By.CSS_SELECTOR, "[class='cursor-pointer relative nikkes-all-item h-[180px] max-h-[180px] w-[102px] max-w-[22%] my-[4.5px] mx-[5px] !w-full !max-w-none !m-[0] aspect-[10/18] !h-[auto] !max-h-none']")
                target_item = current_item_list[i]
                self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", target_item)
                time.sleep(1)
                target_item.click()
                time.sleep(3)

                # 상세 정보 파싱
                parsed_data = self._parseNikkeDetail(name)
                
                # DB 삽입
                columns = list(parsed_data.keys())
                col_str = ", ".join(columns)
                placeholders = ", ".join([f":{c}" for c in columns])
                self.db_cursor.execute(f"INSERT INTO nikke_info_table ({col_str}) VALUES ({placeholders})", parsed_data)
                
                # 리스트 페이지로 복귀
                el_prev_page_btn = self.driver.find_element(By.XPATH, '//*[@id="layout-content"]/div/div[1]/div/div[1]')
                el_prev_page_btn.click()
                time.sleep(5)
                print(f"[{i+1}/{len(item_list)}] {name} ({tier}) 상세 정보 수집 및 저장 완료")

            self.db_connection.commit()

        # 4단계: 레거시 테이블 삭제
        self.db_cursor.execute("DROP TABLE nikke_info_table_legacy")
        self.db_connection.commit()
        self.db_cursor.execute("PRAGMA legacy_alter_table = OFF")
        print("DB 순서 정렬 및 업데이트 완료.")

        # 5단계: 누락된 내용 추가
        self.manualUpdate()

    def parseListItemTier(self, item):
        '''도감 목록의 니케 슬롯(item)에서 등급(SSR/SR/R)을 읽는다(슬롯 테두리 색의 class로 판별). 못 읽으면 빈 문자열.'''
        try:
            for div in item.find_elements(By.TAG_NAME, "div"):
                class_content = div.get_attribute("class")
                if class_content:
                    if "yellow" in class_content: return "SSR"
                    elif "purple" in class_content: return "SR"
                    elif "blue" in class_content: return "R"
        except: pass
        return ""

    def makeProfileImagePath(self, name, tier):
        '''
        니케 프로필 이미지의 db 저장 경로(프로젝트 폴더 기준 상대 경로)를 만든다.
        이름이 같은 니케(예: 사쿠라 SSR/SR)가 있어 파일 이름에 등급을 붙인다 → 'resource/사쿠라_SSR_profile.webp'.
        파일 이름에 ':'를 쓸 수 없어 '-'로 바꾼다 → 'resource/라피 - 레드 후드_SSR_profile.webp'.
        '''
        return f"resource/{name.replace(':', '-')}_{tier}_profile.webp"

    def _parseNikkeDetail(self, name):
        '''
        니케 상세 페이지에서 정보를 파싱하여 dict로 반환
        '''
        # 레어도
        el_tier = self.driver.find_element(By.XPATH, '//*[@id="nikkes-basics"]/div[2]/div[2]/p/img')
        url_tier = el_tier.get_attribute("src")
        text_tier = ""
        if "014d90279e0ca7278b0b8f7e9094a8dd" in url_tier: text_tier = "SSR"
        elif "64ec71645699f5fce79d98cfc20a525f" in url_tier: text_tier = "SR"
        elif "1c5d999c42ed640c95da540af7578667" in url_tier: text_tier = "R"

        # 스쿼드 이름
        el_squad = self.driver.find_element(By.XPATH, '//*[@id="nikkes-basics"]/div[2]/div[2]/div/p/span')
        text_squad = el_squad.text
        
        # 속성
        el_code = self.driver.find_element(By.XPATH, '//*[@id="nikkes-basics"]/div[2]/div[3]/div[1]/p[1]/img')
        url_code = el_code.get_attribute("src")
        text_code = ""
        if "code-fire" in url_code: text_code = "작열"
        elif "code-water" in url_code: text_code = "수냉"
        elif "code-electronic" in url_code: text_code = "전격"
        elif "code-wind" in url_code: text_code = "풍압"
        elif "code-iron" in url_code: text_code = "철갑"

        # 클래스
        el_class = self.driver.find_element(By.XPATH, '//*[@id="nikkes-basics"]/div[2]/div[3]/div[3]/p[1]/img')
        url_class = el_class.get_attribute("src")
        text_class = ""
        if "6337300758895c79f9afe8139500adcf" in url_class: text_class = "화력형"
        elif "6c30fdc2d204905edd388fc958359ec3" in url_class: text_class = "방어형"
        elif "e0ff69fe9e3cf29c55232d3590135811" in url_class: text_class = "지원형"

        # 기업
        el_company = self.driver.find_element(By.XPATH, '//*[@id="nikkes-basics"]/div[2]/div[3]/div[4]/p[1]/img')
        url_company = el_company.get_attribute("src")
        text_company = ""
        if "4aebd6d57c8afd17d334f37130ddc6e1" in url_company: text_company = "엘리시온"
        elif "f3d8f861af997ddc5e1ee81b715ae314" in url_company: text_company = "미실리스"
        elif "8c41a47786bc43a7db71bd7218f4989d" in url_company: text_company = "테트라"
        elif "29d8e5c949e2141795446ca64a867b9b" in url_company: text_company = "필그림"
        elif "2e913befc4fc5e84f9807f8a455232dc" in url_company: text_company = "어브노멀"

        # 버스트 단계
        el_bust_step = self.driver.find_element(By.XPATH, '//*[@id="nikkes-basics"]/div[2]/div[3]/div[5]/p[1]/img')
        url_bust_step = el_bust_step.get_attribute("src")
        text_bust_step = ""
        if "448005ff9513c9a8afabbece6ad2a05b" in url_bust_step: text_bust_step = "I"
        elif "50f0e38e5306ba4ddae04e494921216a" in url_bust_step: text_bust_step = "II"
        elif "3b98d581bb5776454e084075adf00c4c" in url_bust_step: text_bust_step = "III"
        elif "25bb6a298db42bf89464714a7fd6f159" in url_bust_step: text_bust_step = "A"

        # 무기 정보
        el_weapon_btn = self.driver.find_element(By.XPATH, '//*[@id="nikkes-basics"]/div[2]/div[3]/div[2]')
        el_weapon_btn.click()
        time.sleep(2)

        text_weapon_type = self.driver.find_element(By.XPATH,'//*[@id="nikkes-weapon"]/div/div/div[1]/p[1]/span').text
        text_weapon_magazine_capacity = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[2]/div[2]/div/div[1]/p[1]/span[2]').text
        text_weapon_reload_time = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[2]/div[2]/div/div[1]/p[2]/span[2]').text.split("s")[0].strip()
        text_weapon_operation_type = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[2]/div[2]/div/div[1]/p[3]/span[2]').text
        text_weapon_spec = self.driver.find_element(By.CSS_SELECTOR, '#nikkes-weapon > div > div > div:nth-child(2) > div.w-80.border-top-gray.pt-10 > div > div.w-half.pl-5.mt-10 > div').get_attribute("innerText")

        self.driver.find_element(By.XPATH, '/html/body/div[4]/div/a').click() # 닫기
        time.sleep(1)

        # 스킬 정보
        el_skill = self.driver.find_element(By.CSS_SELECTOR, "[id='nikkes-weapon']")
        self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", el_skill)
        time.sleep(2)

        # 스킬 1
        text_skill_1_name = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[1]/div[1]/div[1]/p').text
        self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[1]/div[1]/div[1]/div[2]/a[4]').click()
        time.sleep(0.5)
        text_skill_1_spec = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[1]/div[1]/div[2]/div').text
        
        # 스킬 2
        text_skill_2_name = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[2]/div[1]/div[1]/p').text
        self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[2]/div[1]/div[1]/div[2]/a[4]').click()
        time.sleep(0.5)
        text_skill_2_spec = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[2]/div[1]/div[2]/div').text

        # 버스트
        text_skill_bust_name = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[3]/div[1]/div[1]/p').text
        text_skill_bust_time = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[3]/div[1]/div[1]/div[3]/span').text.split("s")[0].strip()
        self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[3]/div[1]/div[1]/div[2]/a[4]').click()
        time.sleep(0.5)
        text_skill_bust_spec = self.driver.find_element(By.XPATH, '//*[@id="nikkes-weapon"]/div/div/div[3]/div[1]/div[2]/div').text

        return {
            self.db_name: name,
            self.db_path_img_profile: self.makeProfileImagePath(name, text_tier),
            self.db_tier: text_tier,
            self.db_squad: text_squad,
            self.db_code: text_code,
            self.db_weapon_type: text_weapon_type,
            self.db_weapon_magazine_capacity: int(text_weapon_magazine_capacity),
            self.db_weapon_reload_time: text_weapon_reload_time,
            self.db_weapon_operation_type: text_weapon_operation_type,
            self.db_weapon_description: text_weapon_spec,
            self.db_class: text_class,
            self.db_company: text_company,
            self.db_skill_bust_step: text_bust_step,
            self.db_skill_1_name: text_skill_1_name,
            self.db_skill_1_description: text_skill_1_spec,
            self.db_skill_2_name: text_skill_2_name,
            self.db_skill_2_description: text_skill_2_spec,
            self.db_skill_bust_name: text_skill_bust_name,
            self.db_skill_bust_time: text_skill_bust_time,
            self.db_skill_bust_description: text_skill_bust_spec
        }

    def DownloadResource(self):
        '''
        필수 - accessPage 후 진행
        1. db에 등록된 프로필 이미지 경로 중 실제 파일이 없는 것을 (이름, 등급)으로 선별
        2. 도감 목록에서 (이름, 등급)이 같은 니케의 이미지를 그 경로로 다운로드
           (이름이 같은 니케가 있어 이름만으로는 구분할 수 없다 — 예: 사쿠라 SSR/SR)
        '''
        self.db_cursor.execute(f"SELECT {self.db_name}, {self.db_tier}, {self.db_path_img_profile} FROM nikke_info_table")

        # db에 있지만 실제로 없는 이미지 선별 : (이름, 등급) -> 저장할 절대 경로
        download_targets = {}
        for name, tier, img_path in self.db_cursor.fetchall():
            if img_path and not (self.path_project / img_path).exists():
                download_targets[(name, tier)] = self.path_project / img_path
        if not download_targets:
            print("다운로드할 프로필 이미지가 없습니다.")
            return

        # 니케 리스트 가져오기.
        # 니케 슬릇에 적용된 css로 전부 가져와서 리스트로 만듦
        self.wait(By.XPATH, '//*[@data-cname="all-item"]')
        item_list = self.driver.find_elements(By.XPATH, '//*[@data-cname="all-item"]')
        
        time.sleep(0.5)
        # 니케 프로필 이미지가 모두 로딩되도록 스크롤 제어
        for i in range(33, len(item_list), 24):
            temp = self.driver.find_element(By.XPATH, f'//*[@id="layout-content"]/div/div[2]/div[2]/div/div/div[3]/div[{i}]')
            self.driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", temp)
            time.sleep(2)

        for item in item_list:
            # 니케 이름과 등급 가져오기
            el_span = item.find_element(By.XPATH, ".//div[contains(@class, 'nikkes-all-item-btm')]/span")
            text_name = el_span.text
            text_tier = self.parseListItemTier(item)

            file_path = download_targets.get((text_name, text_tier))
            if file_path is None:
                continue

            el_img = item.find_element(By.CSS_SELECTOR, "img[class='nikkes-all-item-img']")
            url_img: str = el_img.get_attribute("src")# 프로필 이미지 링크
        
            response = requests.get(url_img)
            if response.status_code == 200: # http 200 -> 성공
                with open(file_path, 'wb') as file:
                    file.write(response.content)
                print(f"프로필 이미지 저장: {file_path.name}")
                time.sleep(0.5)
            else:
                print(f"프로필 이미지 Http error - {response.status_code}")

    def wait(self, by_type, element, timeout=10):
        '''
        해당 요소가 위치할 때 까지 대기
        '''
        WebDriverWait(self.driver, timeout=timeout).until(EC.visibility_of_element_located((by_type, element)))

    def manualUpdate(self):
        '''
        수동으로 DB 업데이트 진행
        - [라피 : 레드후드],[레드 후드] 버스트 스킬 내용 추가
        '''
        ## 1. 라피 : 레드 후드 버스트 스킬 내용 업데이트
        qeury = f"SELECT {self.db_skill_bust_description} FROM nikke_info_table WHERE Name='라피 : 레드 후드'"
        self.db_cursor.execute(qeury)
        description:str = self.db_cursor.fetchall()[0][0]
        if description == "None":
            return
        elif "\n\n" in description:
            parts = description.split("\n\n")
            updated_description = ("<Step 1에서 사용 시 : 부대 보조 행동 양식>\n" + 
                                parts[0] + "\n" + 
                                "<Step 3에서 사용 시 : 고기동 근접형 결전 병기>\n" + 
                                parts[1])
            qeury = f"UPDATE nikke_info_table SET {self.db_skill_bust_description}='{updated_description}' WHERE Name='라피 : 레드 후드'"
            self.db_cursor.execute(qeury)
            self.db_connection.commit()
        
        ## 2. 레드 후드 버스트 스킬 내용 업데이트
        qeury = f"SELECT {self.db_skill_bust_description} FROM nikke_info_table WHERE Name='레드 후드'"
        self.db_cursor.execute(qeury)
        description:str = self.db_cursor.fetchall()[0][0]
        if description == "None":
            return
        elif "\n\n" in description:
            parts = description.split("\n\n")
            updated_description = ("<Step 1에서 사용 시 : 비스트 케이지>\n" + 
                                parts[0] + "\n" + 
                                "<Step 2에서 사용 시 : 엔드 하울링>\n" + 
                                parts[1] + "\n" +
                                "<Step 3에서 사용 시 : 레드 울프>\n" +
                                parts[2])
            qeury = f"UPDATE nikke_info_table SET {self.db_skill_bust_description}='{updated_description}' WHERE Name='레드 후드'"
            self.db_cursor.execute(qeury)
            self.db_connection.commit()

class KoreanAnalyzer:
    # 공용 표현 : 100% 사람이 수작업으로 관리하는 목록. 원본은 여기이며, 수정은 이 목록에서 한다.
    COMMON_EXPRESSIONS = [
        '주목', '불굴', '디코이', '순차 공격', '포커싱', '사격 집중', '펠릿', '생존',
        '부착형 발사체', '발사체 부착 대미지', '발사체 폭발 대미지', '무기 변경', '사용 무기 변경',
        '장탄 수 무한', '도발', '타겟', '저장', '파츠 대미지', '연속 공격', '해로운 효과', '기절',
        '분배 대미지', '1인 공격 대상에서 제외 직접 피격 시 해제', '강제 재장전', '무적', '삭제',
        '랩쳐', '낙인',
        # 헤더 틀(고유 속성 규칙 8)에서 흔히 걸려 나오는 일반 명사 — 니케 전용 이름이 아니다
        '체력', '스킬', '버스트',
    ]

    # 본문 요소 판정 규칙 3(명칭 단독 대괄호) : 이 라벨들이 뒤따르면 스펙 시트로 보고 통째로 폐기
    SPEC_LABELS = [
        '대미지', '차지 시간', '풀 차지 대미지', '최대 장탄 수', '공격 간격', '공격 속도',
        '펠릿 개수', '유지 시간', '유지시간', '발사체 부착 대미지', '발사체 폭발 대미지',
        '최초 대미지', '일반 대미지',
    ]

    # 본문 요소 판정 규칙 7 : 통째로 폐기하는 설명문 라벨(규칙 3의 범위 판정과 무관하게 항상 적용)
    DISCARD_LABELS = ['기능', '특이사항', '특이 사항', '해제 조건', '제거 시점']

    # 부가 옵션 : 키워드 대상이 아닌 요소. 정확히 일치해야 하는 고정 패턴 + '~별 효과' 계열 패턴
    OPTION_PATTERNS = [
        r'^\d+(\.\d+)?\s*초\s*유지$',
        r'^\d+발\s*유지$',
        r'^\d+(\.\d+)?\s*초\s*간격$',
        r'^\d+\s*중첩$',
        r'^지속$',
        r'^해제\s*불가$',
        r'^중첩\s*불가$',
        r'^중복\s*적용\s*불가$',
        r'^중복\s*적용$',
        r'^전투\s*중\s*(?:단\s*)?\d+회\s*발동$',
        r'^\d+\s*회(?:\s*발동)?$',
        r'^최대\s*\d+(\.\d+)?%$',
        r'^최대\s*\d+\s*중첩$',
        r'^하위\s*효과\s*중복\s*적용$',
        r'^해당되는\s*효과만\s*적용$',
        r'^단계별\s*효과만\s*적용$',
        r'^최대\s*\d+\s*단계$',
        r'^최대\s*\d+\s*개$',
        # 횟수/수량 상한·하한 표기 : '최대 2회', '최대 4개 보유', '최소 1개 유지', '최대 11레벨까지 성장',
        # '최대 2장까지 발권 가능', '최대 MP 12 축적'
        r'^(?:최대|최소)\s*\d+\s*(?:회|개|장|레벨|발|기)(?:까지)?(?:\s*(?:보유|유지|성장|발권 가능|저장|축적))?$',
        r'^최대\s*[A-Za-z가-힣 ]+?\s*\d+\s*(?:축적|저장)$',
        # 제거/삭제/해제 시점 표기 : '마지막 탄환 공격 시 제거', '풀버스트 진입 시 제거', '변신 삭제 시 해제'
        # ('시/후' 바로 뒤에 동작어만 오는 경우. 공용 표현 '1인 공격 대상에서 제외 직접 피격 시 해제'는 isOptionText에서 먼저 걸러냄)
        r'^.+(?:시|후)\s*(?:제거|삭제|해제|초기화)$',
        # '상태에 맞는 바니 모드만 적용' 같은 상태별 분기 안내 문구
        r'^상태에\s*맞는\s*.+만\s*적용$',
        # 단순 해제/삭제/제거 표현('디스트로이 해제', '급소 공략 삭제', '폐품 제거') : 최종 목록에 넣지 않는다
        # (수량이 붙는 '해로운 효과 해제 1개', '탄환 100% 제거'는 해당 없음. 공용 표현은 isOptionText에서 먼저 걸러짐)
        r'^[^\d%▲▼]+\s(?:해제|삭제|제거)$',
        # 적용 횟수/지속 기간 안내 문구 : '재장전 1회 적용', '풀 버스트 타임 동안 지속'
        r'^재장전\s*\d+회\s*적용$',
        r'^유지$',   # '부활 시 유지'는 헤더 문장 '부활 시'를 떼고 남은 '유지'만 옵션으로 본다
        r'^풀\s*버스트\s*타임\s*동안\s*지속$',
        r'.*별\s*(?:추가\s*)?효과$',
    ]

    # 헤더 슬롯 분해(종류 기반)의 경계 어미. 서브 헤더 탐지에도 재사용한다.
    BOUNDARY_ENDING = r'(?:(?<![무임])시|후|마다|(?<!명)(?<!아군 )(?<!적 )중|라면|다면|경우|확률로|이면|으면|했다면|한하여|때(?!\s*(?:한하여|마다))|에게|들에게)$'

    # 고유 속성 규칙 6(효과 문구 내부에 시점+대상이 내장된 경우) 판정에 쓰는 시점/조건성 어미.
    # BOUNDARY_ENDING과 달리 '에게'/'들에게'는 제외한다 — 대상 어미는 별도로 뒤쪽에서 확인해야 하기 때문.
    EMBEDDED_TIMING_TOKENS = r'(?:시|후|마다|중|라면|다면|경우|확률로|이면|으면|했다면|한하여)'

    # 고유 속성 제외 규칙 2 : 공용 표현 + 시스템 제어어 결합은 고유 속성이 아닌 공용 표현 파생어로 본다.
    CONTROL_SUFFIXES = ['면역', '해제', '제거', '방해']

    # 본문 요소 판정 규칙 4(이름+번호 라벨)에서 라벨 자체가 범용 분기어일 때는 '추가 효과'와 같은
    # 이유로 명칭 후보로 삼지 않는다(예: 민트 '상태 1 :'/'상태 2 :' — 진짜 기믹 이름이 아니라 case 1/2일 뿐).
    GENERIC_BRANCH_LABELS = ['상태', '경우']

    # 치환 우선순위(최장 일치) 동점 시 우선순위. 숫자가 작을수록 우선.
    SUBSTITUTION_PRIORITY = {
        '[고유 속성]': 0,
        '[스킬 1]': 1, '[스킬 2]': 1, '[버스트 스킬]': 1,
        '[특정 니케]': 2,
    }

    # 스킬 종류 -> 치환 토큰
    SKILL_SLOT_TOKENS = {'스킬 1': '[스킬 1]', '스킬 2': '[스킬 2]', '버스트 스킬': '[버스트 스킬]'}

    # 헤더 슬롯 분해(종류 기반)/선행 행동 : 헤더를 경계 어미 기준으로 절 단위로 쪼개는 패턴.
    # 각 절은 (시/후/마다/중/라면/다면/경우/확률로/이면/으면/했다면/한하여/에게/들에게) 중 하나로 끝나고
    # 바로 뒤에 공백이나 문자열 끝이 와야 한다('아군 중'/'적 중'의 '중'은 "~ 가운데"라는 범위 표현이라 어미가 아니다.
    # '명중'의 끝글자 '중', '무시'/'임시'의 끝글자 '시'도 어미로 안 본다.
    # 예: '명중 대상에게', '방어력 무시 대미지로 적용', '임시 개조' —
    # 또 '시와'처럼 어미 뒤에 글자가 바로 붙으면 경계가 아님 —
    # 라피 : 레드 후드의 '전투 시작 시와 풀 버스트 타임 종료 시'가 안 쪼개지는 이유).
    HEADER_CLAUSE_PATTERN = re.compile(
        r'(.+?)((?<![무임])시|후|마다|(?<!명)(?<!아군 )(?<!적 )중|라면|다면|경우|확률로|이면|으면|했다면|한하여|때(?!\s*(?:한하여|마다))|들에게|에게)(?=\s|$)'
    )
    PRECEDING_ACTION_ENDING = '후'
    TIMING_ENDINGS = {'시', '마다', '중'}
    # 발동 시점이 'A 시와 B 시'로 이어진 경우(예: '전투 시작 시와 풀 버스트 타임 종료 시')는 위 패턴으로는 한 절이지만
    # 시점 두 개이므로 decomposeHeader가 A/B로 쪼갠다. 뒤 절에서 '시'를 뺀 부분이 한 단어뿐이면('랩쳐 등장 시와 사망 시'의 '사망 시') 앞 절의
    # 주어를 공유한다는 뜻이라 앞 절의 주어를 붙여 쪼갠다('랩쳐 사망 시'). 한 블록에 시점이 여럿이면 줄바꿈(TIMING_JOINER)으로 이어 저장한다.
    TIMING_CONJUNCTION_PATTERN = re.compile(r'^(.+?시)와 (.+시)$')
    TIMING_JOINER = '\n'
    CONDITION_ENDINGS = {'라면', '다면', '경우', '확률로', '이면', '으면', '했다면', '한하여', '때'}
    TARGET_ENDINGS = {'에게', '들에게'}

    # 숫자 치환(7단계) : 숫자를 뽑아내는 패턴과, 스켈레톤에서 숫자 자리를 표시하는 자리표시자
    # (게임 텍스트에 나올 일이 없는 제어 문자를 써서 원본 텍스트와 절대 충돌하지 않게 한다).
    NUMBER_PATTERN = re.compile(r'\d+(?:\.\d+)?')
    NUMBER_PLACEHOLDER = '\x00'

    # 토큰 키워드 제외(9단계) : 이 토큰이 하나라도 들어있으면 대개 그 니케만 쓰는 표현이라 제외한다.
    EXCLUSION_TOKENS = ['[고유 속성]', '[스킬 1]', '[스킬 2]', '[버스트 스킬]', '[특정 니케]']

    # 효과 카테고리 분류표(10단계) : 공용 표현과 같은 성격의 100% 수동 관리 입력. 원본은 여기이며
    # 수정은 이 목록에서 한다. 키는 7단계(숫자 치환)까지 끝난 효과 키워드 문자열 그대로여야 한다.
    # 분류표에 없는 새 키워드는 '미분류'로 출력되므로 그때그때 아래 목록에 추가한다.
    # 분류 기준은 '스킬 키워드 분석 규칙.md'의 8대 표준 카테고리 설명·예시를 따르고, 규칙 문서가 침묵하는
    # 항목은 기존(Gemini 시절) 마스터 시트의 분류를 시드로 이어받았다.
    EFFECT_CATEGORIES = {
        '공격류': [
            'n회 순차 공격', '공격 대미지 n% ▲', '공격 속도 n% ▲', '공격력 n% ▲', '공격력 n% ▼', '관통 대미지 n% ▲', '관통 특화',
            "대상 설명이 '~ 적 1기에게'로 끝나는 버스트 스킬 대미지 38.91% ▲", "대상 설명이 '적 전체에게' 로 끝나는 버스트 스킬 대미지 n% ▲", '발사체 부착 대미지 n% ▲',
            '발사체 폭발 대미지 n% ▲', '방어력 무시 대미지 n% ▲', '보호막에 가하는 대미지 n% ▲', '분배 대미지', '분배 대미지 n% ▲', '순차 공격 대미지 158.4% ▲',
            '시전자 기준 공격력 8.94% ▼', '시전자 기준 공격력 n% ▲', '시전자 기준 차지 속도 n% ▲', '시전자의 최종 최대 체력의 10%를 공격력으로 합산한 1372.8% 대미지',
            '시전자의 최종 최대 체력의 n%만큼 공격력 ▲', '우월 코드 공격 대미지 n% ▲',
            '유지 시간동안 해당 적의 받는 대미지 일괄 누적. 유지 시간 만료 후 누적 대미지만큼 적 전체에 분배 대미지. 대미지 최대 누적량은 시전자 최종 공격력의 8900.83%만큼 가능',
            '일반 공격 대미지 40% 배율 ▼', '일반 공격 대미지 n% 배율 ▲', '일반 공격 대미지 배율 100% ▲', '일반 공격 크리티컬 확률 n% ▲',
            '잃은 체력 1% 당 공격력 0.96% ▲', '자신의 공격 대미지의 50%만큼 누적', '자신이 가한 피해량의 58.99%만큼 본체에 고정 대미지', '저지 부위 공격 대미지 n% ▲',
            '지속 대미지 n% ▲', '차지 대미지 9.59% 배율 ▲', '차지 대미지 n% ▲', '차지 속도 n% ▲', '차지 속도 n% ▼', '차지 시간 0.18초 ▼',
            '차지 시간 n초로 고정', '초과된 값의 240% 차지 대미지 ▲', '최대 대미지 누적량은 시전자 최종 공격력의 1536%',
            '최종 공격력 492.3% 방어력 무시 버스트 스킬 대미지', '최종 공격력 833.79% 코어 명중 대미지', '최종 공격력 n% 대미지', '최종 공격력 n% 방어력 무시 대미지',
            '최종 공격력 n% 버스트 스킬 대미지', '최종 공격력 n% 분배 대미지', '최종 공격력 n% 지속 대미지', '최종 공격력 n% 추가 대미지',
            '최종 최대 장탄 수 1발 당 차지 대미지 2.01% ▲', '최종 최대 체력의 5%를 공격력으로 환산한 1150.84% 대미지', '코어 대미지 n% ▲', '크리티컬 대미지 n% ▲',
            '크리티컬 확률 n% ▲', '파츠 대미지 n% ▲', '해당 발사체에 가하는 대미지 n% ▲',
        ],
        '공격보조류': [
            '강제 재장전', '공격력이 가장 높은 아군 1기의 공격력 복제 8.81%', '관통 범위 n% 확장', '명중률 n% ▲', '명중률 n% ▼', '버스트 게이지 충전 n%',
            '버스트 게이지 충전 속도 n% ▲', '사격 집중', '우월 코드 대미지 적용', '일반 공격을 방어력 무시 대미지로 변경', '일반 공격을 방어력 무시 대미지로 적용',
            '장탄 수 무한', '재장전 비율 50% ▼', '재장전 속도 50% 감소 상태로 고정', '재장전 속도 n% ▲', '재장전 속도 n% 증가 상태로 고정', '재장전 시간 3초로 고정',
            '초과된 값을 차지 대미지로 변환', '최대 장탄 수 n% ▲', '최대 장탄 수 n% ▼', '최대 장탄 수 n발 ▲', '탄환 n% 제거', '탄환 소모 40발', '탄환 충전 n%',
            '탄환 충전 n발', '펠릿 개수 1개로 고정', '펠릿 개수 n개 ▲', '폭발 범위 n% ▲', '풀 버스트 타임 지속 시간 5초 ▼', '풀 버스트 타임 지속 시간 n초 ▲',
        ],
        '방어생존류': [
            '1인 공격 대상에서 제외 직접 피격 시 해제', '다음 보호막 체력 n% ▲', '도발', '무적', '받는 대미지 n% ▲', '받는 대미지 n% ▼', '받는 대미지 균등 분배',
            '받는 대미지 차등 분배', '방어력 8.15%', '방어력 n% ▲', '방어력 n% ▼', '불굴', '수냉 코드 적으로부터 받는 대미지 17.95% ▼',
            '시전자 기준 방어력 1.9% ▼', '시전자 기준 방어력 n% ▲', '시전자 기준 최대 체력 n% ▲', '시전자 기준 최대 체력만 n% ▲',
            '시전자 방어력 비례 20.9% 방어력 ▲', '시전자 최종 최대 체력 비례 10.08% 보호막', '시전자의 최대 체력 비례 엄폐물 최대 체력 n% ▲',
            '시전자의 최종 최대 체력 비례 20.59% 공용 보호막', '시전자의 최종 최대 체력 비례 n% 보호막', '시전자의 최종 최대 체력 비례 n% 아군 공용 보호막',
            '시전자의 최종 최대 체력 비례 최대 체력 27.87% ▲', '엄폐물 방어력 128.57% ▲', '잃은 체력 1% 당 방어력 3.51% ▲', '자신이 설치한 보호막 무적',
            '적 전체 도발', '적 전체를 도발', '전격 코드 적으로부터 받는 대미지 70.4% ▼', '최대 체력 n% ▲', '최대 체력만 n% ▲',
            '최대 체력이 가장 높은 니케의 최대 체력 복제 12.42%', '최대 체력이 가장 높은 아군 1기의 최대 체력 복제 15.03%', '최종 최대 체력 비례 40.12% 보호막',
            '풍압 코드 적으로부터 받는 대미지 90.72% ▼', '현재 체력 n% ▼',
        ],
        '체력회복류': [
            '공격 대미지 비례 n% 회복', '받는 체력 회복량 78.93% ▼', '받는 체력 회복량 n% ▲', '시전자 기준 최대 체력 n%까지 초과 받는 체력 회복량 저장',
            '시전자 최대 체력 비례 n% 회복', '시전자 최종 최대 체력 비례 n% 회복', '시전자의 최종 최대 체력 비례 n% 지속 회복', '시전자의 최종 최대 체력 비례 n% 회복',
            '시전자의 최종 최대 체력 비례 디코이 지속 회복 2.27%', '시전자의 최종 최대 체력 비례 디코이 회복 2.1%', '시전자의 최종 최대 체력 비례 보호막 체력 n% 회복',
            '시전자의 최종 최대 체력 비례 보호막 체력 회복 2.85%', '시전자의 최종 최대 체력 비례 엄폐물 체력 지속 회복 2.88%',
            '시전자의 최종 최대 체력 비례 엄폐물 체력 회복 21.41%', '엄폐물 체력 회복 n%', '주는 체력 회복량 10% ▼', '주는 체력 회복량 n% ▲', '체력 회복 균등 분배',
            '체력 회복량 30.96% ▲', '최대 충전량은 시전자 최종 공격력의 165.28%', '충전량 회복',
        ],
        '유틸리티류': [
            '공격 범위 중심으로 강제 이동', '기절', '기절 면역', '버스트 1단계로 변경', '버스트 스킬 재사용 시간 21초 ▲', '버스트 스킬 재사용 시간 n초 ▼',
            '버스트 재진입 1단계로 변경', '버스트 재진입 n단계', '스킬 1의 풀 차지 공격 횟수 조건이 1회/2회/3회 로 변경', '스킬 2 강제 사용',
            '스킬 2 명중 횟수 조건 20회 ▼', '스킬 2 재사용 시간 18초 ▼', '스킬 2 재사용 시간 n% ▼', '시전자의 최종 최대 체력 비례 n% 분신',
            '엄폐물 체력 n%로 엄폐물 부활', '이동 불가', '이로운 효과 해제 5개', '적정 최대 사거리 20% ▲', '적정 최소 사거리 n% ▲',
            '중첩 가능 이로운 효과 중첩량 1개 ▲', '중첩 가능 해로운 효과 중첩량 1개 ▼', '중첩량 증감 효과 면역', '차지 속도 감소 효과 면역', '차지 속도 증가 효과 면역',
            '체력 n%로 부활', '해로운 효과 면역 1개', '해로운 효과 면역 ∞개', '해로운 효과 해제 1개',
        ],
    }

    def __init__(self, path_db):
        self.current_file = Path(__file__).resolve()
        self.current_dir = self.current_file.parent
        self.path_excel = self.current_dir / "nikke_skill_keyword.xlsx"

        self.db_connection = sqlite3.connect(path_db)
        self.db_cursor = self.db_connection.cursor()

        # 복합 상태명(고유 속성 규칙 7) 집합 : 3단계 시작 전에 전체 설명을 훑어 채운다(buildSkillBlockSheet).
        self.compound_state_names = set()

    def __del__(self):
        try:
            self.db_connection.close()
        except Exception:
            pass

    def createExcelWorkbook(self):
        '''
        [엑셀 뼈대 준비]
        nikke_skill_keyword.xlsx를 전체 새로 생성한다(매번 전체 재생성 원칙 — 기존 파일이 있어도 덮어씀).
        '''
        try:
            wb = openpyxl.Workbook()
            default_sheet = wb.active

            # db를 건들지 않으면서 니케 스킬 설명 후처리용(2단계 산출물)
            sheet_info = wb.create_sheet('니케 스킬 설명')
            sheet_info.append(['니케 id', '니케 이름', '티어', '스킬 1 이름', '스킬 1 설명',
                                '스킬 2 이름', '스킬 2 설명', '버스트 스킬 이름', '버스트 스킬 설명'])

            # 스킬 키워드에서 고유 속성으로 필터링되지 않게 하기 위한 게임 공용 표현 시트(100% 수동, 원본은 COMMON_EXPRESSIONS)
            sheet_common = wb.create_sheet('공용 표현')
            sheet_common.append(['표현목록'])
            for expr in self.COMMON_EXPRESSIONS:
                sheet_common.append([expr])

            # 구조 분해 산출물(3단계) — 블록 단위로 효과/부가 옵션/명칭을 나눠 기록.
            # '치환된 헤더'/'치환된 내용'은 5단계(텍스트 치환), '선행 행동'/'발동 시점'/'발동 조건'/'적용 대상'은
            # 6단계(헤더 슬롯 분해) 산출물 — 3단계에서는 비워두고 각 단계에서 채움.
            sheet_block = wb.create_sheet('스킬 블록')
            sheet_block.append(['니케 id', '니케 이름', '스킬 종류', '블록 번호', '헤더', '종류', '내용', '순번',
                                 '치환된 헤더', '치환된 내용',
                                 '선행 행동', '발동 시점', '발동 조건', '적용 대상'])

            # 스킬 설명에 니케 고유 매커니즘을 필터링하기 위한 키워드 시트(4단계 산출물)
            sheet_unique = wb.create_sheet('고유 속성')
            sheet_unique.append(['니케 이름', '등급', '고유 속성'])

            # 스킬 키워드 마스터에 기록하기 전, 디버깅용으로 스킬 키워드 목록을 기록하는 시트(8단계 산출물)
            sheet_debug = wb.create_sheet('스킬 키워드 디버깅 목록')
            # '효과 카테고리'는 10단계 산출물 — 5번째 열(효과)과 같은 행에 대응(효과 열이 독립 목록이라 가능)
            sheet_debug.append(['발동 시점', '발동 조건', '선행 행동', '적용 대상', '효과', '효과 카테고리'])

            # 고유 속성은 아니지만 한 니케 전용인 이름 목록(4단계 산출물) — 9단계가 이 이름이 든 키워드를 제외한다
            sheet_own_names = wb.create_sheet('전용 이름')
            sheet_own_names.append(['니케 이름', '등급', '전용 이름', '근거'])

            # 토큰/선행 행동이 든 키워드를 제외하면서 사유를 남기는 시트(9단계 산출물)
            sheet_excluded = wb.create_sheet('제외 키워드')
            sheet_excluded.append(['슬롯', '키워드', '제외 사유'])

            # 유사 표현 통합 시트(11단계 산출물) — db의 skill_keyword_unify_table에서 만들고, 승환님이 수작업으로 편집(가로형).
            # 비슷한 표현 열 개수는 실제 그룹 크기에 맞춰 buildUnifySheet가 다시 쓴다.
            sheet_unify = wb.create_sheet('키워드 통합')
            sheet_unify.append(['열', '통합 표현', '비슷한 표현 1'])

            # 통합 시트와 현재 키워드 목록이 어긋난 곳을 알리는 점검 시트(11단계 산출물)
            sheet_unify_check = wb.create_sheet('통합 점검')
            sheet_unify_check.append(['열', '구분', '표현', '비고'])

            # 실질적인 스킬 키워드 목록 시트(11단계 산출물)
            sheet_master = wb.create_sheet('스킬 키워드 마스터')
            sheet_master.append(['발동 시점', '발동 조건', '적용 대상', '효과 - 공격류', '효과 - 공격보조류',
                                  '효과 - 방어생존류', '효과 - 체력회복류', '효과 - 유틸리티류'])

            # 니케와 마스터 키워드를 잇는 링크 시트(12단계 산출물) — 키워드 id는 13단계(db 적재)에서 정해지므로 문자열로 적는다
            sheet_link = wb.create_sheet('니케-키워드 링크')
            sheet_link.append(['니케 id', '니케 이름', '열', '키워드', '출처 스킬'])

            # 파이프라인 전체 검증 결과 시트(14단계 산출물) — 오류/경고와 내용
            sheet_validation = wb.create_sheet('검증 결과')
            sheet_validation.append(['구분', '항목', '내용'])

            # 마지막 db 반영(RUN_APPLY) 이후 바뀐 니케/키워드/링크 목록(RUN_WORKBOOK 마지막 단계 산출물) — 색으로도 표시
            sheet_change = wb.create_sheet(self.CHANGE_SHEET_NAME)
            sheet_change.append(self.BASELINE_HEADER)

            # 마지막 db 반영 시점의 마스터/링크/스킬 설명(RUN_APPLY가 쓰고, RUN_WORKBOOK이 다음 엑셀로 넘겨 비교 기준으로 씀)
            sheet_baseline = wb.create_sheet(self.BASELINE_SHEET_NAME)
            sheet_baseline.append(self.BASELINE_HEADER)

            wb.remove(default_sheet)
            wb.save(self.path_excel)
            print(f"[SUCCESS] createExcelWorkbook 완료: {self.path_excel.name}")
            return True
        except Exception as error_msg:
            print(f"[ERROR] createExcelWorkbook 실행 중 오류 발생: {error_msg}")
            return False

    def normalizeSpacing(self, text):
        '''
        [DB 구축 2단계: 띄어쓰기 통일]
        정규식 규칙만 사용한다(사전 없음 — '스킬 키워드 분석 규칙.md' "띄어쓰기 통일 규칙" 절 참고).
        '''
        if not text:
            return ""
        try:
            result_text = text
            result_text = re.sub(r'[ \t]+', ' ', result_text)              # 연속 공백/탭 정리
            # 아래 규칙들의 공백 매칭은 전부 [ \t]로 한정한다 — \s는 줄바꿈까지 먹어서 별도 줄이던 대괄호들이
            # 한 줄로 합쳐지고(예: 에이다의 '[특수 개조]'가 앞 효과 줄에 붙어버림), 줄 단위로 판정하는 3단계
            # (이름 단독 대괄호/스펙 시트 판정 등)가 틀어지는 버그가 있었다.
            result_text = re.sub(r'\][ \t]*\[', '] [', result_text)           # ...][ ... -> ...] [ ...
            result_text = re.sub(r'(\d+)[ \t]*중첩', r'\1 중첩', result_text)  # 30중첩 -> 30 중첩
            result_text = re.sub(r'(%)[ \t]*([▲▼])', r'\1 \2', result_text)   # 1.4%▲ -> 1.4% ▲
            # '에게'/'들에게' 뒤에 괄호로 부연 설명이 붙는 경우(공백 유무 둘 다 존재) 통째로 제거.
            # 부연 설명이라 필터링에 불필요하고, 괄호 유무로 표현이 갈리면 같은 대상이 다른 키워드로
            # 갈라지므로 아예 없는 표현으로 통합한다(예: '적 전체에게(파츠 포함)' -> '적 전체에게').
            # 부수 효과: 이 괄호 때문에 '~에게'로 안 끝나 보여서 서브 헤더로 인식 못 하던 문제도 해결됨
            # (신데렐라 : 크리스탈 웨이브의 '자신이 저격 모드라면 적 전체에게 (파츠 포함)'에서 확인).
            result_text = re.sub(r'(에게|들에게)[ \t]*\([^)\n]*\)', r'\1', result_text)
            # 표기 통일(데이터 전수 조사로 확인된 갈림만 정규식 규칙으로 처리)
            #  - '스킬2' -> '스킬 2' (라피 : 레드 후드, 아르카나만 붙여 쓰고 나머지는 전부 띄어 씀)
            #  - '수▲' -> '수 ▲' (솔린 : 프로스트 티켓의 '티켓 수▲'만 붙여 씀)
            #  - 중첩 대괄호 평탄화 : '[최대 [록 온] 대상 수 10 ▲]' -> '[최대 록 온 대상 수 10 ▲]' (대괄호 파싱이 중첩을 못 다룸)
            #  - 대괄호 뒤에 붙는 부연 설명 괄호 제거 : '[10초 유지] (동일 효과 중복 불가)' -> '[10초 유지]'
            #  - '마다'는 앞말에 붙여 쓴다. '2초 마다' 10건/'2초마다' 4건, '될 때 마다'/'될 때마다'가 섞여 있어 같은 시점이
            #    다른 키워드로 갈라졌다(승환님 결정: 붙여 쓰는 쪽으로 통일)
            result_text = re.sub(r'(?<=\S)[ \t]+마다(?=\s|$)', '마다', result_text)
            #  - '풀버스트'/'풀차지' -> '풀 버스트'/'풀 차지' (각각 3건/2건만 붙여 씀), '사용 했다면' -> '사용했다면'(1건),
            #    '적들 에게' -> '적들에게'(1건), '% 만큼' -> '%만큼'(5건 대 4건, 승환님 결정: 단위 뒤는 붙여 씀)
            result_text = re.sub(r'풀(버스트|차지)', r'풀 \1', result_text)
            result_text = result_text.replace('사용 했다면', '사용했다면').replace('적들 에게', '적들에게')
            result_text = re.sub(r'%[ \t]+만큼', '%만큼', result_text)
            #  - '% 까지' -> '%까지' ('%만큼'과 같은 기준으로 단위 뒤는 붙임; 2건 대 1건), '유지시간' -> '유지 시간'(1건 대 다수)
            result_text = re.sub(r'%[ \t]+까지', '%까지', result_text)
            result_text = result_text.replace('유지시간', '유지 시간')
            result_text = re.sub(r'스킬(\d)', r'스킬 \1', result_text)
            result_text = re.sub(r'([가-힣])([▲▼])', r'\1 \2', result_text)
            while True:
                flattened_text = re.sub(r'\[([^\[\]]*)\[([^\[\]]*)\]([^\[\]]*)\]', r'[\1\2\3]', result_text)
                if flattened_text == result_text:
                    break
                result_text = flattened_text
            result_text = re.sub(r'(\])[ \t]*\([^)\n]*\)', r'\1', result_text)
            # 말줄임표 통일 : 같은 이름이 효과 쪽에서는 '…'(한 글자), 헤더 쪽에서는 '...'(점 3개)로 갈리는 경우가
            # 있다(길티 '빌려 갈게에….'가 효과의 [이름 : 효과]에는 '….', 헤더에는 '....'로 표기). 점 3개로 통일.
            result_text = result_text.replace('…', '...')
            return result_text.strip()
        except Exception as error_msg:
            print(f"[ERROR] normalizeSpacing 중 오류: {error_msg}")
            return text.strip()

    def buildDescriptionSheet(self):
        '''
        [DB 구축 2단계]
        nikke_info_table의 스킬 설명들을 띄어쓰기 통일한 뒤 '니케 스킬 설명' 시트에 저장.
        이후 단계는 전부 이 시트의 값을 입력으로 사용한다(DB를 다시 조회하지 않음).
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_info = workbook['니케 스킬 설명']

            self.db_cursor.execute(
                "SELECT id, Name, Tier, Skill_1_name, Skill_1_description, "
                "Skill_2_name, Skill_2_description, Skill_bust_name, Skill_bust_description "
                "FROM nikke_info_table ORDER BY id ASC"
            )
            nikke_rows = self.db_cursor.fetchall()

            for (nikke_id, name, tier, s1_name, s1_desc, s2_name, s2_desc, sb_name, sb_desc) in nikke_rows:
                sheet_info.append([
                    nikke_id, name, tier,
                    s1_name, self.normalizeSpacing(s1_desc),
                    s2_name, self.normalizeSpacing(s2_desc),
                    sb_name, self.normalizeSpacing(sb_desc),
                ])

            workbook.save(self.path_excel)
            print(f"[SUCCESS] buildDescriptionSheet 완료: {len(nikke_rows)}명 처리")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildDescriptionSheet 실행 중 오류 발생: {error_msg}")
            return False

    def isOptionText(self, text):
        '''
        [3단계 보조] 대괄호 안 텍스트가 부가 옵션(키워드 대상 아님)인지 판정
        '''
        stripped_text = text.strip()
        if stripped_text in self.COMMON_EXPRESSIONS:
            return False
        for pattern in self.OPTION_PATTERNS:
            if re.match(pattern, stripped_text):
                return True
        return False

    # 복합 상태명 뒤에 붙을 수 있는 동작어('[무대 파트 : 보컬 삭제]'의 '삭제' 등). 이 한 단어만 떼고 비교한다.
    COMPOUND_ACTION_TAILS = ['삭제', '해제', '제거', '적용', '초기화', '변경', '지정']

    def collectCompoundStateNames(self, description_texts):
        '''
        [3단계 보조: 고유 속성 규칙 7] 복합 상태명(`상위개념 : 하위개념`) 목록을 전체 설명에서 모은다.
        대괄호 `[A : B ...]`의 `A : B`가 대괄호 바깥(헤더 등) 어디에선가 글자 그대로 언급되면 복합 상태명으로 본다
        (예: `[무대 파트 : 보컬]`/`[무대 파트 : 보컬 삭제]` ↔ 헤더 '자신이 무대 파트 : 보컬 상태라면').
        바깥에서 언급되지 않는 `[이름 : 효과]`(예: `[딸기 사탕: 최대 장탄 수 56.7% ▲]`)는 평범한 명칭:효과 쌍이므로 제외.
        '''
        outside_parts = []
        bracket_contents = []
        for description_text in description_texts:
            if not description_text:
                continue
            outside_parts.append(re.sub(r'\[[^\[\]]*\]', '\n', description_text))
            bracket_contents.extend(re.findall(r'\[([^\[\]]+)\]', description_text))
        outside_corpus = re.sub(r'[ \t]*[:：][ \t]*', ' : ', '\n'.join(outside_parts))

        compound_names = set()
        for bracket_inner in bracket_contents:
            if ':' not in bracket_inner:
                continue
            left_text, right_text = (part.strip() for part in bracket_inner.split(':', 1))
            if not left_text or not right_text:
                continue
            if left_text in self.COMMON_EXPRESSIONS or right_text in self.COMMON_EXPRESSIONS:
                continue
            right_words = right_text.split()
            candidate_rights = [' '.join(right_words)]
            if len(right_words) >= 2 and right_words[-1] in self.COMPOUND_ACTION_TAILS:
                candidate_rights.append(' '.join(right_words[:-1]))
            for candidate_right in candidate_rights:
                candidate_name = f'{left_text} : {candidate_right}'
                if candidate_name in outside_corpus:
                    compound_names.add(candidate_name)
                    break
        return compound_names

    def findCompoundName(self, bracket_inner):
        '''
        [3단계 보조] 대괄호 안 텍스트가 복합 상태명(이나 '복합 상태명 + 동작어')이면 그 복합 상태명을,
        아니면 None을 돌려준다.
        '''
        if ':' not in bracket_inner:
            return None
        left_text, right_text = (part.strip() for part in bracket_inner.split(':', 1))
        normalized = f'{left_text} : {right_text}'
        for compound_name in self.compound_state_names:
            if normalized == compound_name:
                return compound_name
            if normalized.startswith(compound_name + ' ') and \
                    normalized[len(compound_name):].strip() in self.COMPOUND_ACTION_TAILS:
                return compound_name
        return None

    def matchesCompoundName(self, bracket_inner):
        '''[3단계 보조] 대괄호 안 텍스트가 복합 상태명(이나 '복합 상태명 + 동작어')으로 시작하는지 판정.'''
        return self.findCompoundName(bracket_inner) is not None

    def isPureOption(self, text):
        '''
        [3단계 보조] 부가 옵션이면서 앞에 헤더 문장이 붙어 있지도 않은 문구. `버스트 스킬 사용 시 중첩 제거`처럼
        `<헤더 문장> <단순 제거/삭제/해제>`는 문구 전체를 옵션으로 버리면 헤더 문장(발동 시점)까지 같이 사라지므로
        일단 효과로 남기고, expandEmbeddedHeaders가 헤더 문장을 서브 헤더로 떼어낸 뒤 남은 부분만 옵션으로 바꾼다.
        '''
        if not self.isOptionText(text):
            return False
        if self.splitEmbeddedHeader(text)[0] is None:
            return True
        # 헤더 문장처럼 보여도 문구 전체가 고정 옵션이면(`전투 중 1회 발동`의 '전투 중') 옵션이다. 헤더 문장을 떼어
        # 내야 하는 건 옵션 판정이 "단순 제거/삭제/해제/유지" 패턴에서 나온 경우뿐이다.
        return not re.search(r'(?:해제|삭제|제거|초기화|유지)$', text.strip())

    def splitBracketContents(self, line_text):
        '''
        [3단계 보조] 한 줄에 있는 [ ] 내용을 전부 꺼내 (종류, 내용) 목록으로 반환
        - 콜론이 있으면 왼쪽을 명칭, 오른쪽을 효과로 나눈다(예: '딸기 사탕: 최대 장탄 수 56.7% ▲')
          단, 복합 상태명(`[무대 파트 : 보컬 삭제]` 등)이면 콜론에서 쪼개지 않고 통째로 효과 하나로 둔다
        - 콜론이 없으면 부가 옵션 여부를 먼저 판정하고, 아니면 효과로 취급
        '''
        results = []
        for bracket_inner in re.findall(r'\[([^\[\]]*)\]', line_text):
            bracket_inner = bracket_inner.strip()
            if not bracket_inner:
                continue
            if ':' in bracket_inner and self.matchesCompoundName(bracket_inner):
                left_text, right_text = (part.strip() for part in bracket_inner.split(':', 1))
                results.append(('효과', f'{left_text} : {right_text}'))
                continue
            if ':' in bracket_inner:
                left_text, right_text = bracket_inner.split(':', 1)
                left_text = left_text.strip()
                right_text = right_text.strip()
                if left_text:
                    results.append(('명칭', left_text))
                if right_text:
                    if self.isPureOption(right_text):
                        results.append(('부가옵션', right_text))
                    else:
                        results.append(('효과', right_text))
            else:
                if self.isPureOption(bracket_inner):
                    results.append(('부가옵션', bracket_inner))
                else:
                    results.append(('효과', bracket_inner))
        return results

    def isBracketOnlyLine(self, line_text):
        '''
        [3단계 보조] 줄 전체가 대괄호(들)로만 이루어져 있는지(뒤에 콜론 하나만 허용)
        '''
        return bool(re.fullmatch(r'(?:\s*\[[^\[\]]+\]\s*)+:?\s*', line_text.strip()))

    def looksLikeSpecSheetStart(self, body_lines, start_index):
        '''
        [3단계 보조: 본문 요소 판정 규칙 3] start_index부터 훑으며, 대괄호로만 이루어진 줄(다른 명칭성 마커)은
        건너뛰고 처음 만나는 '실질적인' 줄이 `기능 :` 또는 SPEC_LABELS로 시작하는지 확인한다.
        이 판정이 참일 때만 직전의 명칭 단독 대괄호(들)를 "명칭 단독 대괄호 스펙 시트"로 취급한다.
        '''
        for idx in range(start_index, len(body_lines)):
            candidate = body_lines[idx].strip()
            if self.isBracketOnlyLine(candidate):
                continue
            # `기능 :` 뿐 아니라 `효과 :`/`효과 N :`이 이어질 때도 직전 대괄호는 이름 단독 라벨이다
            # (예: 엠마 : 택티컬 업 '[노출] [해제 불가]' 다음 줄 '효과 : [주목 : 적 전체 도발]').
            if re.match(r'^(?:기능|효과\s*\d*)\s*[:：]', candidate):
                return True
            for spec_label in self.SPEC_LABELS:
                if re.match(rf'^{re.escape(spec_label)}\s*[:：]', candidate):
                    return True
            return False
        return False

    def firstBracketLooksLikeEffect(self, line_text):
        '''
        [3단계 보조] 줄의 첫 대괄호가 수치 효과(▲/▼ 또는 n%)이면 True. 이런 줄은 뒤에 `기능 :`이 이어져도 이름 라벨이 아니라
        그냥 효과 줄이다(예: 스노우 화이트 : 헤비암즈 버스트의 `[공격 대미지 84.48% ▲] [10초 유지]` 바로 다음 줄이
        이름 단독 대괄호 `[세븐스 드워프 풀 액티브]`) — 앞 줄까지 이름으로 삼으면 효과가 통째로 사라진다.
        '''
        first_match = re.search(r'\[([^\[\]]+)\]', line_text)
        if not first_match:
            return False
        first_group = first_match.group(1)
        return bool(re.search(r'[▲▼]|\d+(?:\.\d+)?\s*%', first_group))

    def followedByLabeledRows(self, body_lines, start_index):
        '''
        [3단계 보조] start_index부터 처음 만나는 '대괄호로만 된 줄이 아닌' 줄이 `라벨 : [효과]` 행(라벨이 숫자로
        시작하지 않는 것)인지 확인한다 — 직전의 대괄호가 표 머리글인지 판정하는 데 쓴다.
        '''
        # 바로 다음 줄만 본다 — 대괄호 줄들을 건너뛰며 찾으면 '[효과] [옵션] :' 뒤에 라벨 행이 오는 평범한
        # 구조(예: 이사벨 버스트)에서 앞의 진짜 효과 대괄호까지 표 머리글로 오인해 버린다.
        if start_index >= len(body_lines):
            return False
        candidate = body_lines[start_index].strip()
        row_match = re.match(r'^([^\[\]:：]+?)\s*[:：]\s*\[', candidate)
        if not row_match:
            return False
        # 숫자 라벨('1회 :')과 `효과`/`추가 효과`/`기능`/`대상` 라벨 행은 표 행이 아니라 일반 본문 라벨이다
        return not re.match(r'^(?:\d|효과|추가\s*효과|기능|대상)', row_match.group(1).strip())

    def stripTargetLabel(self, header_text):
        '''
        [3단계 보조] 서브 헤더 앞에 붙은 `대상 :` 라벨을 뗀다(예: 은화 : 택티컬 업 버스트의 `대상 : 명중 대상에게`).
        라벨이 남으면 '명중'의 끝 글자 '중'이 경계 어미로 오인돼 `대상 : 명중`이 발동 시점으로 갈라진다.
        '''
        return re.sub(r'^대상\s*[:：]\s*', '', header_text.strip())

    def splitEmbeddedHeader(self, effect_text):
        '''
        [3단계 보조] 효과 문구 맨 앞에 헤더 문장(시점/조건/선행 행동/대상 절)이 붙어 있으면 (헤더 문장, 나머지 효과)로 나눈다.
        예) '효과 적용 후 사슬 감기 삭제' -> ('효과 적용 후', '사슬 감기 삭제'),
            '유지 시간 만료 후 아군 전체에게 충전량 회복' -> ('유지 시간 만료 후 아군 전체에게', '충전량 회복').
        헤더 문장은 호출부가 서브 헤더(독립 블록)로 승격해 다른 헤더와 똑같이 6단계(헤더 슬롯 분해)를 거치게 한다.
        오탐을 줄이는 가드 : 헤더 문장이 2단어 이상이고(`임시 개조 중첩 복사`의 '임시' 같은 단어 끝 글자 오인 방지),
        여러 문장(`. `)이 아니고, 나머지가 동작어 하나(제거/삭제/해제/초기화)뿐이 아니어야 한다.
        '''
        if '. ' in effect_text:
            return None, effect_text
        clauses = self.tokenizeHeaderClauses(effect_text)
        if len(clauses) < 2 or clauses[-1][1] is not None:
            return None, effect_text
        prefix_clauses = clauses[:-1]
        if any(ending is None for _, ending in prefix_clauses):
            return None, effect_text
        prefix_text = ' '.join(clause_text for clause_text, _ in prefix_clauses)
        remainder_text = clauses[-1][0]
        if len(prefix_text.split()) < 2 or not remainder_text:
            return None, effect_text
        if remainder_text in ('제거', '삭제', '해제', '초기화'):
            return None, effect_text
        return prefix_text, remainder_text

    def expandEmbeddedHeaders(self, blocks):
        '''
        [3단계 보조] 블록들의 '효과' 항목 중 맨 앞에 헤더 문장이 붙은 것(splitEmbeddedHeader)을 찾아, 그 헤더 문장을
        서브 헤더로 하는 새 블록으로 쪼갠다. 그 항목 뒤에 이어지던 항목들은 새 블록에 딸린다.
        '''
        expanded_blocks = []
        for block in blocks:
            current_block = {'header': block['header'], 'items': []}
            in_sub_block = False
            for kind, content in block['items']:
                if kind == '효과':
                    sub_header, body_text = self.splitEmbeddedHeader(content)
                    if sub_header:
                        expanded_blocks.append(current_block)
                        current_block = {'header': sub_header, 'items': []}
                        in_sub_block = True
                        content = body_text
                        if self.isOptionText(content):
                            kind = '부가옵션'   # 단순 제거/삭제/해제는 키워드가 아니므로 헤더 문장(시점/선행 행동)만 남긴다
                    elif in_sub_block:
                        expanded_blocks.append(current_block)
                        current_block = {'header': block['header'], 'items': []}
                        in_sub_block = False
                elif in_sub_block and kind != '부가옵션':
                    # 서브 블록에는 그 효과에 딸린 부가 옵션([지속] 등)만 붙고, 새 효과/명칭부터는 원래 헤더로 돌아간다
                    expanded_blocks.append(current_block)
                    current_block = {'header': block['header'], 'items': []}
                    in_sub_block = False
                current_block['items'].append((kind, content))
            expanded_blocks.append(current_block)
        return expanded_blocks

    def stripWrappingBracket(self, text):
        '''
        [3단계 보조] '[관통 특화]'처럼 텍스트 전체가 대괄호(들)로만 감싸져 있으면 벗겨서 (종류, 내용) 목록으로 반환.
        대괄호로만 이루어져 있지 않으면(앞에 다른 서술이 붙어 있으면) 원문 그대로 효과 하나로 돌려준다
        (효과 N: 규칙 — 대상/시점 서술이 섞여 있으면 분리하지 않는다는 결정과 일치).
        '''
        text = text.strip()
        if not text:
            return []
        if self.isBracketOnlyLine(text):
            return self.splitBracketContents(text)
        return [('효과', text)]

    def isSubHeaderLine(self, line_text):
        '''
        [3단계 보조] 대괄호로 시작하지 않고 경계 어미로 끝나면, 하나의 `■` 안에 있는 별도 헤더(서브 헤더)로 본다.
        접두 라벨(숫자+회/단계+콜론, 이름+번호+콜론)이 붙어 있으면 라벨을 뗀 나머지로 판정한다.
        '''
        text = line_text.strip()
        if not text or text.startswith('['):
            return False
        text = re.sub(r'^\d+\s*(?:회|단계)(?:이하|이상)?\s*[:：]\s*', '', text)
        text = re.sub(r'^[가-힣A-Za-z]+\s*\d+\s*[:：]\s*', '', text)
        if not text or text.startswith('['):
            return False
        return bool(re.search(self.BOUNDARY_ENDING, text))

    def splitPrefixAndBrackets(self, content_text):
        '''
        [3단계 보조] '<접두 서술> [대괄호] [대괄호]...' 형태를 (접두 서술, 대괄호 부분)으로 나눈다.
        접두 서술이 경계 어미(시/후/마다/중/…/에게/들에게)로 끝나는 헤더 문장이고 나머지가 대괄호로만 이루어져 있을 때만
        분리하고, 아니면 (None, 원문)을 돌려준다. 접두 서술은 호출부가 서브 헤더(독립 블록)로 승격해
        다른 헤더와 똑같이 6단계(헤더 슬롯 분해)를 거치게 한다.
        (예: `효과 4 : 적 사망 시 [폐품 2개 ▲] [최대 10개] [지속]` -> 헤더 '적 사망 시' + 대괄호들)
        '''
        content_text = content_text.strip()
        first_bracket = content_text.find('[')
        if first_bracket <= 0:
            return None, content_text
        prefix_text = content_text[:first_bracket].strip()
        rest_text = content_text[first_bracket:].strip()
        if prefix_text and self.isBracketOnlyLine(rest_text) and re.search(self.BOUNDARY_ENDING, prefix_text):
            return prefix_text, rest_text
        return None, content_text

    def tryRescueAddOnLabel(self, line_text):
        '''
        [3단계 보조: 본문 요소 판정 규칙 2/3] '추가 효과'/'효과 N'/'효과 N 대상'만 구제한다.
        - 번호가 있는 이름 라벨(추가 효과 1 등)은 번호를 떼고 라벨 자체는 버림(명칭으로 삼지 않음 — '추가 효과'는 범용 단어라 고유 속성 후보로 부적절)
        - 콜론 없이 단독으로만 나온 '추가 효과'는 구분선이므로 버림([])
        - 뒤에 오는 내용이 대괄호(들)로만 이루어져 있으면 각각 분리, 서술이 섞여 있으면 통째로 효과 하나로 둔다
        - 뒤에 오는 내용이 '헤더 문장 + 대괄호(들)'이면 헤더 문장을 서브 헤더로 분리한다(splitPrefixAndBrackets)
        반환값은 (서브 헤더 또는 None, (종류, 내용) 튜플의 리스트). 구제할 게 없으면 (None, []).
        '''
        text = line_text.strip()

        # '효과 N 대상 : ...' 은 부연설명이라 여기서 구제하지 않는다(규칙 7에서 통째 폐기)
        if re.match(r'^효과\s*\d+\s*대상\s*[:：]', text):
            return None, []

        content = None
        match_effect_n = re.match(r'^효과\s*\d+\s*[:：]\s*(.*)$', text)
        if match_effect_n:
            content = match_effect_n.group(1).strip()
        else:
            match_addon = re.match(r'^추가\s*효과\s*\d*\s*[:：]\s*(.*)$', text)
            if match_addon:
                content = match_addon.group(1).strip()
        if content is None:
            return None, []

        sub_header, rest_text = self.splitPrefixAndBrackets(content)
        return sub_header, self.stripWrappingBracket(rest_text)

    def chunkSkillDescription(self, description_text):
        '''
        [DB 구축 3단계: 구조 분해]
        스킬 설명 텍스트를 블록(헤더, 본문 요소 목록) 리스트로 분해한다.
        본문 요소는 (종류, 내용) 튜플이며 종류는 '효과'/'부가옵션'/'명칭'/'스킬 별칭'(<Step n> 전용, 아래 참고) 중 하나.
        의미 해석 없이 '■'/줄바꿈/'[ ]'/라벨 형태만으로 결정적으로 처리한다.
        ('스킬 키워드 분석 규칙.md'의 "본문 요소 판정 규칙" 절이 이 함수의 근거)
        '''
        blocks = []
        if not description_text:
            return blocks

        try:
            raw_sections = re.split(r'■\s*', description_text)

            # 첫 '■' 이전 텍스트(raw_sections[0])는 보통 빈 내용이라 버리지만, <Step n에서 사용 시 : 이름>이
            # 설명 전체의 맨 첫 줄로 올 때(레드 후드의 버스트 'Step 1')는 이 자리에 와서 그냥 버리면 이름이
            # 통째로 유실된다. 그런 이름만 건져서 뒤이어 나오는 첫 블록에 '스킬 별칭'으로 붙여준다.
            pending_step_aliases = []
            if raw_sections and raw_sections[0]:
                for pre_line in raw_sections[0].split('\n'):
                    pre_line = pre_line.strip()
                    if not pre_line:
                        continue
                    pre_step_match = re.match(r'^<[^<>:：]+[:：]\s*([^<>]+)>$', pre_line)
                    if pre_step_match:
                        pre_step_name = pre_step_match.group(1).strip()
                        if pre_step_name:
                            pending_step_aliases.append(('스킬 별칭', pre_step_name))

            for raw_section in raw_sections[1:]:  # 첫 조각은 '■' 이전 텍스트라 사용 안 함
                lines = [line.strip() for line in raw_section.split('\n') if line.strip()]
                if not lines:
                    continue

                current_header = lines[0]
                current_items = []
                body_lines = lines[1:]
                # 직전 줄이 '명칭 단독 대괄호 스펙 시트'의 스펙 라벨이었는지(규칙 3의 범위를 대신하는 최소 상태)
                # 이 값이 True인 동안 숫자 단계 라벨('1단계이하 :' 등)이 나오면 그 스펙의 세부값으로 보고 같이 폐기한다.
                last_label_was_spec = False

                line_index = 0
                while line_index < len(body_lines):
                    line = body_lines[line_index]
                    line_index += 1

                    # 6. <Step n에서 사용 시 : 이름> 분기 — 콜론 앞은 무시. 콜론 뒤 이름은 '명칭'이 아니라
                    # '스킬 별칭'으로 남긴다 — 이 이름은 [이름:효과] 브래킷처럼 텍스트로 효과를 직접 부여하는
                    # 게 아니라 버스트 스킬의 특정 Step(분기)을 가리키는 것뿐이라, 고유 속성 규칙 3(효과를
                    # 통한 부여)을 충족하지 못한다. 대신 5단계(텍스트 치환)에서 이 니케의 스킬 이름과 동급
                    # 별칭으로 취급하기 위한 후보로 남겨둔다(4단계 고유 속성 후보에서는 자동으로 제외됨).
                    step_match = re.match(r'^<[^<>:：]+[:：]\s*([^<>]+)>$', line)
                    if step_match:
                        step_name = step_match.group(1).strip()
                        if step_name:
                            current_items.append(('스킬 별칭', step_name))
                        last_label_was_spec = False
                        continue
                    if re.match(r'^<[^<>]+>$', line):
                        last_label_was_spec = False
                        continue

                    # 7. 통째로 폐기하는 설명문 : 기능/특이사항/해제 조건/제거 시점/효과 N 대상 (항상 적용)
                    discard_match = re.match(r'^([^:：\[\]]{1,20}?)\s*[:：]\s*(.*)$', line)
                    if discard_match and discard_match.group(1).strip() in self.DISCARD_LABELS:
                        last_label_was_spec = False
                        continue
                    if re.match(r'^효과\s*\d+\s*대상\s*[:：]', line):
                        last_label_was_spec = False
                        continue

                    # 규칙 3 보강 : `최대 [록 온] 대상 수 : 5개`, `최대 [오토 파이어 레디] 장전 수 : 5개` —
                    # 명칭 단독 대괄호 스펙 시트에 딸린 수량 상한 줄이므로 스펙 라벨과 같이 폐기한다.
                    if re.match(r'^최대\s*\[[^\[\]]+\][^:：\[\]]*[:：]', line):
                        continue

                    # 2. `효과N :` / `추가 효과` 서술 줄은 스펙 범위 안팎을 가리지 않고 항상 구제.
                    # 뒤에 '헤더 문장 + 대괄호'가 오면 그 헤더 문장은 서브 헤더(독립 블록)로 승격한다.
                    addon_header, addon_items = self.tryRescueAddOnLabel(line)
                    if addon_header or addon_items:
                        if addon_header:
                            blocks.append({'header': current_header, 'items': current_items})
                            current_header = addon_header
                            current_items = []
                        current_items.extend(addon_items)
                        last_label_was_spec = False
                        continue

                    # 대괄호 하나짜리 `[<헤더 문장> 추가 효과]` 라벨(예: 은화/엠마 택티컬 업의
                    # `[자신에게 포메이션 AS 적용 중 추가 효과]`)은 뒤따르는 `효과 N :` 줄들이 어떤 조건에서
                    # 적용되는지 알리는 서브 헤더다 — '추가 효과'를 떼고 헤더 문장을 서브 헤더로 승격한다.
                    extra_effect_label = re.fullmatch(r'\s*\[([^\[\]]+?)\s*추가\s*효과\]\s*', line)
                    if extra_effect_label and re.search(self.BOUNDARY_ENDING, extra_effect_label.group(1).strip()):
                        blocks.append({'header': current_header, 'items': current_items})
                        current_header = extra_effect_label.group(1).strip()
                        current_items = []
                        last_label_was_spec = False
                        continue

                    # 표 머리글 라벨 : 대괄호로만 된 줄 바로 뒤에 `라벨 : [효과]` 행들이 이어지면(예: D : 킬러 와이프
                    # '[명중 부위] [명중 시 적용되는 이로운 효과]' 다음의 '파츠를 명중시킨 아군 : [...]') 그 대괄호는
                    # 열 이름 설명이지 효과가 아니므로, 부가 옵션/명칭:효과 쌍이 아닌 평범한 대괄호는 버린다.
                    if self.isBracketOnlyLine(line) and not line.rstrip().endswith((':', '：')) \
                            and self.followedByLabeledRows(body_lines, line_index):
                        for table_group in [g.strip() for g in re.findall(r'\[([^\[\]]+)\]', line) if g.strip()]:
                            if ':' in table_group:
                                for kind, content in self.splitBracketContents(f'[{table_group}]'):
                                    current_items.append((kind, content))
                            elif self.isOptionText(table_group):
                                current_items.append(('부가옵션', table_group))
                        last_label_was_spec = False
                        continue

                    # 규칙 3 : 명칭 단독 대괄호(들)로만 이루어진 줄. 뒤에 실제로 `기능 :`/`효과 :`/스펙 라벨이 이어지거나
                    # 줄 끝에 콜론이 붙어 있을 때만 "이름 라벨"로 보고 이름을 뽑아 쓴다(룩어헤드 없이는 일반
                    # [명칭 : 효과] 패턴과 구분 불가 — 예: '[주목 : 적 전체 도발]'은 이름 라벨이 아니라 평범한 명칭:효과 쌍).
                    # 줄 끝 콜론 예) `[구경났냐 짜식들아!!] :` 뒤에 `[부끄러움 면역] [10초 유지]` 같은 효과 줄이 이어지는 경우.
                    ends_with_colon = line.rstrip().endswith((':', '：'))
                    is_spec_sheet = self.isBracketOnlyLine(line) and self.looksLikeSpecSheetStart(body_lines, line_index)                         and not self.firstBracketLooksLikeEffect(line)
                    if self.isBracketOnlyLine(line) and (is_spec_sheet or ends_with_colon):
                        bracket_groups = [g.strip() for g in re.findall(r'\[([^\[\]]+)\]', line) if g.strip()]
                        if bracket_groups:
                            first_group = bracket_groups[0]
                            if ':' in first_group:
                                left_part, right_part = (p.strip() for p in first_group.split(':', 1))
                                # 예) '사용 무기 변경 : 메티스 버스트 버스터' -> 진짜 이름은 공용 표현이 아닌 쪽
                                name_candidate = left_part if right_part == '사용 무기 변경' else right_part
                            else:
                                name_candidate = first_group
                            # '사용 무기 변경'처럼 구체적인 이름 없이 공용 표현 자체만 있으면 명칭 후보로 삼지 않는다.
                            # `[시작 횟수 별 효과] [하위 효과 중복 적용] :`처럼 첫 대괄호가 부가 옵션이면 이름이 아니다.
                            if name_candidate and self.isOptionText(name_candidate):
                                current_items.append(('부가옵션', name_candidate))
                            elif name_candidate and name_candidate not in self.SPEC_LABELS and name_candidate not in self.COMMON_EXPRESSIONS:
                                current_items.append(('명칭', name_candidate))
                            for extra_group in bracket_groups[1:]:
                                if self.isOptionText(extra_group):
                                    current_items.append(('부가옵션', extra_group))
                                elif ':' in extra_group:
                                    left_text, right_text = (part.strip() for part in extra_group.split(':', 1))
                                    if left_text:
                                        current_items.append(('명칭', left_text))
                                    if right_text:
                                        current_items.append(('부가옵션' if self.isOptionText(right_text) else '효과', right_text))
                                elif extra_group:
                                    current_items.append(('효과', extra_group))
                        last_label_was_spec = is_spec_sheet
                        continue

                    # 규칙 3 : 대괄호 없이 'SPEC 라벨 : 내용' 형태 — 사용 무기 변경류 스펙이므로 통째 폐기
                    if discard_match and discard_match.group(1).strip() in self.SPEC_LABELS:
                        last_label_was_spec = True
                        continue

                    # 1. 접두 라벨 제거(순수 숫자 라벨) : '1회 :', '2단계 :', '5단계이상 :'
                    num_label_match = re.match(r'^(\d+\s*(?:회|단계)(?:이하|이상)?)\s*[:：]\s*(.*)$', line)
                    if num_label_match:
                        label_text = num_label_match.group(1)
                        rest = num_label_match.group(2).strip()
                        if last_label_was_spec and '단계' in label_text:
                            # 직전 SPEC 라벨(예: 차지 시간)의 단계별 세부값 — 독립된 효과가 아니라 같이 폐기
                            continue
                        last_label_was_spec = False
                        rest_header, rest_brackets = self.splitPrefixAndBrackets(rest)
                        if rest_header:
                            blocks.append({'header': current_header, 'items': current_items})
                            current_header = rest_header
                            current_items = []
                            rest = rest_brackets
                        if rest.startswith('['):
                            for kind, content in self.splitBracketContents(rest):
                                current_items.append((kind, content))
                        elif rest and self.isSubHeaderLine(rest):
                            blocks.append({'header': current_header, 'items': current_items})
                            current_header = self.stripTargetLabel(rest)
                            current_items = []
                        elif rest:
                            current_items.append(('효과', rest))
                        continue

                    last_label_was_spec = False

                    # 5. 조건/대상형 라벨 : `명중 대상에게 : [..]`, `자신에게 : [..]`처럼 콜론 앞이 이미 경계 어미로 끝나는
                    # 헤더 문장이면 일반 서브 헤더와 똑같이 취급한다(라벨을 버리면 대상/시점 정보가 사라진다).
                    cond_label_match = re.match(r'^([^\[\]:：]+?)\s*[:：]\s*(\[.*)$', line)
                    if cond_label_match and re.search(self.BOUNDARY_ENDING, cond_label_match.group(1).strip()) \
                            and self.isBracketOnlyLine(cond_label_match.group(2)):
                        label_text = cond_label_match.group(1).strip()
                        if re.search(r'(?:라면|다면|경우|확률로|이면|으면|했다면|한하여|때)$', label_text):
                            # 상태 단계별 임계 라벨(`70% 이상일 때 한하여 : [...]`)은 쓸모없는 표현이라 라벨만 버리고
                            # 뒤의 내용만 상위 블록에 남긴다. 대상/시점 라벨만 서브 헤더로 승격한다.
                            for kind, content in self.splitBracketContents(cond_label_match.group(2)):
                                current_items.append((kind, content))
                            continue
                        blocks.append({'header': current_header, 'items': current_items})
                        current_header = label_text
                        current_items = []
                        for kind, content in self.splitBracketContents(cond_label_match.group(2)):
                            current_items.append((kind, content))
                        continue

                    # 4. 이름+번호 라벨 : '생존본능 1 :', '타겟 마킹 2 :'. '상태 1 :'처럼 라벨 자체가
                    # 범용 분기어(GENERIC_BRANCH_LABELS)면 번호만 떼고 명칭 후보로는 삼지 않는다
                    # (진짜 기믹 이름이 아니라 case 1/2 같은 분기 번호일 뿐이기 때문 — '추가 효과'와 동일한 이유).
                    name_num_label_match = re.match(r'^([가-힣A-Za-z][가-힣A-Za-z\s]*?)\s*\d+\s*[:：]\s*(.*)$', line)
                    if name_num_label_match:
                        name_candidate = name_num_label_match.group(1).strip()
                        rest = name_num_label_match.group(2).strip()
                        if name_candidate and name_candidate not in self.GENERIC_BRANCH_LABELS:
                            current_items.append(('명칭', name_candidate))
                        # '상태 1 : 자신이 무대 파트 : 댄스 상태라면 [무대 파트 : 보컬] [지속]'처럼 라벨 뒤가
                        # '헤더 문장 + 대괄호'이면 헤더 문장을 서브 헤더로 승격한다.
                        rest_header, rest_brackets = self.splitPrefixAndBrackets(rest)
                        if rest_header:
                            blocks.append({'header': current_header, 'items': current_items})
                            current_header = rest_header
                            current_items = []
                            rest = rest_brackets
                        if rest.startswith('['):
                            for kind, content in self.splitBracketContents(rest):
                                current_items.append((kind, content))
                        elif rest:
                            current_items.append(('효과', rest))
                        continue

                    # 여러 헤더 : 대괄호로 시작하지 않고 경계 어미로 끝나는 조건/대상 문장은 새 서브 헤더
                    if self.isSubHeaderLine(line):
                        blocks.append({'header': current_header, 'items': current_items})
                        current_header = self.stripTargetLabel(line)
                        current_items = []
                        continue

                    # 그 외 : 대괄호가 있으면 전부 꺼내 분류, 없으면 무시(자연어 서술 잔여분)
                    if '[' in line:
                        for kind, content in self.splitBracketContents(line):
                            current_items.append((kind, content))

                blocks.append({'header': current_header, 'items': current_items})

            blocks = self.expandEmbeddedHeaders(blocks)

            if pending_step_aliases and blocks:
                blocks[0]['items'] = pending_step_aliases + blocks[0]['items']
        except Exception as error_msg:
            print(f"[ERROR] chunkSkillDescription 실행 중 오류 발생: {error_msg}")

        return blocks

    def buildSkillBlockSheet(self):
        '''
        [DB 구축 3단계]
        '니케 스킬 설명' 시트를 입력으로 각 니케의 스킬 1/2/버스트 설명을 블록으로 분해해 '스킬 블록' 시트에 저장.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_info = workbook['니케 스킬 설명']
            sheet_block = workbook['스킬 블록']

            skill_slot_columns = [
                ('스킬 1', 5),   # (니케 스킬 설명 시트에서 스킬 1 설명이 있는 열 번호, 1-indexed)
                ('스킬 2', 7),
                ('버스트 스킬', 9),
            ]

            # 복합 상태명(고유 속성 규칙 7)은 전체 설명을 한 번 훑어야 알 수 있다 — 대괄호 `[A : B]`가 평범한
            # 명칭:효과 쌍인지 복합 상태명인지는 'A : B'가 대괄호 바깥(헤더)에서 언급되는지로 가른다.
            all_descriptions = []
            for row in sheet_info.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                for _, col_index in skill_slot_columns:
                    all_descriptions.append(row[col_index - 1])
            self.compound_state_names = self.collectCompoundStateNames(all_descriptions)

            processed_count = 0
            for row in sheet_info.iter_rows(min_row=2, values_only=True):
                nikke_id, nikke_name = row[0], row[1]
                if nikke_id is None:
                    continue

                for skill_slot_name, col_index in skill_slot_columns:
                    description_text = row[col_index - 1]
                    blocks = self.chunkSkillDescription(description_text)
                    for block_no, block in enumerate(blocks, start=1):
                        if not block['items']:
                            # 본문 요소가 하나도 없는 블록(예: 뒤따르는 `효과 N :` 줄들의 조건만 알리는 서브 헤더)도
                            # 헤더는 의미가 있으므로 '컨테이너' 행으로 남긴다 — 그래야 고유 속성의 "언급" 판정과
                            # 헤더 슬롯 분해에서 이 헤더가 빠지지 않는다.
                            sheet_block.append([
                                nikke_id, nikke_name, skill_slot_name, block_no,
                                block['header'], '컨테이너', None, 1,
                            ])
                            continue
                        for item_no, (kind, content) in enumerate(block['items'], start=1):
                            sheet_block.append([
                                nikke_id, nikke_name, skill_slot_name, block_no,
                                block['header'], kind, content, item_no,
                            ])
                processed_count += 1

            workbook.save(self.path_excel)
            print(f"[SUCCESS] buildSkillBlockSheet 완료: {processed_count}명 처리")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildSkillBlockSheet 실행 중 오류 발생: {error_msg}")
            return False

    def isControlSuffixOfCommonExpression(self, candidate_text):
        '''
        [4단계 보조: 고유 속성 제외 규칙 2] '기절 면역'처럼 공용 표현 뒤에 면역/해제/제거/방해 같은
        시스템 제어어가 결합된 형태인지 판정. 이런 후보는 언급 여부와 무관하게 고유 속성에서 제외한다.
        '''
        for suffix in self.CONTROL_SUFFIXES:
            if candidate_text.endswith(suffix):
                base_text = candidate_text[:-len(suffix)].strip()
                if base_text in self.COMMON_EXPRESSIONS:
                    return True
        return False

    def hasStateChangeStructure(self, effect_text):
        '''
        [4단계 보조: 고유 속성 규칙 1] `X를 Y로 변경` 형태의 효과 문구는 X(변경 대상)와 Y(변경 결과)가 모두 상태 이름이라
        적용 대상 성격의 표현이다(예: 트리나 `뻗은 뿌리를 시든 뿌리로 변경`) — 규칙 1의 "언급"에 해당한다.
        '''
        return bool(re.search(r'\S+[을를]\s+\S.*로\s*변경$', effect_text))

    def hasEmbeddedGimmickStructure(self, effect_text):
        '''
        [4단계 보조: 고유 속성 규칙 6] 효과 문구 하나에 시점/조건성 어미가 문장 중간에 있고
        그 뒤로 적용 대상 어미(~에게/~들에게)가 이어지면, 그 효과 문구 자체가 발동 시점·적용 대상
        성격을 내장한 것으로 본다(예: 트로니 '누적 폭발 스킬 폭파 시 공격 범위 내 적들에게 분배 대미지').
        대괄호가 섞여 있으면 적용하지 않는다 — "N초마다 아군 전체에게 [이름 : 효과]"처럼 흔한
        타이밍+대상 상투구 뒤에 평범한 [명칭:효과] 브래킷이 붙었을 뿐인 경우(예: 유키코 '메디아')까지
        걸려서, 그 명칭이 자기 자신이 담긴 이 문장 안에서 "언급됨"으로 자기참조 오탐되는 것을 막기 위함.
        진짜 규칙 6 사례(트로니)는 애초에 대괄호 없는 생짜 문장이라 영향받지 않는다.
        '''
        if '[' in effect_text or ']' in effect_text:
            return False
        match = re.search(rf'\s{self.EMBEDDED_TIMING_TOKENS}\s', effect_text)
        if not match:
            return False
        remainder = effect_text[match.end():]
        return bool(re.search(r'(?:에게|들에게)', remainder))

    # 고유 속성 규칙 8(헤더 틀 후보)에서 쓰는 상수
    # - 효과 문구 안에서 명사구 뒤에 곧바로 이어지면 "자원/상태 이름"으로 보는 수량·동작 단서
    #   (숫자, ▲▼, 'X'(곱셈 표기), 문구 끝, 수/개수/중첩, 충전·소환·포획·제작·소모·축적·삭제·해제·제거·발권·만큼·지정,
    #   '이름 - 부제' 하이픈 구분(`페르소나 - 요한나`))
    RESOURCE_CUE_PATTERN = (r'(?=\s*(?:\d|[▲▼\]]|X(?:\s|$)|$|수(?:\s|$)|개수|중첩|제작|획득|충전|소환|포획|소모|축적'
                            r'|삭제|해제|제거|발권|만큼|지정|-\s))')
    # 고유 속성 규칙 9(수량 자원형)에서 쓰는 패턴 : 효과 문구가 `<명사구> n개/n장/n발/n기`, `중첩량`, `개수`, `갯수만큼`,
    # `n 충전`, `의 ○○ 소모`로 이어지거나(A), `… X <명사구> (중첩) 수`처럼 곱셈 표기 뒤에 쓰인다(B).
    RESOURCE_UNIT_PATTERN = (r'(?:\d+(?:\.\d+)?\s*(?:개|장|발|기)(?![가-힣])|중첩량|개수|갯수만큼|\d+\s*충전'
                             r'|의\s+\S+\s+소모)')
    RESOURCE_PREFIX_REGEX = re.compile(
        r'^(?:보유\s+|공격 당\s+)?(?P<noun>[가-힣A-Za-z][가-힣A-Za-z ]*?)\s*' + RESOURCE_UNIT_PATTERN)
    RESOURCE_MULTIPLIER_REGEX = re.compile(
        r'X\s+(?P<noun>[가-힣A-Za-z][가-힣A-Za-z ]*?)\s+(?:중첩\s*수|수)(?:\s|▲|▼|$)')

    # - 헤더 틀에서 명사구 꼬리에 붙는 수식어(떼고 본체만 후보로 삼는다) : '티켓 보유 상태', '메이드 정신 최대 중첩 상태'
    FRAME_QUALIFIER_TOKENS = {'최대', '최소', '중첩', '보유', '제작', '개수', '수', '단계', '수량'}
    # - 명사구가 이런 단어로 끝나면 이름이 아니라 문법 조각이므로 후보에서 뺀다
    FRAME_GRAMMAR_TAILS = {'적용', '발동', '획득', '포획', '충전', '사용', '해제', '종료', '시작', '도달', '증가', '감소',
                           '효과', '대상', '공격', '명중', '유지', '소환', '삭제', '제거', '지정'}

    def compactText(self, text):
        '''공백을 전부 뗀 문자열. 같은 이름이 '메이드정신'/'메이드 정신'처럼 공백만 달리 쓰이는 경우를 같게 취급하기 위함.'''
        return re.sub(r'\s+', '', text or '')

    def buildFlexiblePattern(self, noun_text):
        '''noun_text의 글자 사이에 공백이 있어도/없어도 맞는 정규식 조각을 만든다.'''
        return r'\s*'.join(re.escape(char) for char in noun_text if not char.isspace())

    def extractHeaderFrameNouns(self, header_text):
        '''
        [4단계 보조: 고유 속성 규칙 8] 헤더 문장에서 "자원/상태 이름이 오는 자리"(틀)를 찾아 그 앞의 명사구 후보들을 돌려준다.
        틀 : `N 상태…`, `N 적용/발동/획득/포획/충전/사용 시·후·중`, `N 적용될/소환된…`, `N 개수가`/`N 수가`,
        `N이/가 <숫자 또는 최대/최소>…`(예: `MP가 0 이라면`, `배터리가 100% 도달 시`).
        틀 바로 앞 1~4 단어를 창(window)으로 잘라 후보로 내고, 꼬리 수식어(보유/제작/최대 중첩…)를 뗀 형태도 함께 낸다.
        어느 창이 진짜 이름인지는 호출부가 "효과 문구에서 수량·동작 단서와 함께 쓰이는가"로 가린다.
        반환값은 {틀 번호: [창 후보들(긴 것 먼저)]}.
        '''
        tokens = header_text.split()
        frames = []

        def collectWindows(end_index, strip_particle):
            windows = []
            for window_size in range(1, 5):
                start_index = end_index - window_size + 1
                if start_index < 0:
                    break
                words = tokens[start_index:end_index + 1]
                variants = [words]
                if strip_particle:
                    stripped_last = re.sub(r'(?:이|가)$', '', words[-1])
                    if stripped_last and stripped_last != words[-1]:
                        variants.append(words[:-1] + [stripped_last])
                for variant in variants:
                    trimmed = list(variant)
                    while trimmed and (trimmed[-1] in self.FRAME_QUALIFIER_TOKENS
                                       or re.match(r'^[\d.]+\S*$', trimmed[-1]) or '%' in trimmed[-1]):
                        trimmed.pop()
                    # 이름에는 개수가 들어가지 않으므로 숫자가 낀 창은 버린다(`폐품 10개`가 아니라 `폐품`)
                    if not trimmed or trimmed[-1] in self.FRAME_GRAMMAR_TAILS \
                            or any(re.search(r'[\d%]', word) for word in trimmed):
                        continue
                    candidate_text = ' '.join(trimmed)
                    if len(self.compactText(candidate_text)) >= 2 and candidate_text not in windows:
                        windows.append(candidate_text)
            windows.sort(key=lambda text: -len(text))
            return windows

        for index, token in enumerate(tokens):
            next_token = tokens[index + 1] if index + 1 < len(tokens) else ''
            if token.startswith('상태'):
                frames.append(collectWindows(index - 1, True))
            elif token in ('적용', '발동', '획득', '포획', '충전', '사용') and next_token.startswith(('시', '후', '중')):
                frames.append(collectWindows(index - 1, True))
            elif token in ('적용될', '발동될', '소환된'):
                frames.append(collectWindows(index - 1, True))
            elif token in ('개수가', '수가'):
                frames.append(collectWindows(index - 1, False))
            elif re.search(r'[이가]$', token) and re.match(r'^(?:\d|최대|최소)', next_token):
                frames.append(collectWindows(index, True))
        return frames

    def buildUniqueAttributeSheet(self):
        '''
        [DB 구축 4단계: 고유 속성 분석]
        '니케 스킬 설명'/'스킬 블록' 시트를 입력으로 니케별 고유 속성을 100% 자동 판정해
        '고유 속성' 시트에 저장한다. ('스킬 키워드 분석 규칙.md'의 "고유 속성" 절이 이 함수의 근거)

        - 후보 소스 1 : 스킬 블록의 '명칭' 항목, 명칭이 공용 표현일 때 바로 다음 항목(규칙 5 우측 결합),
          복합 상태명 `A : B`(규칙 7 — 3단계에서 대괄호를 쪼개지 않고 효과로 남긴 것), 그리고 등록된 스킬
          1/2/버스트 이름 중 "본인"의 효과(종류='효과') 문구에도 실제로 등장하는 것(규칙 2 — 예:
          치사토 '초감각'은 자신의 효과 "초감각 100% 충전"에 등장해 인정). '스킬 별칭'(<Step n> 전용
          종류)과, 효과 문구에는 전혀 안 나오고 헤더/<Step n> 라벨에만 있는 스킬 이름(예: 레드 후드
          '레드 울프')은 후보로 삼지 않는다 — 효과를 통한 부여(규칙 3)가 텍스트로 확인되지 않기 때문.
        - 후보 소스 2(규칙 8) : 이름 표기(`[이름 : 효과]`) 없이 효과 문구 속에만 나오는 자원/상태 이름(MP, 폐품, 티켓,
          배터리, 유령 …). 헤더에서 "자원/상태 이름이 오는 틀"(extractHeaderFrameNouns)로 명사구를 뽑고, 그 명사구가
          어떤 니케의 효과 문구에서 수량·동작 단서와 함께 쓰이거나(`MP 1 충전`) 문구 전체가 그 명사구(`[미래 예지]`)일 때만
          그 니케의 고유 속성으로 인정한다(규칙 3: 효과 문구로 다루는 쪽이 소유자).
        - 언급(규칙 1) 판정 : 모든 블록의 헤더(발동 시점/조건/대상이 아직 안 갈라진 원문) 전체 +
          시점·대상이 내장된 효과 문구(규칙 6) 안에 후보 문자열이 부분 일치하면 언급된 것으로 본다.
          공백 유무는 무시한다('메이드정신' = '메이드 정신').
        - 소유권(규칙 3) : 후보가 어느 니케의 블록에서 '명칭'/효과로 나왔는지로 결정.
          동일 후보를 여러 니케가 각자 정의하면(규칙 3의 "여러 니케가 가질 수 있음") 모두 별도 행으로 남는다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_info = workbook['니케 스킬 설명']
            sheet_block = workbook['스킬 블록']
            sheet_unique = workbook['고유 속성']

            # 니케id -> (니케 이름, 등급) / 니케id -> [등록된 스킬 1/2/버스트 이름들]
            nikke_meta = {}
            skill_names_by_nikke = {}
            for row in sheet_info.iter_rows(min_row=2, values_only=True):
                nikke_id = row[0]
                if nikke_id is None:
                    continue
                nikke_meta[nikke_id] = (row[1], row[2])
                skill_names_by_nikke[nikke_id] = [self.normalizeSpacing(name) for name in (row[3], row[5], row[7]) if name]

            # (니케id, 스킬 종류, 블록 번호) 단위로 그룹핑, 순번 순으로 정렬
            blocks = {}
            for row in sheet_block.iter_rows(min_row=2, values_only=True):
                nikke_id, nikke_name, skill_slot, block_no, header, kind, content, item_no = row[:8]
                if nikke_id is None:
                    continue
                block_key = (nikke_id, skill_slot, block_no)
                block = blocks.setdefault(block_key, {'header': header, 'items': []})
                block['items'].append((item_no, kind, content))
            for block in blocks.values():
                block['items'].sort(key=lambda item: item[0])

            # 규칙 1/4/6 언급 판정용 전역 텍스트 풀(니케 구분 없이 전부 모음, 공백 제거본)
            header_texts = set()
            gimmick_effect_texts = set()
            for block in blocks.values():
                if block['header']:
                    header_texts.add(block['header'])
                for _, kind, content in block['items']:
                    if kind == '효과' and (self.hasEmbeddedGimmickStructure(content)
                                           or self.hasStateChangeStructure(content)):
                        gimmick_effect_texts.add(content)
            compact_mention_pool = [self.compactText(text) for text in (header_texts | gimmick_effect_texts)]

            def isMentioned(candidate_text):
                compact_candidate = self.compactText(candidate_text)
                return any(compact_candidate in pool_text for pool_text in compact_mention_pool)

            # 후보 문자열 -> 소유(정의) 니케id 집합
            candidate_owners = {}
            canonical_spelling = {}   # (소유 니케id, 공백 뗀 표기) -> 먼저 등록된 표기('메이드정신'/'메이드 정신' 중복 방지)

            def registerCandidate(candidate_text, owner_nikke_id):
                candidate_text = (candidate_text or '').strip()
                if len(self.compactText(candidate_text)) < 2:
                    return
                if candidate_text in self.COMMON_EXPRESSIONS:
                    return
                if self.isControlSuffixOfCommonExpression(candidate_text):
                    return
                spelling_key = (owner_nikke_id, self.compactText(candidate_text))
                if canonical_spelling.setdefault(spelling_key, candidate_text) != candidate_text:
                    return
                candidate_owners.setdefault(candidate_text, set()).add(owner_nikke_id)

            for (nikke_id, skill_slot, block_no), block in blocks.items():
                items = block['items']
                for idx, (item_no, kind, content) in enumerate(items):
                    # 규칙 7 : 복합 상태명(3단계에서 쪼개지 않고 효과로 남긴 `A : B`/`A : B 삭제`)
                    if kind == '효과':
                        compound_name = self.findCompoundName(content)
                        if compound_name:
                            registerCandidate(compound_name, nikke_id)
                        continue
                    if kind != '명칭':
                        continue
                    registerCandidate(content, nikke_id)
                    # 규칙 5 우측 결합 : 명칭 자체가 공용 표현이면, 바로 다음 항목이 진짜 고유 속성 후보
                    if content.strip() in self.COMMON_EXPRESSIONS and idx + 1 < len(items):
                        next_kind, next_content = items[idx + 1][1], items[idx + 1][2]
                        if next_kind in ('효과', '부가옵션'):
                            registerCandidate(next_content, nikke_id)

            # 규칙 2(스킬 이름 연계) : 등록된 스킬 이름 자체는, 그 니케 "본인"의 효과(종류='효과') 문구
            # 안에 그 이름이 실제로 등장할 때만 후보로 인정한다 — 치사토 '초감각'(자신의 효과
            # "초감각 100% 충전"/"초감각 1% ▼"에 직접 등장)처럼 본인이 효과로 그 이름을 다루는 경우는
            # 규칙 3(효과를 통한 부여)을 충족하지만, 레드 후드 '레드 울프'처럼 헤더 조건과 <Step n> 라벨
            # (이제 '스킬 별칭')에만 등장하고 어떤 효과 문구에도 안 나오면 규칙 3을 충족하지 못해 제외한다.
            # (참고로 일레그 '붐 인스톨'처럼 텍스트 안에 [이름:효과]로 독립 정의된 경우는 이 검사와 무관하게
            # 위의 '명칭' 루프로 이미 등록돼 있다.)
            effect_texts_by_nikke = {}
            evidence_items_by_nikke = {}   # 규칙 8의 증거 후보 : (종류, 내용) — 효과 문구 + 명칭
            for (nikke_id, skill_slot, block_no), block in blocks.items():
                for item_no, kind, content in block['items']:
                    if kind == '효과':
                        effect_texts_by_nikke.setdefault(nikke_id, []).append(content)
                    if kind in ('효과', '명칭') and content:
                        evidence_items_by_nikke.setdefault(nikke_id, []).append((kind, content))

            for nikke_id, skill_names in skill_names_by_nikke.items():
                own_compact_effects = [self.compactText(text) for text in effect_texts_by_nikke.get(nikke_id, [])]
                for skill_name in skill_names:
                    compact_skill_name = self.compactText(skill_name)
                    if any(compact_skill_name in effect_text for effect_text in own_compact_effects):
                        registerCandidate(skill_name, nikke_id)

            # 규칙 8 : 헤더 틀로 뽑은 명사구가 효과 문구에서 자원/상태 이름으로 쓰이는 니케를 소유자로 등록
            all_compact_evidence = self.compactText(' '.join(
                text for items in evidence_items_by_nikke.values() for _, text in items))
            evidence_regex_cache = {}

            def findEvidenceOwners(noun_text):
                # 증거 4형태 : (d) 명칭이 그 명사구와 같다 / (c) 효과 문구 전체가 그 명사구다 /
                # (a) 효과·명칭 문구가 그 명사구로 시작하고 곧바로 수량·동작 단서가 이어진다(`MP 1 충전`) /
                # (b) `... X <명사구> 수 ▲`처럼 곱셈 표기 뒤에 단서와 함께 쓰인다(`X 용사 레벨 수 ▲`).
                # 더 긴 이름의 꼬리(예: 델타 '인법 카모플라쥬' 속 '카모플라쥬')는 증거가 되지 않는다.
                compact_noun = self.compactText(noun_text)
                if compact_noun not in all_compact_evidence:
                    return []
                if noun_text not in evidence_regex_cache:
                    flexible = self.buildFlexiblePattern(noun_text)
                    evidence_regex_cache[noun_text] = re.compile(
                        r'^\s*' + flexible + self.RESOURCE_CUE_PATTERN + r'|X\s+' + flexible + self.RESOURCE_CUE_PATTERN)
                regex = evidence_regex_cache[noun_text]
                owners = []
                for owner_id, items in evidence_items_by_nikke.items():
                    for item_kind, item_text in items:
                        if self.compactText(item_text) == compact_noun or regex.search(item_text):
                            owners.append(owner_id)
                            break
                return owners

            for header_text in header_texts:
                for window_candidates in self.extractHeaderFrameNouns(header_text):
                    for window_text in window_candidates:   # 긴 창부터 — 증거가 있는 가장 긴 것 하나만 본다
                        owners = findEvidenceOwners(window_text)
                        if owners:
                            # 가장 긴 증거 창이 공용 표현이면 이 틀은 고유 속성이 아니다(더 짧은 창으로 내려가지 않음)
                            if window_text not in self.COMMON_EXPRESSIONS and \
                                    not self.isControlSuffixOfCommonExpression(window_text):
                                for owner_id in owners:
                                    registerCandidate(window_text, owner_id)
                            break

            # 규칙 9 : 헤더에서 언급되지 않아도, 한 니케만 효과 문구에서 수량 단서와 함께 쓰는 전용 자원 이름
            # (포획 사슬, 휘슬, 소중한 추억 …)은 고유 속성으로 인정한다. 근거가 2개 이상(명칭 1 + 효과 문구들 또는
            # 서로 다른 효과 문구 2개 이상)이어야 하고, 두 니케 이상이 쓰는 명사(`최대 장탄 수`, `탄환 충전`)는 일반 용어라 제외한다.
            resource_evidence = {}   # 공백 뗀 명사구 -> {표기: ..., 'owners': {니케id: set(근거)}}
            for nikke_id, items in evidence_items_by_nikke.items():
                for item_kind, item_text in items:
                    if item_kind == '명칭':
                        continue
                    for resource_regex in (self.RESOURCE_PREFIX_REGEX, self.RESOURCE_MULTIPLIER_REGEX):
                        match = resource_regex.search(item_text)
                        if not match:
                            continue
                        noun_text = match.group('noun').strip()
                        if len(self.compactText(noun_text)) < 2 or noun_text in self.COMMON_EXPRESSIONS:
                            continue
                        entry = resource_evidence.setdefault(
                            self.compactText(noun_text), {'text': noun_text, 'owners': {}, 'reference': False})
                        entry['owners'].setdefault(nikke_id, set()).add(item_text)
                        # 수량 참조형 : 수량이 다른 효과의 변수로 쓰인다(`… X 소중한 추억 중첩 수 ▲`, `포획 사슬 갯수만큼 공격`)
                        if resource_regex is self.RESOURCE_MULTIPLIER_REGEX or '갯수만큼' in item_text:
                            entry['reference'] = True
            name_texts_by_nikke = {}
            for (nikke_id, skill_slot, block_no), block in blocks.items():
                for item_no, kind, content in block['items']:
                    if kind == '명칭':
                        name_texts_by_nikke.setdefault(nikke_id, set()).add(self.compactText(content))
            mention_exempt = set()
            extra_own_names = {}   # 고유 속성은 못 됐지만 한 니케 전용인 자원 이름 -> (소유 니케id, 근거)
            for compact_noun, entry in resource_evidence.items():
                if len(entry['owners']) != 1:
                    continue
                owner_id, evidence_set = next(iter(entry['owners'].items()))
                evidence_count = len(evidence_set) + (1 if compact_noun in name_texts_by_nikke.get(owner_id, set()) else 0)
                if evidence_count < 2:
                    continue
                if entry['text'].split()[-1] in self.FRAME_GRAMMAR_TAILS or self.isControlSuffixOfCommonExpression(entry['text']):
                    continue
                if entry['reference']:
                    registerCandidate(entry['text'], owner_id)
                    mention_exempt.add(entry['text'])
                else:
                    # 충전·소모·중첩량 증감만 있고 다른 효과가 변수로 참조하지 않는 전용 자원 이름 -> 고유 속성이 아닌 '전용 이름'
                    extra_own_names[entry['text']] = (owner_id, '전용 자원 이름(수량 참조 없음)')

            # 규칙 1(+6) 언급 필터를 통과한 후보만 최종 채택(규칙 9 후보는 헤더 언급이 필요 없다)
            final_rows = []
            owned_nikke_ids = set()
            for candidate_text, owner_ids in candidate_owners.items():
                if not isMentioned(candidate_text) and candidate_text not in mention_exempt:
                    continue
                for owner_id in owner_ids:
                    nikke_name, tier = nikke_meta[owner_id]
                    final_rows.append((nikke_name, tier, candidate_text))
                    owned_nikke_ids.add(owner_id)

            final_rows.sort(key=lambda row: (row[0], row[2]))
            for row in final_rows:
                sheet_unique.append(list(row))

            # 전용 이름 시트 : 고유 속성이 되지 못한 이름(명칭) + 수량 참조 없는 전용 자원 이름. 9단계(토큰/전용 이름 제외)가
            # 이 이름이 든 키워드를 마스터에서 뺀다. 고유 속성이 된 이름과 공용 표현은 여기에 넣지 않는다.
            sheet_own_names = workbook['전용 이름']
            final_unique_compacts = {self.compactText(row[2]) for row in final_rows}
            own_name_rows = {}
            for (nikke_id, skill_slot, block_no), block in blocks.items():
                for item_no, kind, content in block['items']:
                    if kind != '명칭' or not content:
                        continue
                    compact_name = self.compactText(content)
                    if len(compact_name) < 2 or content in self.COMMON_EXPRESSIONS                             or compact_name in final_unique_compacts or self.isOptionText(content):
                        continue
                    own_name_rows.setdefault((nikke_id, compact_name), (content, '명칭(고유 속성 아님)'))
            for noun_text, (owner_id, reason) in extra_own_names.items():
                compact_name = self.compactText(noun_text)
                if compact_name not in final_unique_compacts:
                    own_name_rows.setdefault((owner_id, compact_name), (noun_text, reason))
            for (owner_id, compact_name), (name_text, reason) in sorted(
                    own_name_rows.items(), key=lambda item: (nikke_meta[item[0][0]][0], item[1][0])):
                sheet_own_names.append([nikke_meta[owner_id][0], nikke_meta[owner_id][1], name_text, reason])

            # 고유 속성이 하나도 없는 니케는 None으로 한 행 표시
            no_attr_rows = []
            for nikke_id, (nikke_name, tier) in nikke_meta.items():
                if nikke_id not in owned_nikke_ids:
                    no_attr_rows.append((nikke_name, tier))
            no_attr_rows.sort(key=lambda row: row[0])
            for nikke_name, tier in no_attr_rows:
                sheet_unique.append([nikke_name, tier, None])

            workbook.save(self.path_excel)
            print(f"[SUCCESS] buildUniqueAttributeSheet 완료: 고유 속성 {len(final_rows)}건, "
                  f"고유 속성 없는 니케 {len(no_attr_rows)}명, 전체 니케 {len(nikke_meta)}명")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildUniqueAttributeSheet 실행 중 오류 발생: {error_msg}")
            return False

    def substituteTokens(self, text, candidates):
        '''
        [DB 구축 5단계 보조: 치환 우선순위(최장 일치)]
        text를 왼쪽부터 훑으며, 그 위치에서 시작하는 candidates(문자열, 토큰) 중 가장 긴 것을 골라
        치환한다. 길이가 같으면 candidates가 이미 (길이 내림차순, 우선순위) 순으로 정렬돼 있다고
        가정하고 그중 먼저 나오는 쪽을 쓴다 — [고유 속성] → [스킬 1/2/버스트 스킬] → [특정 니케] 순.
        공백 유무는 무시하고 맞춘다('메이드정신'/'메이드 정신', '저격모드'/'저격 모드', '1more'/'1 more') —
        text와 후보 양쪽에서 공백을 뗀 문자열끼리 비교하고, 치환은 원문 위치로 되돌려서 한다.
        '''
        if not text:
            return text
        compact_chars = []
        index_map = []   # 공백을 뗀 문자열의 i번째 글자가 원문에서 몇 번째 글자였는지
        for index, char in enumerate(text):
            if not char.isspace():
                compact_chars.append(char)
                index_map.append(index)
        compact_text = ''.join(compact_chars)

        compiled_candidates = [(self.compactText(candidate_text), token) for candidate_text, token in candidates]

        result_parts = []
        last_original_end = 0
        position = 0
        compact_length = len(compact_text)
        while position < compact_length:
            matched = None
            for compact_candidate, token in compiled_candidates:
                if compact_candidate and compact_text.startswith(compact_candidate, position):
                    matched = (compact_candidate, token)
                    break
            if matched:
                compact_candidate, token = matched
                original_start = index_map[position]
                original_end = index_map[position + len(compact_candidate) - 1] + 1
                result_parts.append(text[last_original_end:original_start])
                result_parts.append(token)
                last_original_end = original_end
                position += len(compact_candidate)
            else:
                position += 1
        result_parts.append(text[last_original_end:])
        return ''.join(result_parts)

    def buildSubstitutionCandidates(self, global_unique_attributes, own_skill_names, own_aliases_by_slot):
        '''
        [DB 구축 5단계 보조] 한 니케 전체를 기준으로 치환 후보 풀을 만들어
        (길이 내림차순, 우선순위) 순으로 정렬해 반환한다. 치환 우선순위 규칙이 "니케마다 후보 풀을
        하나로 모은다"고 명시하므로, 어느 스킬 슬롯의 텍스트를 치환하든 이 니케의 스킬 1/2/버스트
        이름과 각 슬롯의 '스킬 별칭'을 전부 같은 풀에 넣는다(예: 레드 후드 스킬 2 헤더 "자신이
        비스트 케이지 적용 시"의 '비스트 케이지'는 버스트 스킬에서 뽑힌 별칭이지만 스킬 2 텍스트에도
        똑같이 적용돼야 함 — 슬롯별로 후보를 나누면 이런 교차 참조를 놓친다).
        - 전역 고유 속성 전체(규칙 4 때문에 다른 니케 것도 포함)
        - 이 니케의 스킬 1/2/버스트 이름 3개 전부
        - 이 니케의 스킬 1/2/버스트에서 <Step n>으로 뽑힌 '스킬 별칭' 전부(해당 슬롯의 토큰으로 치환)
        - '나의 왕'/'나의 우상' -> [특정 니케]
        '''
        candidates = [(name, '[고유 속성]') for name in global_unique_attributes]
        for slot, token in self.SKILL_SLOT_TOKENS.items():
            skill_name = own_skill_names.get(slot)
            if skill_name:
                candidates.append((skill_name, token))
            for alias in own_aliases_by_slot.get(slot, []):
                candidates.append((alias, token))
        candidates.append(('나의 왕', '[특정 니케]'))
        candidates.append(('나의 우상', '[특정 니케]'))
        candidates.sort(key=lambda item: (-len(self.compactText(item[0])), self.SUBSTITUTION_PRIORITY.get(item[1], 9)))
        return candidates

    def applyTextSubstitution(self):
        '''
        [DB 구축 5단계: 텍스트 치환]
        '고유 속성'/스킬 이름(+스킬 별칭)/특정 니케 치환을 하나의 통합 최장 일치 패스로 수행한다.
        헤더는 블록(니케id, 스킬종류, 블록번호) 단위로 한 번만 치환해 같은 블록의 모든 행에 공유하고,
        내용은 '효과' 종류일 때만 치환한다(부가옵션/명칭/스킬 별칭은 최종 마스터로 가지 않으므로 불필요).
        결과는 '스킬 블록' 시트의 '치환된 헤더'/'치환된 내용' 열에 저장한다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_info = workbook['니케 스킬 설명']
            sheet_unique = workbook['고유 속성']
            sheet_block = workbook['스킬 블록']

            global_unique_attributes = sorted({
                row[2] for row in sheet_unique.iter_rows(min_row=2, values_only=True) if row[2]
            })

            skill_names_by_nikke = {}
            for row in sheet_info.iter_rows(min_row=2, values_only=True):
                nikke_id = row[0]
                if nikke_id is None:
                    continue
                skill_names_by_nikke[nikke_id] = {
                    '스킬 1': self.normalizeSpacing(row[3]), '스킬 2': self.normalizeSpacing(row[5]),
                    '버스트 스킬': self.normalizeSpacing(row[7])}

            # 스킬 별칭 수집 : (니케id, 스킬종류) -> [별칭, ...]
            all_rows = list(sheet_block.iter_rows(min_row=2))
            aliases_by_nikke_slot = {}
            for row_cells in all_rows:
                nikke_id, skill_slot, kind, content = (row_cells[0].value, row_cells[2].value,
                                                         row_cells[5].value, row_cells[6].value)
                if kind == '스킬 별칭' and content:
                    aliases_by_nikke_slot.setdefault((nikke_id, skill_slot), []).append(content)

            candidate_cache = {}  # 니케id -> 정렬된 후보 리스트(니케 전체 공통, 슬롯 구분 없음)
            header_cache = {}     # (니케id, 스킬종류, 블록번호) -> 치환된 헤더

            def getCandidates(nikke_id):
                if nikke_id not in candidate_cache:
                    own_skill_names = skill_names_by_nikke.get(nikke_id, {})
                    own_aliases = {
                        slot: aliases_by_nikke_slot.get((nikke_id, slot), [])
                        for slot in self.SKILL_SLOT_TOKENS
                    }
                    candidate_cache[nikke_id] = self.buildSubstitutionCandidates(
                        global_unique_attributes, own_skill_names, own_aliases)
                return candidate_cache[nikke_id]

            processed_count = 0
            for row_cells in all_rows:
                nikke_id, skill_slot, block_no, header, kind, content = (
                    row_cells[0].value, row_cells[2].value, row_cells[3].value,
                    row_cells[4].value, row_cells[5].value, row_cells[6].value)

                block_key = (nikke_id, skill_slot, block_no)
                if block_key not in header_cache:
                    header_cache[block_key] = self.substituteTokens(header, getCandidates(nikke_id))
                row_cells[8].value = header_cache[block_key]

                if kind == '효과':
                    row_cells[9].value = self.substituteTokens(content, getCandidates(nikke_id))

                processed_count += 1

            workbook.save(self.path_excel)
            print(f"[SUCCESS] applyTextSubstitution 완료: {processed_count}행 처리")
            return True
        except Exception as error_msg:
            print(f"[ERROR] applyTextSubstitution 실행 중 오류 발생: {error_msg}")
            return False

    def tokenizeHeaderClauses(self, text):
        '''
        [6단계 보조] 헤더(치환된) 텍스트를 경계 어미 기준으로 (절 텍스트, 어미) 목록으로 쪼갠다.
        경계 어미로 끝나지 않는 나머지가 있으면 (절 텍스트, None)으로 남겨 호출부가 미분류로 판정하게 한다.
        '''
        clauses = []
        if not text:
            return clauses
        position = 0
        for match in self.HEADER_CLAUSE_PATTERN.finditer(text):
            clause_text = text[position:match.end()].strip()
            if clause_text:
                clauses.append((clause_text, match.group(2)))
            position = match.end()
        if position < len(text):
            leftover = text[position:].strip()
            if leftover:
                clauses.append((leftover, None))
        return clauses

    def splitTimingConjunction(self, timing_text):
        '''[6단계 보조] 'A 시와 B 시' 형태의 발동 시점을 [A 시, B 시]로 쪼갠다(쪼개지 않는 경우는 [원문] 그대로).'''
        match = self.TIMING_CONJUNCTION_PATTERN.match(timing_text)
        if not match:
            return [timing_text]
        first_part, second_part = match.group(1), match.group(2)
        if len(second_part.split()) >= 3:
            return [first_part, second_part]
        # 뒤 절이 '<동사> 시' 한 단어뿐이면 앞 절의 주어(마지막 '<동사> 시'를 뺀 앞부분)를 물려받는다
        # (예: '랩쳐 등장 시와 사망 시' -> '랩쳐 등장 시', '랩쳐 사망 시'). 물려받을 주어가 없으면 쪼개지 않는다.
        subject_words = first_part.split()[:-2]
        if subject_words:
            return [first_part, ' '.join(subject_words) + ' ' + second_part]
        return [timing_text]

    def joinTimingParts(self, timing_parts):
        '''
        [6단계 보조] 발동 시점 절들을 슬롯 값 하나로 합친다. 'A 시와 B 시'에서 쪼개진 시점은 서로 다른 시점이라
        TIMING_JOINER로, 그 외 같은 종류 절은 기존대로 공백으로 이어붙인다.
        '''
        plain_parts, split_values = [], []
        for timing_text in timing_parts:
            pieces = self.splitTimingConjunction(timing_text)
            if len(pieces) > 1:
                split_values.extend(pieces)
            else:
                plain_parts.append(timing_text)
        values = ([' '.join(plain_parts)] if plain_parts else []) + split_values
        return self.TIMING_JOINER.join(values) or None

    def decomposeHeader(self, header_text):
        '''
        [DB 구축 6단계: 헤더 슬롯 분해(종류 기반) + 선행 행동]
        치환된 헤더를 (선행 행동, 발동 시점, 발동 조건, 적용 대상, 미분류)로 분류한다.
        "앞이 시점 뒤가 조건" 같은 위치 가정 없이, 각 절의 어미 종류로만 판단한다
        ('스킬 키워드 분석 규칙.md'의 "선행 행동"/"헤더 슬롯 분해(종류 기반)" 절이 근거).
        - '~후'로 끝나는 절은 예외 없이 선행 행동으로 먼저 뗀다(이 분해보다 먼저 처리).
        - 같은 종류의 절이 여럿이면(인접하지 않아도) 공백으로 이어붙여 하나의 슬롯 값으로 합친다 —
          최종 산출물에는 슬롯당 열이 하나뿐이라 순서는 결과에 영향이 없다(규칙 문서 명시).
        '''
        result = {'선행 행동': None, '발동 시점': None, '발동 조건': None, '적용 대상': None, '미분류': None}
        if not header_text:
            return result

        clauses = self.tokenizeHeaderClauses(header_text)

        preceding_parts, timing_parts, condition_parts, target_parts, unclassified_parts = [], [], [], [], []
        for clause_text, ending in clauses:
            if ending == self.PRECEDING_ACTION_ENDING:
                preceding_parts.append(clause_text)
            elif ending in self.TIMING_ENDINGS:
                timing_parts.append(clause_text)
            elif ending in self.CONDITION_ENDINGS:
                condition_parts.append(clause_text)
            elif ending in self.TARGET_ENDINGS:
                target_parts.append(clause_text)
            else:
                unclassified_parts.append(clause_text)

        result['선행 행동'] = ' '.join(preceding_parts) or None
        result['발동 시점'] = self.joinTimingParts(timing_parts)
        result['발동 조건'] = ' '.join(condition_parts) or None
        result['적용 대상'] = ' '.join(target_parts) or None
        result['미분류'] = ' '.join(unclassified_parts) or None
        return result

    def buildHeaderSlots(self):
        '''
        [DB 구축 6단계: 헤더 슬롯 분해]
        '스킬 블록' 시트의 '치환된 헤더'를 블록(니케id, 스킬종류, 블록번호) 단위로 한 번씩 분해해
        '선행 행동'/'발동 시점'/'발동 조건'/'적용 대상' 열에 저장한다(같은 블록의 모든 행이 공유).
        경계 어미로 끝나지 않아 미분류로 남는 헤더가 있으면 경고로 출력한다(정상 데이터면 0건이어야 함).
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_block = workbook['스킬 블록']

            all_rows = list(sheet_block.iter_rows(min_row=2))
            slot_cache = {}  # (니케id, 스킬종류, 블록번호) -> decomposeHeader 결과
            unclassified_examples = []

            processed_count = 0
            for row_cells in all_rows:
                nikke_id, nikke_name, skill_slot, block_no = (
                    row_cells[0].value, row_cells[1].value, row_cells[2].value, row_cells[3].value)
                substituted_header = row_cells[8].value

                block_key = (nikke_id, skill_slot, block_no)
                if block_key not in slot_cache:
                    slots = self.decomposeHeader(substituted_header)
                    slot_cache[block_key] = slots
                    if slots['미분류']:
                        unclassified_examples.append((nikke_name, skill_slot, substituted_header, slots['미분류']))
                else:
                    slots = slot_cache[block_key]

                row_cells[10].value = slots['선행 행동']
                row_cells[11].value = slots['발동 시점']
                row_cells[12].value = slots['발동 조건']
                row_cells[13].value = slots['적용 대상']
                processed_count += 1

            workbook.save(self.path_excel)
            print(f"[SUCCESS] buildHeaderSlots 완료: {processed_count}행 처리, "
                  f"미분류 헤더 {len(unclassified_examples)}건")
            for nikke_name, skill_slot, header, leftover in unclassified_examples[:10]:
                print(f"  [WARN] 미분류: {nikke_name}/{skill_slot} 헤더='{header}' 남은부분='{leftover}'")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildHeaderSlots 실행 중 오류 발생: {error_msg}")
            return False

    def buildNumberSkeleton(self, text):
        '''
        [7단계 보조] 텍스트에서 숫자를 전부 뽑아 자리표시자로 바꾼 '스켈레톤'과, 뽑힌 숫자 목록을
        (원문 순서대로) 반환한다. 스켈레톤이 같은 텍스트끼리 "같은 구조"로 간주해 묶는다.
        '''
        numbers = self.NUMBER_PATTERN.findall(text)
        skeleton = self.NUMBER_PATTERN.sub(self.NUMBER_PLACEHOLDER, text)
        return skeleton, numbers

    def substituteGroupNumbers(self, texts):
        '''
        [DB 구축 7단계: 숫자 치환]
        같은 구조(숫자를 뗀 나머지가 동일)의 텍스트들을 모아, 숫자 값이 니케마다 다르면 그 자리만
        'n'으로, 전체 데이터에서 값이 전부 같으면 그 숫자를 그대로 유지한다(화이트리스트 없이
        전체 데이터 기준으로 일괄 판정 — '스킬 키워드 분석 규칙.md' "DB 구축 순서" 7번이 근거).
        texts와 같은 길이의 치환된 텍스트 리스트를 반환한다.
        '''
        groups = {}
        for index, text in enumerate(texts):
            skeleton, numbers = self.buildNumberSkeleton(text)
            groups.setdefault(skeleton, []).append((index, numbers))

        results = [None] * len(texts)
        for skeleton, entries in groups.items():
            number_count = len(entries[0][1])
            kept_values = []
            for position in range(number_count):
                values_at_position = {numbers[position] for _, numbers in entries}
                kept_values.append(values_at_position.pop() if len(values_at_position) == 1 else None)
            for index, numbers in entries:
                rebuilt = skeleton
                for position in range(number_count):
                    final_value = numbers[position] if kept_values[position] is not None else 'n'
                    rebuilt = rebuilt.replace(self.NUMBER_PLACEHOLDER, final_value, 1)
                results[index] = rebuilt
        return results

    def applySlotNumberSubstitution(self):
        '''
        [DB 구축 7단계: 숫자 치환]
        '스킬 블록' 시트의 '선행 행동'/'발동 시점'/'발동 조건'/'적용 대상'(블록 단위로 공유)과
        '치환된 내용'(종류='효과'인 행만)에 숫자 치환을 적용해 같은 열에 덮어쓴다. 네 블록 슬롯과
        효과 내용은 서로 다른 '도메인'이므로 각각 독립적으로 전체 데이터를 모아 판정한다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_block = workbook['스킬 블록']
            all_rows = list(sheet_block.iter_rows(min_row=2))

            # 블록 단위 슬롯은 블록당 한 번만 모은다
            block_slot_cache = {}
            for row_cells in all_rows:
                block_key = (row_cells[0].value, row_cells[2].value, row_cells[3].value)
                if block_key not in block_slot_cache:
                    block_slot_cache[block_key] = (
                        row_cells[10].value or '', row_cells[11].value or '',
                        row_cells[12].value or '', row_cells[13].value or '')

            block_keys = list(block_slot_cache.keys())
            preceding_result = self.substituteGroupNumbers([block_slot_cache[k][0] for k in block_keys])
            # 발동 시점은 한 칸에 여러 시점이 TIMING_JOINER로 들어 있을 수 있어 값 단위로 풀어 치환한 뒤 다시 합친다
            timing_values = [value for k in block_keys for value in block_slot_cache[k][1].split(self.TIMING_JOINER)]
            timing_substituted = iter(self.substituteGroupNumbers(timing_values))
            timing_result = [self.TIMING_JOINER.join(next(timing_substituted)
                                                     for _ in block_slot_cache[k][1].split(self.TIMING_JOINER))
                             for k in block_keys]
            condition_result = self.substituteGroupNumbers([block_slot_cache[k][2] for k in block_keys])
            target_result = self.substituteGroupNumbers([block_slot_cache[k][3] for k in block_keys])
            new_block_slots = {
                key: (preceding_result[i] or None, timing_result[i] or None,
                      condition_result[i] or None, target_result[i] or None)
                for i, key in enumerate(block_keys)
            }

            # 효과 항목(치환된 내용)은 행 단위로 모은다
            effect_row_positions = [i for i, row_cells in enumerate(all_rows) if row_cells[5].value == '효과']
            effect_texts = [all_rows[i][9].value or '' for i in effect_row_positions]
            effect_result = self.substituteGroupNumbers(effect_texts)

            processed_count = 0
            for row_cells in all_rows:
                block_key = (row_cells[0].value, row_cells[2].value, row_cells[3].value)
                preceding, timing, condition, target = new_block_slots[block_key]
                row_cells[10].value, row_cells[11].value = preceding, timing
                row_cells[12].value, row_cells[13].value = condition, target
                processed_count += 1
            for list_index, row_position in enumerate(effect_row_positions):
                all_rows[row_position][9].value = effect_result[list_index] or None

            workbook.save(self.path_excel)
            print(f"[SUCCESS] applySlotNumberSubstitution 완료: {processed_count}행 처리, "
                  f"효과 {len(effect_row_positions)}건 숫자 치환")
            return True
        except Exception as error_msg:
            print(f"[ERROR] applySlotNumberSubstitution 실행 중 오류 발생: {error_msg}")
            return False

    def buildDebuggingList(self):
        '''
        [DB 구축 8단계: 중복 제거]
        7단계 결과('스킬 블록' 시트의 선행 행동/발동 시점/발동 조건/적용 대상/치환된 내용)에서
        종류별로 중복을 제거하고 오름차순 정렬해 '스킬 키워드 디버깅 목록' 시트에 저장한다.
        마스터 시트와 같은 원칙으로 열마다 독립된 어휘 목록이라 행끼리 서로 대응하지 않는다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_block = workbook['스킬 블록']
            sheet_debug = workbook['스킬 키워드 디버깅 목록']

            timing_set, condition_set, preceding_set, target_set, effect_set = set(), set(), set(), set(), set()
            seen_blocks = set()

            for row in sheet_block.iter_rows(min_row=2, values_only=True):
                nikke_id, skill_slot, block_no = row[0], row[2], row[3]
                kind, substituted_content = row[5], row[9]
                preceding, timing, condition, target = row[10], row[11], row[12], row[13]

                block_key = (nikke_id, skill_slot, block_no)
                if block_key not in seen_blocks:
                    seen_blocks.add(block_key)
                    if preceding:
                        preceding_set.add(preceding)
                    if timing:
                        timing_set.update(timing.split(self.TIMING_JOINER))
                    if condition:
                        condition_set.add(condition)
                    if target:
                        target_set.add(target)
                if kind == '효과' and substituted_content:
                    effect_set.add(substituted_content)

            timing_list = sorted(timing_set)
            condition_list = sorted(condition_set)
            preceding_list = sorted(preceding_set)
            target_list = sorted(target_set)
            effect_list = sorted(effect_set)

            max_length = max(len(timing_list), len(condition_list), len(preceding_list),
                              len(target_list), len(effect_list))
            for i in range(max_length):
                sheet_debug.append([
                    timing_list[i] if i < len(timing_list) else None,
                    condition_list[i] if i < len(condition_list) else None,
                    preceding_list[i] if i < len(preceding_list) else None,
                    target_list[i] if i < len(target_list) else None,
                    effect_list[i] if i < len(effect_list) else None,
                ])

            workbook.save(self.path_excel)
            print(f"[SUCCESS] buildDebuggingList 완료: 발동시점 {len(timing_list)}건, "
                  f"발동조건 {len(condition_list)}건, 선행행동 {len(preceding_list)}건, "
                  f"적용대상 {len(target_list)}건, 효과 {len(effect_list)}건")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildDebuggingList 실행 중 오류 발생: {error_msg}")
            return False

    def findContainedTokens(self, keyword_text):
        '''[9단계 보조] 키워드 문자열에 EXCLUSION_TOKENS 중 들어있는 것을 전부 찾아 리스트로 반환.'''
        return [token for token in self.EXCLUSION_TOKENS if token in keyword_text]

    def buildExcludedKeywordsSheet(self):
        '''
        [DB 구축 9단계: 토큰/선행 행동 제외]
        '스킬 키워드 디버깅 목록'의 발동 시점/발동 조건/적용 대상/효과에서 [고유 속성] 등 토큰이
        든 키워드를 제외하고, 선행 행동 열은 전부(예외 없이) 제외한다. 디버깅 목록 자체는 그대로
        두고, 제외된 키워드만 사유(들어있던 토큰 / 선행 행동)와 함께 '제외 키워드' 시트에 남긴다
        ('스킬 키워드 분석 규칙.md'의 "토큰 키워드 제외" 절이 근거).
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_debug = workbook['스킬 키워드 디버깅 목록']
            sheet_excluded = workbook['제외 키워드']

            slot_columns = [('발동 시점', 0), ('발동 조건', 1), ('선행 행동', 2), ('적용 대상', 3), ('효과', 4)]

            # 전용 이름(고유 속성은 아니지만 한 니케 전용인 이름) 목록 — 공백을 뗀 형태로 키워드와 비교한다
            own_names = []
            for row in workbook['전용 이름'].iter_rows(min_row=2, values_only=True):
                if row[2]:
                    own_names.append((row[2], self.compactText(row[2])))

            excluded_rows = []
            for slot_name, col_index in slot_columns:
                seen_keywords = set()
                for row in sheet_debug.iter_rows(min_row=2, values_only=True):
                    keyword = row[col_index]
                    if not keyword or keyword in seen_keywords:
                        continue
                    seen_keywords.add(keyword)
                    if slot_name == '선행 행동':
                        excluded_rows.append((slot_name, keyword, '선행 행동'))
                        continue
                    contained_tokens = self.findContainedTokens(keyword)
                    if contained_tokens:
                        excluded_rows.append((slot_name, keyword, '토큰: ' + ', '.join(contained_tokens)))
                        continue
                    compact_keyword = self.compactText(keyword)
                    contained_names = sorted({name for name, compact_name in own_names if compact_name in compact_keyword})
                    if contained_names:
                        excluded_rows.append((slot_name, keyword, '전용 이름: ' + ', '.join(contained_names)))

            excluded_rows.sort(key=lambda row: (row[0], row[1]))
            for row in excluded_rows:
                sheet_excluded.append(list(row))

            workbook.save(self.path_excel)
            print(f"[SUCCESS] buildExcludedKeywordsSheet 완료: {len(excluded_rows)}건 제외")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildExcludedKeywordsSheet 실행 중 오류 발생: {error_msg}")
            return False

    def buildEffectCategoryMap(self):
        '''
        [10단계 보조] EFFECT_CATEGORIES(카테고리 -> 키워드 목록)를 {키워드: 카테고리} 사전으로 뒤집는다.
        같은 키워드가 두 카테고리에 들어 있으면(분류표 입력 실수) 경고를 출력하고 먼저 나온 쪽을 쓴다.
        '''
        category_map = {}
        for category, keywords in self.EFFECT_CATEGORIES.items():
            for keyword in keywords:
                if keyword in category_map and category_map[keyword] != category:
                    print(f"  [WARN] 분류표 중복: '{keyword}'가 '{category_map[keyword]}'와 '{category}' 양쪽에 있음")
                    continue
                category_map[keyword] = category
        return category_map

    def assignEffectCategories(self):
        '''
        [DB 구축 10단계: 효과 카테고리 배정]
        '스킬 키워드 디버깅 목록'의 효과 키워드 하나하나에 분류표(EFFECT_CATEGORIES)로 카테고리를 배정해
        같은 행의 '효과 카테고리' 열에 적는다. 9단계에서 제외된 효과는 '제외', 분류표에 없으면 '미분류'로
        표시하고 미분류 목록을 출력한다(검증 단계에서 미분류 0건이 목표 — 그때그때 분류표에 추가).
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_debug = workbook['스킬 키워드 디버깅 목록']
            sheet_excluded = workbook['제외 키워드']

            excluded_effects = {
                row[1] for row in sheet_excluded.iter_rows(min_row=2, values_only=True) if row[0] == '효과'
            }
            category_map = self.buildEffectCategoryMap()

            counts = {}
            unclassified = []
            for row_index in range(2, sheet_debug.max_row + 1):
                effect = sheet_debug.cell(row=row_index, column=5).value
                if not effect:
                    continue
                if effect in excluded_effects:
                    category = '제외'
                elif effect in category_map:
                    category = category_map[effect]
                else:
                    category = '미분류'
                    unclassified.append(effect)
                sheet_debug.cell(row=row_index, column=6).value = category
                counts[category] = counts.get(category, 0) + 1

            # 분류표에만 있고 현재 데이터에는 없는 키워드(오타이거나 더 이상 안 쓰는 표현) 점검
            current_effects = {
                sheet_debug.cell(row=r, column=5).value for r in range(2, sheet_debug.max_row + 1)
            }
            stale_keys = [keyword for keyword in category_map if keyword not in current_effects]

            workbook.save(self.path_excel)
            print(f"[SUCCESS] assignEffectCategories 완료: {counts}")
            if stale_keys:
                print(f"  [WARN] 분류표에는 있으나 현재 효과 목록에 없는 키워드 {len(stale_keys)}건: {stale_keys[:10]}")
            if unclassified:
                print(f"  [WARN] 미분류 효과 {len(unclassified)}건 (분류표에 추가 필요)")
            return True
        except Exception as error_msg:
            print(f"[ERROR] assignEffectCategories 실행 중 오류 발생: {error_msg}")
            return False

    # 스킬 키워드 마스터 시트의 열 순서(발동 시점/발동 조건/적용 대상 + 효과 카테고리 5종)
    MASTER_EFFECT_COLUMNS = [
        ('공격류', '효과 - 공격류'), ('공격보조류', '효과 - 공격보조류'), ('방어생존류', '효과 - 방어생존류'),
        ('체력회복류', '효과 - 체력회복류'), ('유틸리티류', '효과 - 유틸리티류'),
    ]

    # 마스터 시트/통합 시트가 쓰는 열 이름(통합은 같은 열 안에서만 허용 — 열 사이 통합 금지)
    MASTER_COLUMN_NAMES = ['발동 시점', '발동 조건', '적용 대상'] + [name for _, name in MASTER_EFFECT_COLUMNS]

    # '키워드 통합' 시트/통합 테이블에서 통합 표현 자리에 이 표시를 쓰면 그 행의 원래 표현들은 어디에도 통합하지 않고
    # 마스터에서 뺀다(검색에 쓸 필요가 없는 키워드). 링크도 만들지 않는다.
    UNIFY_EXCLUDE_MARK = '[제외]'

    def collectMasterCandidates(self, workbook):
        '''
        [11단계 보조] '스킬 키워드 디버깅 목록'에서 '제외 키워드'를 뺀 마스터 후보(통합 적용 전 원래 표현)를
        {마스터 열 이름: 키워드 집합}으로 모은다. 효과는 10단계의 카테고리 열로 나뉘고, 카테고리를 알 수 없는 효과는
        (효과, 카테고리) 쌍으로 따로 돌려준다.
        '''
        sheet_debug = workbook['스킬 키워드 디버깅 목록']
        sheet_excluded = workbook['제외 키워드']
        excluded_keys = {(row[0], row[1]) for row in sheet_excluded.iter_rows(min_row=2, values_only=True)}

        candidates = {name: set() for name in self.MASTER_COLUMN_NAMES}
        effect_columns = {category: name for category, name in self.MASTER_EFFECT_COLUMNS}
        unclassified = []
        for row in sheet_debug.iter_rows(min_row=2, values_only=True):
            timing, condition, preceding, target, effect, category = row[:6]
            for slot_name, keyword in (('발동 시점', timing), ('발동 조건', condition), ('적용 대상', target)):
                if keyword and (slot_name, keyword) not in excluded_keys:
                    candidates[slot_name].add(keyword)
            if effect and ('효과', effect) not in excluded_keys:
                if category in effect_columns:
                    candidates[effect_columns[category]].add(effect)
                else:
                    unclassified.append((effect, category))
        return candidates, unclassified

    def loadUnifyGroupsFromDb(self):
        '''
        [11단계 보조] db의 skill_keyword_unify_table(통합 그룹)과 skill_keyword_master_table(이전에 반영된 마스터)을 읽는다.
        테이블이 없거나 비어 있으면(최초 구축) 빈 결과 — 이 경우 통합 시트는 빈 채로 만들어진다.
        반환: (groups, db_master_keys). groups = [(열, 통합 표현, [원래 표현...])], db_master_keys = {(열, 키워드)}
        '''
        groups = []
        db_master_keys = set()
        existing_tables = {row[0] for row in self.db_cursor.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if 'skill_keyword_unify_table' in existing_tables:
            grouped = {}
            for category, original, unified in self.db_cursor.execute(
                    "SELECT category, original_keyword, unified_keyword FROM skill_keyword_unify_table ORDER BY id"):
                grouped.setdefault((category, unified), []).append(original)
            order = {name: index for index, name in enumerate(self.MASTER_COLUMN_NAMES)}
            for (category, unified), originals in sorted(grouped.items(), key=lambda item: (order.get(item[0][0], 99), item[0][1])):
                groups.append((category, unified, sorted(originals)))
        if 'skill_keyword_master_table' in existing_tables:
            for category, keyword in self.db_cursor.execute("SELECT category, keyword FROM skill_keyword_master_table"):
                db_master_keys.add((category, keyword))
        return groups, db_master_keys

    def readUnifySheet(self, workbook):
        '''
        [11단계 보조] '키워드 통합' 시트(가로형)를 [(열, 통합 표현, [원래 표현...])]로 읽는다.
        열 이름이 잘못된 행 등 시트 자체의 문제는 problems(문자열 목록)로 돌려준다.
        '''
        groups = []
        problems = []
        sheet_unify = workbook['키워드 통합']
        for row_number, row in enumerate(sheet_unify.iter_rows(min_row=2, values_only=True), start=2):
            cells = [str(cell).strip() if cell is not None else '' for cell in row]
            if not any(cells):
                continue
            column_name, unified = cells[0], (cells[1] if len(cells) > 1 else '')
            # 통합 표현을 비슷한 표현 칸에도 적은 행/같은 행 안의 중복은 무해하므로 무시한다
            originals = list(dict.fromkeys(cell for cell in cells[2:] if cell and cell != unified))
            if column_name not in self.MASTER_COLUMN_NAMES:
                problems.append(f"{row_number}행: 열 이름 '{column_name}'이 마스터 열 이름이 아님")
                continue
            if not unified:
                problems.append(f"{row_number}행: 통합 표현이 비어 있음")
                continue
            groups.append((column_name, unified, originals))
        return groups, problems

    def buildUnifyMap(self, groups):
        '''
        [11단계 보조] 그룹 목록을 {(열, 원래 표현): 통합 표현}으로 뒤집으며 구조 오류를 모아 돌려준다.
        오류: 같은 원래 표현이 두 그룹에 있음 / 원래 표현이 다른 그룹의 통합 표현이기도 함(연쇄) / 원래 표현이 자기 자신.
        '''
        unify_map = {}
        problems = []
        unified_keys = {(column_name, unified) for column_name, unified, _ in groups}
        for column_name, unified, originals in groups:
            for original in originals:
                if original == unified:
                    problems.append(f"[{column_name}] '{original}'가 자기 자신으로 통합됨")
                    continue
                key = (column_name, original)
                if key in unify_map and unify_map[key] != unified:
                    problems.append(f"[{column_name}] '{original}'가 두 그룹에 있음: '{unify_map[key]}', '{unified}'")
                    continue
                unify_map[key] = unified
        for (column_name, original), unified in unify_map.items():
            if (column_name, original) in unified_keys:
                problems.append(f"[{column_name}] '{original}'가 통합 표현이면서 '{unified}'의 원래 표현이기도 함(연쇄)")
        return unify_map, problems

    def buildUnifySheet(self):
        '''
        [DB 구축 11단계 전반: 유사 표현 통합 시트]
        db의 통합 테이블(skill_keyword_unify_table)로 '키워드 통합' 시트(가로형 — 열 / 통합 표현 / 비슷한 표현 1..N)를 채운다.
        비슷한 표현 열 개수 N은 실제 최대 그룹 크기 + 여유 3열(최소 5열)로 매번 새로 정한다. db에 통합 테이블이 없거나
        비어 있으면(최초 구축) 빈 시트로 둔다. 또 현재 키워드 목록과 어긋난 곳을 '통합 점검' 시트에 남긴다
        (낡은 원래 표현 / 그룹 전체가 낡음 / 카테고리 불일치 / 신규 표현). 점검 항목은 경고일 뿐 중단 사유는 아니다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            groups, db_master_keys = self.loadUnifyGroupsFromDb()
            candidates, _ = self.collectMasterCandidates(workbook)
            all_current = {}   # 키워드 -> 현재 그것이 속한 열들(효과는 카테고리가 바뀌었는지 확인하려고)
            for column_name, keywords in candidates.items():
                for keyword in keywords:
                    all_current.setdefault(keyword, set()).add(column_name)

            sheet_unify = workbook['키워드 통합']
            sheet_unify.delete_rows(1, sheet_unify.max_row)
            max_group_size = max((len(originals) for _, _, originals in groups), default=0)
            similar_count = max(5, max_group_size + 3)
            sheet_unify.append(['열', '통합 표현'] + [f'비슷한 표현 {index}' for index in range(1, similar_count + 1)])
            for column_name, unified, originals in groups:
                sheet_unify.append([column_name, unified] + originals)

            check_rows = []
            for column_name, unified, originals in groups:
                alive_originals = 0
                for original in originals:
                    if original in candidates[column_name]:
                        alive_originals += 1
                    elif original in all_current:
                        current_columns = ', '.join(sorted(all_current[original]))
                        check_rows.append((column_name, '카테고리 불일치', original,
                                           f"현재는 '{current_columns}'에 있음 — 통합이 적용되지 않음"))
                    else:
                        check_rows.append((column_name, '낡은 원래 표현', original, f"통합 표현 '{unified}'"))
                if alive_originals == 0 and unified not in candidates[column_name]:
                    check_rows.append((column_name, '그룹 전체가 낡음', unified, '원래 표현도 통합 표현도 현재 목록에 없음'))

            # 이전에 db에 반영된 마스터가 있을 때만 '신규'를 가릴 수 있다(최초 구축에는 기준이 없음)
            if db_master_keys:
                known = set(db_master_keys)
                for column_name, unified, originals in groups:
                    known.add((column_name, unified))
                    known.update((column_name, original) for original in originals)
                for column_name in self.MASTER_COLUMN_NAMES:
                    for keyword in sorted(candidates[column_name]):
                        if (column_name, keyword) not in known:
                            check_rows.append((column_name, '신규 표현', keyword, '이전 db 마스터/통합에 없던 표현 — 통합 대상인지 확인'))

            sheet_check = workbook['통합 점검']
            for row in check_rows:
                sheet_check.append(list(row))

            workbook.save(self.path_excel)
            print(f"[SUCCESS] buildUnifySheet 완료: 통합 그룹 {len(groups)}개, 점검 항목 {len(check_rows)}건")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildUnifySheet 실행 중 오류 발생: {error_msg}")
            return False

    def buildMasterSheet(self):
        '''
        [DB 구축 11단계: 스킬 키워드 마스터 분류 완성]
        '스킬 키워드 디버깅 목록'에서 9단계에서 제외된 키워드(토큰/전용 이름/선행 행동)를 빼고 남은 키워드를
        '스킬 키워드 마스터' 시트에 8열(발동 시점/발동 조건/적용 대상/효과 - 5종)로 적재한다. 효과는 10단계에서
        배정한 카테고리 열로 나눈다. 열마다 독립된 어휘 목록이라 행끼리 대응하지 않으며, 각 열은 오름차순 정렬이다.
        '키워드 통합' 시트의 그룹이 있으면 같은 열 안에서 원래 표현을 통합 표현으로 바꿔 적재한다(통합이 없으면 1:1).
        미분류 효과가 있거나 통합 시트에 구조 오류(중복/연쇄/열 이름 오류)가 있으면 적재하지 않고 중단한다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_master = workbook['스킬 키워드 마스터']

            candidates, unclassified = self.collectMasterCandidates(workbook)
            if unclassified:
                print(f"[ERROR] buildMasterSheet 중단: 카테고리를 알 수 없는 효과 {len(unclassified)}건 "
                      f"(분류표에 추가 필요): {unclassified[:5]}")
                return False

            groups, sheet_problems = self.readUnifySheet(workbook)
            unify_map, map_problems = self.buildUnifyMap(groups)
            if sheet_problems or map_problems:
                print(f"[ERROR] buildMasterSheet 중단: '키워드 통합' 시트 오류 {len(sheet_problems) + len(map_problems)}건: "
                      f"{(sheet_problems + map_problems)[:5]}")
                return False

            sorted_columns = []
            for column_name in self.MASTER_COLUMN_NAMES:
                mapped = {unify_map.get((column_name, keyword), keyword) for keyword in candidates[column_name]}
                mapped.discard(self.UNIFY_EXCLUDE_MARK)
                sorted_columns.append(sorted(mapped))

            row_count = max(len(column) for column in sorted_columns)
            for row_index in range(row_count):
                sheet_master.append([column[row_index] if row_index < len(column) else None
                                     for column in sorted_columns])

            workbook.save(self.path_excel)
            sizes = {header: len(column) for header, column in zip(self.MASTER_COLUMN_NAMES, sorted_columns)}
            excluded_count = sum(1 for unified in unify_map.values() if unified == self.UNIFY_EXCLUDE_MARK)
            print(f"[SUCCESS] buildMasterSheet 완료: {sizes}, 총 {sum(sizes.values())}건 "
                  f"(통합 적용 {len(unify_map) - excluded_count}건, 제외 {excluded_count}건)")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildMasterSheet 실행 중 오류 발생: {error_msg}")
            return False

    def checkUnifyConsistency(self):
        '''
        [11단계 점검: 수작업 후 일관성 검사 — 12단계(링크)/13단계(DB 반영) 직전에 실행]
        승환님이 '키워드 통합' 시트와 '스킬 키워드 마스터' 시트를 손으로 고친 뒤, 둘이 서로 맞는지 읽기 전용으로 검사한다.
        오류(하나라도 있으면 False):
          1. 통합 시트로 옮긴 원래 표현이 마스터의 같은 열에 그대로 남아 있음
          2. (살아 있는 그룹의) 통합 표현이 마스터의 같은 열에 없음
          3. 현재 키워드(제외 키워드 제외)가 마스터에도 없고 통합 시트의 원래 표현에도 없음(링크할 곳이 없음)
          4. 통합 시트 구조 오류(같은 원래 표현이 두 그룹에 있음 / 연쇄 / 자기 자신 / 열 이름 오류)
        경고(통과는 함): 마스터에만 있고 현재 키워드도 통합 표현도 아닌 표현.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            candidates, unclassified = self.collectMasterCandidates(workbook)
            groups, sheet_problems = self.readUnifySheet(workbook)
            unify_map, map_problems = self.buildUnifyMap(groups)

            sheet_master = workbook['스킬 키워드 마스터']
            master_header = [cell.value for cell in sheet_master[1]]
            master_columns = {name: set() for name in self.MASTER_COLUMN_NAMES}
            for row in sheet_master.iter_rows(min_row=2, values_only=True):
                for name, cell in zip(master_header, row):
                    if name in master_columns and cell:
                        master_columns[name].add(str(cell).strip())

            errors = list(sheet_problems) + list(map_problems)
            warnings = []
            if unclassified:
                errors.append(f"카테고리를 알 수 없는 효과 {len(unclassified)}건: {unclassified[:3]}")
            originals_by_column = {name: set() for name in self.MASTER_COLUMN_NAMES}
            unified_by_column = {name: set() for name in self.MASTER_COLUMN_NAMES}
            for column_name, unified, originals in groups:
                unified_by_column[column_name].add(unified)
                originals_by_column[column_name].update(originals)
                for original in originals:
                    if original in master_columns[column_name]:
                        errors.append(f"[{column_name}] 1. 원래 표현 '{original}'가 마스터에 남아 있음(통합 표현 '{unified}')")
                alive = unified in candidates[column_name] or any(original in candidates[column_name] for original in originals)
                if unified != self.UNIFY_EXCLUDE_MARK and alive and unified not in master_columns[column_name]:
                    errors.append(f"[{column_name}] 2. 통합 표현 '{unified}'가 마스터에 없음")
            for column_name in self.MASTER_COLUMN_NAMES:
                for keyword in sorted(candidates[column_name]):
                    if keyword not in master_columns[column_name] and keyword not in originals_by_column[column_name]:
                        errors.append(f"[{column_name}] 3. 키워드 '{keyword}'가 마스터에도 통합 시트에도 없음")
                for keyword in sorted(master_columns[column_name]):
                    if keyword not in candidates[column_name] and keyword not in unified_by_column[column_name]:
                        warnings.append(f"[{column_name}] 마스터에만 있는 표현 '{keyword}'(현재 키워드도 통합 표현도 아님)")

            if errors:
                print(f"[ERROR] checkUnifyConsistency 실패: 오류 {len(errors)}건")
                for message in errors[:30]:
                    print(f"  - {message}")
            else:
                print(f"[SUCCESS] checkUnifyConsistency 통과: 통합 그룹 {len(groups)}개, 통합 적용 {len(unify_map)}건")
            if warnings:
                print(f"  [WARN] 경고 {len(warnings)}건: {warnings[:5]}")
            return not errors
        except Exception as error_msg:
            print(f"[ERROR] checkUnifyConsistency 실행 중 오류 발생: {error_msg}")
            return False

    def ensureUnifyTable(self):
        '''[13단계 보조] skill_keyword_unify_table이 없으면 만든다(DBManager.createDBTable과 같은 정의).'''
        self.db_cursor.execute('''CREATE TABLE IF NOT EXISTS skill_keyword_unify_table (
                               id INTEGER PRIMARY KEY AUTOINCREMENT,
                               category TEXT NOT NULL,
                               original_keyword TEXT NOT NULL,
                               unified_keyword TEXT NOT NULL,
                               UNIQUE (category, original_keyword)
                               )''')

    def replaceUnifyRows(self, unify_map):
        '''[13단계 보조] 통합 테이블을 비우고 {(열, 원래 표현): 통합 표현}으로 다시 채운다. commit은 호출부가 한다.'''
        self.db_cursor.execute("DELETE FROM skill_keyword_unify_table")
        # 행만 지우면 AUTOINCREMENT 번호가 이어지므로(다시 적재할 때마다 id가 커짐) 번호 카운터도 초기화해 id를 1부터 매긴다
        self.db_cursor.execute("DELETE FROM sqlite_sequence WHERE name = 'skill_keyword_unify_table'")
        self.db_cursor.executemany(
            "INSERT INTO skill_keyword_unify_table (category, original_keyword, unified_keyword) VALUES (?, ?, ?)",
            [(column_name, original, unified) for (column_name, original), unified in unify_map.items()])

    def writeUnifyGroupsToDb(self, groups):
        '''
        [13단계 보조] 통합 그룹 [(열, 통합 표현, [원래 표현...])]만 db의 skill_keyword_unify_table에 한 트랜잭션으로 다시 적재한다
        (기존 내용은 비움). 13단계 전체 반영(writeKeywordTablesToDb) 없이 통합 테이블만 고칠 때 쓴다.
        구조 오류(중복/연쇄/열 이름)가 있으면 아무것도 쓰지 않고 False를 돌려준다.
        '''
        try:
            unify_map, problems = self.buildUnifyMap(groups)
            if problems:
                print(f"[ERROR] writeUnifyGroupsToDb 중단: 통합 그룹 구조 오류 {len(problems)}건: {problems[:5]}")
                return False
            self.ensureUnifyTable()
            self.replaceUnifyRows(unify_map)
            self.db_connection.commit()
            print(f"[SUCCESS] writeUnifyGroupsToDb 완료: {len(unify_map)}건")
            return True
        except Exception as error_msg:
            self.db_connection.rollback()
            print(f"[ERROR] writeUnifyGroupsToDb 실행 중 오류 발생: {error_msg}")
            return False

    def buildLinkSheet(self):
        '''
        [DB 구축 12단계: 니케-키워드 링크 생성]
        '스킬 블록' 시트의 각 블록 키워드(발동 시점/발동 조건/적용 대상/효과)를 '키워드 통합' 시트로 번역한 뒤
        '스킬 키워드 마스터' 시트의 같은 열에서 찾아 (니케, 열, 키워드) 링크를 '니케-키워드 링크' 시트에 적는다.
        키워드는 서로 독립이다 — 포함 관계로 링크를 퍼뜨리지 않고, 니케 스킬 문구가 말하는 표현(통합 번역 후) 그대로만 연결한다.
        링크는 니케 단위라 같은 (니케, 열, 키워드)는 하나로 합치고 '출처 스킬'에 어느 스킬에서 나왔는지 남긴다(13단계는 무시).
        키워드의 id는 마스터가 db에 적재되는 13단계에서 정해지므로 이 시트에는 키워드 문자열을 적는다.
        번역 순서 : 통합 시트에 원래 표현으로 있으면 통합 표현으로(통합 표현이 [제외]면 링크 없음) → 마스터에 그대로 있으면 그대로
        → '제외 키워드' 시트에 있으면(토큰/전용 이름/선행 행동) 링크 없음 → 그 외는 오류.
        시작 전에 checkUnifyConsistency를 통과해야 하며, 마스터에서 찾지 못한 키워드가 하나라도 있으면 시트를 쓰지 않고 중단한다.
        '''
        try:
            if not self.checkUnifyConsistency():
                print("[ERROR] buildLinkSheet 중단: 통합 일관성 점검을 통과해야 링크를 만들 수 있음")
                return False

            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_block = workbook['스킬 블록']
            sheet_master = workbook['스킬 키워드 마스터']
            sheet_excluded = workbook['제외 키워드']
            sheet_link = workbook['니케-키워드 링크']

            master_header = [cell.value for cell in sheet_master[1]]
            master_columns = {name: set() for name in self.MASTER_COLUMN_NAMES}
            for row in sheet_master.iter_rows(min_row=2, values_only=True):
                for name, cell in zip(master_header, row):
                    if name in master_columns and cell:
                        master_columns[name].add(str(cell).strip())

            groups, _ = self.readUnifySheet(workbook)
            unify_map, _ = self.buildUnifyMap(groups)
            excluded_keys = {(row[0], row[1]) for row in sheet_excluded.iter_rows(min_row=2, values_only=True)}

            # 효과 키워드가 속한 마스터 열 : 10단계 카테고리로 나뉜 후보(제외 키워드 제외)에서 역으로 찾는다
            candidates, _ = self.collectMasterCandidates(workbook)
            effect_column_of = {}
            for _, column_name in self.MASTER_EFFECT_COLUMNS:
                for keyword in candidates[column_name]:
                    effect_column_of[keyword] = column_name

            stats = {'통합 번역': 0, '[제외]로 링크 안 만듦': 0, '제외 키워드로 링크 안 만듦': 0}
            problems = []
            links = {}   # (니케id, 열, 키워드) -> [니케 이름, 출처 스킬 목록]

            def addLink(nikke_id, nikke_name, skill_slot, column_name, raw_keyword, slot_name):
                if (column_name, raw_keyword) in unify_map:
                    target = unify_map[(column_name, raw_keyword)]
                    if target == self.UNIFY_EXCLUDE_MARK:
                        stats['[제외]로 링크 안 만듦'] += 1
                        return
                    stats['통합 번역'] += 1
                    keyword = target
                elif raw_keyword in master_columns[column_name]:
                    keyword = raw_keyword
                elif (slot_name, raw_keyword) in excluded_keys:
                    stats['제외 키워드로 링크 안 만듦'] += 1
                    return
                else:
                    problems.append((nikke_name, skill_slot, column_name, raw_keyword))
                    return
                if keyword not in master_columns[column_name]:
                    problems.append((nikke_name, skill_slot, column_name, f"{raw_keyword} -> {keyword}(마스터에 없음)"))
                    return
                entry = links.setdefault((nikke_id, column_name, keyword), [nikke_name, []])
                if skill_slot not in entry[1]:
                    entry[1].append(skill_slot)

            seen_blocks = set()
            for row in sheet_block.iter_rows(min_row=2, values_only=True):
                nikke_id, nikke_name, skill_slot, block_no = row[0], row[1], row[2], row[3]
                kind, substituted_content = row[5], row[9]
                timing, condition, target = row[11], row[12], row[13]

                block_key = (nikke_id, skill_slot, block_no)
                if block_key not in seen_blocks:
                    seen_blocks.add(block_key)
                    if timing:
                        for timing_value in timing.split(self.TIMING_JOINER):
                            addLink(nikke_id, nikke_name, skill_slot, '발동 시점', timing_value, '발동 시점')
                    if condition:
                        addLink(nikke_id, nikke_name, skill_slot, '발동 조건', condition, '발동 조건')
                    if target:
                        addLink(nikke_id, nikke_name, skill_slot, '적용 대상', target, '적용 대상')
                if kind == '효과' and substituted_content:
                    effect_column = effect_column_of.get(substituted_content)
                    if effect_column:
                        addLink(nikke_id, nikke_name, skill_slot, effect_column, substituted_content, '효과')
                    elif ('효과', substituted_content) in excluded_keys:
                        stats['제외 키워드로 링크 안 만듦'] += 1
                    else:
                        problems.append((nikke_name, skill_slot, '효과', substituted_content))

            if problems:
                print(f"[ERROR] buildLinkSheet 중단: 마스터에서 찾지 못한 키워드 {len(problems)}건")
                for nikke_name, skill_slot, column_name, keyword in problems[:20]:
                    print(f"  - {nikke_name}/{skill_slot} [{column_name}] {keyword}")
                return False

            # apply 모드(사람이 고친 엑셀로 다시 실행)에서는 이전 링크가 시트에 남아 있으므로 머리글만 남기고 비운 뒤 쓴다
            if sheet_link.max_row > 1:
                sheet_link.delete_rows(2, sheet_link.max_row - 1)
            column_order = {name: index for index, name in enumerate(self.MASTER_COLUMN_NAMES)}
            for (nikke_id, column_name, keyword), (nikke_name, skill_slots) in sorted(
                    links.items(), key=lambda item: (item[0][0], column_order[item[0][1]], item[0][2])):
                sheet_link.append([nikke_id, nikke_name, column_name, keyword, ', '.join(skill_slots)])

            workbook.save(self.path_excel)

            linked_keywords = {(column_name, keyword) for (_, column_name, keyword) in links}
            orphan_keywords = [(name, keyword) for name in self.MASTER_COLUMN_NAMES
                               for keyword in sorted(master_columns[name]) if (name, keyword) not in linked_keywords]
            linked_nikke_ids = {nikke_id for (nikke_id, _, _) in links}
            all_nikkes = {row[0]: row[1] for row in workbook['니케 스킬 설명'].iter_rows(min_row=2, values_only=True)
                          if row[0] is not None}
            unlinked_nikkes = [name for nikke_id, name in all_nikkes.items() if nikke_id not in linked_nikke_ids]
            print(f"[SUCCESS] buildLinkSheet 완료: 링크 {len(links)}건, 연결된 니케 {len(linked_nikke_ids)}명 / 전체 {len(all_nikkes)}명, {stats}")
            if orphan_keywords:
                print(f"  [WARN] 링크가 하나도 없는 마스터 키워드 {len(orphan_keywords)}건: {orphan_keywords[:5]}")
            if unlinked_nikkes:
                print(f"  [WARN] 링크가 하나도 없는 니케 {len(unlinked_nikkes)}명: {unlinked_nikkes[:5]}")
            return True
        except Exception as error_msg:
            print(f"[ERROR] buildLinkSheet 실행 중 오류 발생: {error_msg}")
            return False

    def writeKeywordTablesToDb(self):
        '''
        [DB 구축 13단계: DB 반영]
        엑셀의 '스킬 키워드 마스터' → skill_keyword_master_table, '니케-키워드 링크' → nikke_keyword_link_table,
        '키워드 통합' → skill_keyword_unify_table(세로형)을 한 트랜잭션에서 비우고 다시 적재한다.
        마스터 키워드 id는 열 순서(발동 시점 → … → 효과 - 유틸리티류) + 열 안 시트 순서(오름차순)대로 1부터 매긴다
        (앱이 'ORDER BY category, id'로 버튼을 만들므로 열 안에서 시트 순서가 곧 버튼 순서).
        적재 전 검사(하나라도 걸리면 아무것도 쓰지 않음): 통합 일관성 점검 통과 / 링크 시트가 비어 있지 않음 /
        링크의 키워드가 마스터 같은 열에 없음 / 링크의 니케 id가 nikke_info_table에 없음 / 통합 시트 구조 오류.
        (같은 키워드 문자열이 두 열에 있는 것은 허용 — 마스터 테이블의 유일 키가 (category, keyword)이다.)
        이 단계가 끝나야 엑셀을 다시 만들어도 수작업(통합/마스터 수정)이 db에서 복원된다.
        '''
        try:
            if not self.checkUnifyConsistency():
                print("[ERROR] writeKeywordTablesToDb 중단: 통합 일관성 점검을 통과해야 db에 반영할 수 있음")
                return False

            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_master = workbook['스킬 키워드 마스터']
            sheet_link = workbook['니케-키워드 링크']

            # 마스터 : 열별로 시트에 적힌 순서 그대로
            master_header = [cell.value for cell in sheet_master[1]]
            master_rows = []   # [(열, 키워드)]
            master_keys = set()
            for column_index, column_name in enumerate(master_header):
                if column_name not in self.MASTER_COLUMN_NAMES:
                    continue
                for row in sheet_master.iter_rows(min_row=2, values_only=True):
                    cell = row[column_index] if column_index < len(row) else None
                    if cell and (column_name, str(cell).strip()) not in master_keys:
                        master_keys.add((column_name, str(cell).strip()))
                        master_rows.append((column_name, str(cell).strip()))
            column_order = {name: index for index, name in enumerate(self.MASTER_COLUMN_NAMES)}
            master_rows.sort(key=lambda item: column_order[item[0]])   # 안정 정렬 — 열 안 순서는 유지

            # 같은 키워드 문자열이 다른 열에 있어도 된다 — 마스터 테이블은 (category, keyword) 쌍이 유일
            problems = []

            existing_nikke_ids = {row[0] for row in self.db_cursor.execute("SELECT id FROM nikke_info_table")}
            link_rows = []     # [(니케id, 열, 키워드)]
            for row in sheet_link.iter_rows(min_row=2, values_only=True):
                if row[0] is None:
                    continue
                nikke_id, column_name, keyword = row[0], row[2], row[3]
                if (column_name, keyword) not in master_keys:
                    problems.append(f"링크 키워드 [{column_name}] '{keyword}'(니케 {row[1]})가 마스터에 없음")
                if nikke_id not in existing_nikke_ids:
                    problems.append(f"링크 니케 id {nikke_id}({row[1]})가 nikke_info_table에 없음")
                link_rows.append((nikke_id, column_name, keyword))
            if not link_rows:
                problems.append("'니케-키워드 링크' 시트가 비어 있음 — 12단계(buildLinkSheet)를 먼저 실행")

            groups, sheet_problems = self.readUnifySheet(workbook)
            unify_map, map_problems = self.buildUnifyMap(groups)
            problems += sheet_problems + map_problems

            if problems:
                print(f"[ERROR] writeKeywordTablesToDb 중단: 적재 전 검사 오류 {len(problems)}건")
                for message in problems[:20]:
                    print(f"  - {message}")
                return False

            # ---- 한 트랜잭션 ----
            self.ensureUnifyTable()
            self.db_cursor.execute("DELETE FROM nikke_keyword_link_table")
            self.db_cursor.execute("DELETE FROM skill_keyword_master_table")
            self.db_cursor.execute("DELETE FROM sqlite_sequence WHERE name = 'skill_keyword_master_table'")

            keyword_ids = {}
            for column_name, keyword in master_rows:
                self.db_cursor.execute("INSERT INTO skill_keyword_master_table (category, keyword) VALUES (?, ?)",
                                       (column_name, keyword))
                keyword_ids[(column_name, keyword)] = self.db_cursor.lastrowid

            link_pairs = sorted({(nikke_id, keyword_ids[(column_name, keyword)])
                                 for nikke_id, column_name, keyword in link_rows})
            self.db_cursor.executemany("INSERT INTO nikke_keyword_link_table (nikke_id, keyword_id) VALUES (?, ?)", link_pairs)

            self.replaceUnifyRows(unify_map)
            self.db_connection.commit()

            print(f"[SUCCESS] writeKeywordTablesToDb 완료: 마스터 {len(master_rows)}건, 링크 {len(link_pairs)}건, "
                  f"통합 {len(unify_map)}건")
            return True
        except Exception as error_msg:
            self.db_connection.rollback()
            print(f"[ERROR] writeKeywordTablesToDb 실행 중 오류 발생(롤백함): {error_msg}")
            return False

    def validatePipeline(self):
        '''
        [DB 구축 14단계: 검증]
        엑셀 산출물과 db를 대조해 '검증 결과' 시트(구분 / 항목 / 내용)에 남기고, 오류가 하나라도 있으면 False를 돌려준다.
        오류 : 미분류 효과 / 링크 시트가 비었거나 마스터에 없는 키워드를 가리킴 / 스킬(스킬 1·스킬 2·버스트 스킬)에 블록이 없음 /
              링크가 하나도 없는 마스터 키워드 / 링크가 하나도 없는 니케 / db(마스터·링크·통합 테이블)가 엑셀과 다름(13단계 미반영).
              토큰([고유 속성] 등)이 남은 마스터 키워드.
        경고(사람이 확인) : 숫자만 다른 마스터 키워드(통합 후보) / 니케 이름·스킬명이 그대로 남은 마스터 키워드 /
              '통합 점검' 시트 항목.
        정보 : 값이 하나뿐이라 숫자가 남은 마스터 키워드(7단계 규칙상 정상 — 통합할 때 참고).
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            results = []   # (구분, 항목, 내용)

            def addError(item, detail):
                results.append(('오류', item, detail))

            def addWarning(item, detail):
                results.append(('경고', item, detail))

            # ---- 엑셀 산출물 읽기 ----
            sheet_master = workbook['스킬 키워드 마스터']
            master_header = [cell.value for cell in sheet_master[1]]
            master_keys = set()
            for row in sheet_master.iter_rows(min_row=2, values_only=True):
                for name, cell in zip(master_header, row):
                    if name in self.MASTER_COLUMN_NAMES and cell:
                        master_keys.add((name, str(cell).strip()))
            link_keys = set()
            for row in workbook['니케-키워드 링크'].iter_rows(min_row=2, values_only=True):
                if row[0] is not None:
                    link_keys.add((row[0], row[2], row[3]))
            nikke_rows = [row for row in workbook['니케 스킬 설명'].iter_rows(min_row=2, values_only=True) if row[0] is not None]

            # 1. 미분류 효과
            unclassified = sorted({row[4] for row in workbook['스킬 키워드 디버깅 목록'].iter_rows(min_row=2, values_only=True)
                                   if row[4] and row[5] == '미분류'})
            for effect in unclassified:
                addError('미분류 효과', effect)

            # 2. 링크 조회 실패
            if not link_keys:
                addError('링크', "'니케-키워드 링크' 시트가 비어 있음 — 12단계(buildLinkSheet) 미실행 또는 중단")
            for nikke_id, column_name, keyword in sorted(link_keys, key=str):
                if (column_name, keyword) not in master_keys:
                    addError('링크 조회 실패', f"니케 id {nikke_id} [{column_name}] {keyword}")

            # 3. 스킬마다 블록 1개 이상
            block_slots = {(row[0], row[2]) for row in workbook['스킬 블록'].iter_rows(min_row=2, values_only=True) if row[0] is not None}
            for row in nikke_rows:
                for skill_slot in self.SKILL_SLOT_TOKENS:
                    if (row[0], skill_slot) not in block_slots:
                        addError('블록 없는 스킬', f"{row[1]} / {skill_slot}")

            # 4. 링크가 하나도 없는 마스터 키워드 / 니케
            linked_keywords = {(column_name, keyword) for _, column_name, keyword in link_keys}
            for column_name, keyword in sorted(master_keys - linked_keywords):
                addError('링크 없는 마스터 키워드', f"[{column_name}] {keyword}")
            linked_nikke_ids = {nikke_id for nikke_id, _, _ in link_keys}
            for row in nikke_rows:
                if row[0] not in linked_nikke_ids:
                    addError('링크 없는 니케', row[1])

            # 5. db 반영 확인(13단계)
            db_master = {(category, keyword): keyword_id for keyword_id, category, keyword in
                         self.db_cursor.execute("SELECT id, category, keyword FROM skill_keyword_master_table")}
            if set(db_master) != master_keys:
                addError('db 불일치', f"마스터 테이블이 엑셀과 다름(엑셀에만 {len(master_keys - set(db_master))}건, "
                                    f"db에만 {len(set(db_master) - master_keys)}건) — 13단계(writeKeywordTablesToDb) 필요")
            else:
                id_to_key = {keyword_id: key for key, keyword_id in db_master.items()}
                db_links = {(nikke_id,) + id_to_key[keyword_id] for nikke_id, keyword_id in
                            self.db_cursor.execute("SELECT nikke_id, keyword_id FROM nikke_keyword_link_table")
                            if keyword_id in id_to_key}
                if db_links != link_keys:
                    addError('db 불일치', f"링크 테이블이 엑셀과 다름(엑셀에만 {len(link_keys - db_links)}건, "
                                        f"db에만 {len(db_links - link_keys)}건) — 13단계 필요")
            groups, _ = self.readUnifySheet(workbook)
            sheet_unify_map, _ = self.buildUnifyMap(groups)
            self.ensureUnifyTable()
            db_unify_map = {(category, original): unified for category, original, unified in self.db_cursor.execute(
                "SELECT category, original_keyword, unified_keyword FROM skill_keyword_unify_table")}
            if db_unify_map != sheet_unify_map:
                addError('db 불일치', "통합 테이블이 '키워드 통합' 시트와 다름 — 13단계 필요")

            # 6. 치환 잔여물
            #  - 토큰이 남은 마스터 키워드는 오류(9단계에서 빠졌어야 함)
            #  - 숫자는 7단계 규칙상 값이 하나뿐일 때만 남으므로 정상 — '정보'로만 남긴다
            #  - 숫자를 n으로 바꾸면 같은 열의 다른 마스터 키워드와 같아지는 표현은 통합 후보로 경고
            column_index_of = {name: index for index, name in enumerate(self.MASTER_COLUMN_NAMES)}
            sorted_master = sorted(master_keys, key=lambda key: (column_index_of[key[0]], key[1]))
            normalized_groups = {}
            for column_name, keyword in sorted_master:
                contained_tokens = self.findContainedTokens(keyword)
                if contained_tokens:
                    addError('토큰이 남은 키워드', f"[{column_name}] {keyword} ← {', '.join(contained_tokens)}")
                if self.NUMBER_PATTERN.search(keyword):
                    results.append(('정보', '값이 하나뿐이라 숫자가 남은 키워드', f"[{column_name}] {keyword}"))
                normalized = self.NUMBER_PATTERN.sub('n', keyword)
                normalized_groups.setdefault((column_name, normalized), []).append(keyword)
            for (column_name, _), keywords in normalized_groups.items():
                if len(keywords) > 1:
                    addWarning('숫자만 다른 마스터 키워드(통합 후보)', f"[{column_name}] " + ' / '.join(keywords))

            #  - 니케 이름·스킬명이 그대로 남은 키워드 : 단어 단위로 맞춘다(라이플의 '라이', 활성화의 '성화' 같은 오탐 방지).
            #    이름 뒤에 조사가 붙은 경우(라피의, 라피에게)도 잡는다.
            name_candidates = set()
            for row in nikke_rows:
                full_name = row[1] or ''
                name_candidates.add(full_name)
                name_candidates.update(part for part in full_name.split(' : '))
                for skill_name in (row[3], row[5], row[7]):
                    if skill_name:
                        name_candidates.add(self.normalizeSpacing(skill_name))
            common_expressions = set(self.COMMON_EXPRESSIONS)
            name_token_lists = [(name, name.split()) for name in name_candidates
                                if name and len(self.compactText(name)) >= 2 and name not in common_expressions]
            particle_pattern = re.compile(r'^(?:의|에게|에|가|이|를|을|와|과|는|은|로|으로)?$')

            def containsName(keyword_tokens, name_tokens):
                span = len(name_tokens)
                for start in range(len(keyword_tokens) - span + 1):
                    window = keyword_tokens[start:start + span]
                    if window[:-1] != name_tokens[:-1]:
                        continue
                    last_token, last_name = window[-1], name_tokens[-1]
                    if last_token.startswith(last_name) and particle_pattern.match(last_token[len(last_name):]):
                        return True
                return False

            for column_name, keyword in sorted_master:
                keyword_tokens = keyword.split()
                hits = sorted({name for name, name_tokens in name_token_lists if containsName(keyword_tokens, name_tokens)})
                if hits:
                    addWarning('이름이 남은 키워드', f"[{column_name}] {keyword} ← {', '.join(hits)}")
            for row in workbook['통합 점검'].iter_rows(min_row=2, values_only=True):
                if row[0]:
                    addWarning('통합 점검', f"[{row[0]}] {row[1]}: {row[2]}")

            if '검증 결과' in workbook.sheetnames:
                del workbook['검증 결과']
            sheet_result = workbook.create_sheet('검증 결과')
            sheet_result.append(['구분', '항목', '내용'])
            for row in results:
                sheet_result.append(list(row))
            workbook.save(self.path_excel)

            error_count = sum(1 for row in results if row[0] == '오류')
            warning_count = sum(1 for row in results if row[0] == '경고')
            summary = {}
            for kind, item, _ in results:
                summary[(kind, item)] = summary.get((kind, item), 0) + 1
            if error_count:
                print(f"[ERROR] validatePipeline 실패: 오류 {error_count}건, 경고 {warning_count}건 {summary}")
            else:
                print(f"[SUCCESS] validatePipeline 통과: 오류 0건, 경고 {warning_count}건 {summary}")
            return error_count == 0
        except Exception as error_msg:
            print(f"[ERROR] validatePipeline 실행 중 오류 발생: {error_msg}")
            return False


    # ==========================================================================
    # 변경 표시 — 마지막 db 반영(RUN_APPLY) 상태와 새로 만든 엑셀을 비교해 바뀐 곳에 색을 칠한다
    # ==========================================================================
    # '반영 기준' 시트 : RUN_APPLY(13단계 직후)에 그 시점의 마스터/링크/스킬 설명을 남긴다. 니케는 id가 아니라 (이름, 등급)으로
    # 남긴다(크롤링이 니케 테이블을 다시 만들면서 id가 밀리기 때문). RUN_WORKBOOK은 기존 엑셀의 이 시트를 새 엑셀로 그대로 넘기고,
    # 이것과 비교해 색을 칠하므로 RUN_WORKBOOK을 여러 번 돌려도 기준은 다음 RUN_APPLY 전까지 유지된다.
    BASELINE_SHEET_NAME = '반영 기준'
    BASELINE_HEADER = ['구분', '니케 이름', '등급', '열', '내용']
    CHANGE_SHEET_NAME = '변경 점검'
    DESCRIPTION_SLOT_COLUMNS = [('스킬 1', 4), ('스킬 2', 6), ('버스트 스킬', 8)]   # '니케 스킬 설명' 시트의 설명 열(0부터)

    # 색 : 추가(초록) / 변경(노랑) / 삭제(빨강) — 엑셀 기본 '좋음/보통/나쁨' 셀 스타일과 같은 색
    CHANGE_FILL_COLORS = {'추가': 'C6EFCE', '변경': 'FFEB9C', '삭제': 'FFC7CE'}

    def makeChangeFill(self, change_kind):
        '''[변경 표시 보조] 추가/변경/삭제에 맞는 셀 배경색.'''
        color = self.CHANGE_FILL_COLORS[change_kind]
        return PatternFill(start_color=color, end_color=color, fill_type='solid')

    def readNikkeIdentity(self, workbook):
        '''[변경 표시 보조] '니케 스킬 설명' 시트에서 {니케 id: (이름, 등급)}을 만든다.'''
        return {row[0]: (row[1], row[2]) for row in workbook['니케 스킬 설명'].iter_rows(min_row=2, values_only=True)
                if row[0] is not None}

    def buildBaselineRows(self, workbook):
        '''
        [변경 표시 보조] 현재 엑셀의 마스터/링크/스킬 설명을 '반영 기준' 시트 행([구분, 니케 이름, 등급, 열, 내용])으로 만든다.
        - 마스터 : ['마스터', None, None, 열, 키워드]
        - 링크   : ['링크', 니케 이름, 등급, 열, 키워드]
        - 스킬 설명 : ['스킬 설명', 니케 이름, 등급, 스킬 종류, 띄어쓰기 통일된 설명]
        '''
        rows = []
        sheet_master = workbook['스킬 키워드 마스터']
        master_header = [cell.value for cell in sheet_master[1]]
        for row in sheet_master.iter_rows(min_row=2, values_only=True):
            for name, cell in zip(master_header, row):
                if name in self.MASTER_COLUMN_NAMES and cell:
                    rows.append(['마스터', None, None, name, str(cell).strip()])
        identity = self.readNikkeIdentity(workbook)
        for row in workbook['니케-키워드 링크'].iter_rows(min_row=2, values_only=True):
            if row[0] in identity:
                rows.append(['링크', identity[row[0]][0], identity[row[0]][1], row[2], row[3]])
        for row in workbook['니케 스킬 설명'].iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            for skill_slot, column_index in self.DESCRIPTION_SLOT_COLUMNS:
                rows.append(['스킬 설명', row[1], row[2], skill_slot, row[column_index]])
        return rows

    def parseBaselineRows(self, rows):
        '''[변경 표시 보조] '반영 기준' 행들을 (마스터 집합, 링크 집합, {(이름, 등급, 스킬 종류): 설명})으로 나눈다.'''
        master_keys, link_keys, descriptions = set(), set(), {}
        for kind, nikke_name, tier, column_name, content in rows:
            if kind == '마스터':
                master_keys.add((column_name, content))
            elif kind == '링크':
                link_keys.add((nikke_name, tier, column_name, content))
            elif kind == '스킬 설명':
                descriptions[(nikke_name, tier, column_name)] = content
        return master_keys, link_keys, descriptions

    def writeBaselineSheet(self, workbook, rows):
        '''[변경 표시 보조] '반영 기준' 시트를 머리글 + rows로 다시 쓴다.'''
        if self.BASELINE_SHEET_NAME in workbook.sheetnames:
            del workbook[self.BASELINE_SHEET_NAME]
        sheet_baseline = workbook.create_sheet(self.BASELINE_SHEET_NAME)
        sheet_baseline.append(self.BASELINE_HEADER)
        for row in rows:
            sheet_baseline.append(list(row))

    def loadPreviousBaseline(self):
        '''
        [변경 표시 보조] RUN_WORKBOOK이 엑셀을 덮어쓰기 전에, 기존 엑셀의 '반영 기준' 시트를 메모리로 읽어 둔다
        (markWorkbookChanges가 새 엑셀에 다시 넣고 비교에 쓴다). 엑셀이나 시트가 없으면 None.
        '''
        self.previous_baseline_rows = None
        if not self.path_excel.exists():
            return
        workbook = openpyxl.load_workbook(self.path_excel)
        if self.BASELINE_SHEET_NAME not in workbook.sheetnames:
            return
        self.previous_baseline_rows = [list(row[:5]) for row in workbook[self.BASELINE_SHEET_NAME].iter_rows(
            min_row=2, values_only=True) if row and row[0]]

    def saveAppliedBaseline(self):
        '''
        [RUN_APPLY : 13단계 직후] db에 반영한 지금의 마스터/링크/스킬 설명을 '반영 기준' 시트에 남긴다.
        다음 RUN_WORKBOOK이 이것과 비교해 바뀐 곳에 색을 칠한다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            rows = self.buildBaselineRows(workbook)
            self.writeBaselineSheet(workbook, rows)
            workbook.save(self.path_excel)
            counts = {}
            for row in rows:
                counts[row[0]] = counts.get(row[0], 0) + 1
            print(f"[SUCCESS] saveAppliedBaseline 완료: {counts}")
            return True
        except Exception as error_msg:
            print(f"[ERROR] saveAppliedBaseline 실행 중 오류 발생: {error_msg}")
            return False

    def markWorkbookChanges(self):
        '''
        [RUN_WORKBOOK 마지막 단계] 기존 엑셀에서 읽어 둔 '반영 기준'(loadPreviousBaseline)과 새로 만든 엑셀을 비교한다.
        - '니케 스킬 설명' : 새 니케 행 초록, 설명이 바뀐 셀 노랑
        - '스킬 키워드 마스터' : 새로 생긴 키워드 셀 초록
        - '니케-키워드 링크' : 새로 생긴 링크 행 초록
        - '변경 점검' 시트 : 추가/사라진 니케, 설명 변경(키워드 변화 있음/없음), 추가/사라진 키워드, 추가/사라진 링크
          (사라진 것은 새 엑셀에 행이 없어 색을 칠할 수 없으므로 이 시트에서만 보인다)
        '반영 기준' 시트는 새 엑셀에 그대로 다시 넣는다(다음 RUN_APPLY 때 바뀜). 기준이 없으면 색 없이 넘어간다.
        '''
        try:
            workbook = openpyxl.load_workbook(self.path_excel)
            sheet_change = workbook[self.CHANGE_SHEET_NAME]
            baseline_rows = getattr(self, 'previous_baseline_rows', None)
            if not baseline_rows:
                sheet_change.append(['기준 없음', None, None, None,
                                     "'반영 기준' 시트가 없어 비교하지 않음 — RUN_APPLY를 한 번 하면 다음 RUN_WORKBOOK부터 표시"])
                workbook.save(self.path_excel)
                print("[INFO] markWorkbookChanges : 반영 기준이 없어 변경 표시를 건너뜀(RUN_APPLY 후부터 표시)")
                return True

            self.writeBaselineSheet(workbook, baseline_rows)
            base_master, base_links, base_descriptions = self.parseBaselineRows(baseline_rows)
            now_master, now_links, now_descriptions = self.parseBaselineRows(self.buildBaselineRows(workbook))

            fill_add, fill_change, fill_remove = (self.makeChangeFill('추가'), self.makeChangeFill('변경'),
                                                  self.makeChangeFill('삭제'))
            base_nikkes = {(name, tier) for name, tier, _ in base_descriptions}
            now_nikkes = {(name, tier) for name, tier, _ in now_descriptions}
            added_nikkes, removed_nikkes = now_nikkes - base_nikkes, base_nikkes - now_nikkes
            added_master, removed_master = now_master - base_master, base_master - now_master
            added_links, removed_links = now_links - base_links, base_links - now_links
            changed_slots = {}   # (이름, 등급) -> [바뀐 스킬 종류]
            for (name, tier, skill_slot), text in now_descriptions.items():
                if (name, tier) in base_nikkes and base_descriptions.get((name, tier, skill_slot)) != text:
                    changed_slots.setdefault((name, tier), []).append(skill_slot)

            # ---- 색 칠하기 ----
            slot_column_index = dict(self.DESCRIPTION_SLOT_COLUMNS)
            for row_cells in workbook['니케 스킬 설명'].iter_rows(min_row=2):
                key = (row_cells[1].value, row_cells[2].value)
                if key in added_nikkes:
                    for cell in row_cells:
                        cell.fill = fill_add
                for skill_slot in changed_slots.get(key, []):
                    row_cells[slot_column_index[skill_slot]].fill = fill_change
            sheet_master = workbook['스킬 키워드 마스터']
            master_header = [cell.value for cell in sheet_master[1]]
            for row_cells in sheet_master.iter_rows(min_row=2):
                for name, cell in zip(master_header, row_cells):
                    if cell.value and (name, str(cell.value).strip()) in added_master:
                        cell.fill = fill_add
            identity = self.readNikkeIdentity(workbook)
            for row_cells in workbook['니케-키워드 링크'].iter_rows(min_row=2):
                nikke_id = row_cells[0].value
                if nikke_id in identity and (identity[nikke_id] + (row_cells[2].value, row_cells[3].value)) in added_links:
                    for cell in row_cells:
                        cell.fill = fill_add

            # ---- 변경 점검 시트 ----
            def appendChange(kind, fill, nikke_name, tier, column_name, content):
                sheet_change.append([kind, nikke_name, tier, column_name, content])
                for cell in sheet_change[sheet_change.max_row]:
                    cell.fill = fill

            link_changed_nikkes = {(name, tier) for name, tier, _, _ in added_links | removed_links}
            for name, tier in sorted(added_nikkes):
                appendChange('추가된 니케', fill_add, name, tier, None, None)
            for name, tier in sorted(removed_nikkes):
                appendChange('사라진 니케', fill_remove, name, tier, None, None)
            for (name, tier), skill_slots in sorted(changed_slots.items()):
                kind = '설명 변경(키워드 변화 있음)' if (name, tier) in link_changed_nikkes else '설명 변경(키워드 변화 없음)'
                appendChange(kind, fill_change, name, tier, ', '.join(skill_slots), None)
            column_order = {name: index for index, name in enumerate(self.MASTER_COLUMN_NAMES)}
            for column_name, keyword in sorted(added_master, key=lambda key: (column_order.get(key[0], 99), key[1])):
                appendChange('추가된 키워드', fill_add, None, None, column_name, keyword)
            for column_name, keyword in sorted(removed_master, key=lambda key: (column_order.get(key[0], 99), key[1])):
                appendChange('사라진 키워드', fill_remove, None, None, column_name, keyword)
            link_sort = lambda key: (key[0], str(key[1]), column_order.get(key[2], 99), key[3])
            for name, tier, column_name, keyword in sorted(added_links, key=link_sort):
                appendChange('추가된 링크', fill_add, name, tier, column_name, keyword)
            for name, tier, column_name, keyword in sorted(removed_links, key=link_sort):
                appendChange('사라진 링크', fill_remove, name, tier, column_name, keyword)
            if sheet_change.max_row == 1:
                sheet_change.append(['변화 없음', None, None, None, '마지막 RUN_APPLY 이후 바뀐 니케/키워드/링크 없음'])

            workbook.save(self.path_excel)
            print(f"[SUCCESS] markWorkbookChanges 완료: 니케 +{len(added_nikkes)}/-{len(removed_nikkes)}, "
                  f"설명 변경 {len(changed_slots)}명, 키워드 +{len(added_master)}/-{len(removed_master)}, "
                  f"링크 +{len(added_links)}/-{len(removed_links)} (자세한 내용은 '{self.CHANGE_SHEET_NAME}' 시트)")
            return True
        except Exception as error_msg:
            print(f"[ERROR] markWorkbookChanges 실행 중 오류 발생: {error_msg}")
            return False

    # ==========================================================================
    # 파이프라인 실행(진입점) — '스킬 키워드 분석 규칙.md'의 "DB 구축 순서" 2~14단계
    # ==========================================================================
    # 엑셀 생성 파이프라인(2~12단계) : 각 단계는 앞 단계가 만든 시트를 읽어 다음 시트를 만든다.
    # 엑셀은 매번 전체를 새로 만들며, 사람이 관리하는 입력(공용 표현/효과 분류표는 코드, 통합은 db)으로 다시 채워진다.
    WORKBOOK_PIPELINE_STEPS = [
        ('엑셀 뼈대 생성', 'createExcelWorkbook'),
        ('2단계 띄어쓰기 통일', 'buildDescriptionSheet'),
        ('3단계 스킬 설명 구조 분해', 'buildSkillBlockSheet'),
        ('4단계 고유 속성 분석', 'buildUniqueAttributeSheet'),
        ('5단계 텍스트 치환', 'applyTextSubstitution'),
        ('6단계 헤더 슬롯 분해', 'buildHeaderSlots'),
        ('7단계 숫자 치환', 'applySlotNumberSubstitution'),
        ('8단계 중복 제거(디버깅 목록)', 'buildDebuggingList'),
        ('9단계 제외 키워드', 'buildExcludedKeywordsSheet'),
        ('10단계 효과 카테고리 배정', 'assignEffectCategories'),
        ('11단계 통합 시트', 'buildUnifySheet'),
        ('11단계 스킬 키워드 마스터', 'buildMasterSheet'),
        ('12단계 니케-키워드 링크', 'buildLinkSheet'),
        ('변경 표시(마지막 반영 대비)', 'markWorkbookChanges'),
    ]

    # db 반영 파이프라인(12~14단계) : 사람이 엑셀(통합 시트/마스터 시트)을 고친 뒤 실행한다.
    # 고친 내용으로 링크를 다시 만들고(12) db에 한 트랜잭션으로 적재한 뒤(13) 전체를 검증한다(14).
    APPLY_PIPELINE_STEPS = [
        ('12단계 니케-키워드 링크', 'buildLinkSheet'),
        ('13단계 DB 반영', 'writeKeywordTablesToDb'),
        ('13단계 반영 기준 저장', 'saveAppliedBaseline'),
        ('14단계 검증', 'validatePipeline'),
    ]

    def runSteps(self, steps):
        '''[파이프라인 보조] (설명, 메서드 이름) 목록을 순서대로 실행하고, 하나라도 실패하면 그 자리에서 멈춘다.'''
        for description, method_name in steps:
            print(f"\n===== {description} ({method_name}) =====")
            if not getattr(self, method_name)():
                print(f"[ERROR] 파이프라인 중단: '{description}' 단계 실패 — 위 로그를 확인")
                return False
        return True

    def findUnappliedWorkbookEdits(self):
        '''
        [파이프라인 보조] 기존 엑셀에 db에 아직 반영하지 않은 수작업이 있는지 확인한다(엑셀 재생성 전 안전장치).
        엑셀은 db가 원본이므로, 반영 전에 다시 만들면 승환님의 수작업이 사라진다("유사 표현 통합" 절 5번 작업 순서).
        - '키워드 통합' 시트의 통합 내용이 db의 통합 테이블과 다르면 → 반영 안 된 통합 작업
        - '스킬 키워드 마스터' 시트가 그 엑셀의 키워드 목록 + 통합 시트로 계산한 마스터와 다르면 → 손으로 고친 마스터
          (새 니케가 추가돼 마스터가 db와 달라진 것은 수작업이 아니므로 잡지 않는다)
        엑셀이 없거나 예전 형식이라 해당 시트가 없으면 그 항목은 건너뛴다. 발견한 사유 목록을 돌려준다(없으면 빈 목록).
        '''
        reasons = []
        if not self.path_excel.exists():
            return reasons
        workbook = openpyxl.load_workbook(self.path_excel)

        if '키워드 통합' in workbook.sheetnames:
            groups, _ = self.readUnifySheet(workbook)
            sheet_unify_map, _ = self.buildUnifyMap(groups)
            db_groups, _ = self.loadUnifyGroupsFromDb()
            db_unify_map, _ = self.buildUnifyMap(db_groups)
            if sheet_unify_map != db_unify_map:
                only_sheet = len(set(sheet_unify_map.items()) - set(db_unify_map.items()))
                only_db = len(set(db_unify_map.items()) - set(sheet_unify_map.items()))
                reasons.append(f"'키워드 통합' 시트가 db와 다름(시트에만 {only_sheet}건, db에만 {only_db}건)")

            needed_sheets = {'스킬 키워드 마스터', '스킬 키워드 디버깅 목록', '제외 키워드'}
            if needed_sheets <= set(workbook.sheetnames):
                candidates, _ = self.collectMasterCandidates(workbook)
                derived = set()
                for column_name in self.MASTER_COLUMN_NAMES:
                    for keyword in candidates[column_name]:
                        mapped = sheet_unify_map.get((column_name, keyword), keyword)
                        if mapped != self.UNIFY_EXCLUDE_MARK:
                            derived.add((column_name, mapped))
                sheet_master = workbook['스킬 키워드 마스터']
                master_header = [cell.value for cell in sheet_master[1]]
                actual = set()
                for row in sheet_master.iter_rows(min_row=2, values_only=True):
                    for name, cell in zip(master_header, row):
                        if name in self.MASTER_COLUMN_NAMES and cell:
                            actual.add((name, str(cell).strip()))
                if actual != derived:
                    reasons.append(f"'스킬 키워드 마스터' 시트를 손으로 고친 흔적(계산 결과와 {len(actual ^ derived)}건 다름)")
        return reasons

    def runWorkbookPipeline(self, force=False):
        '''
        [진입점: 엑셀 생성] 2~12단계를 실행해 nikke_skill_keyword.xlsx를 db 기준으로 새로 만든다.
        db에 반영하지 않은 수작업이 기존 엑셀에 있으면(findUnappliedWorkbookEdits) 엑셀을 덮어쓰지 않고 멈춘다.
        수작업을 버려도 될 때만 force=True로 실행한다.
        다음 할 일 : 엑셀의 '통합 점검'/'키워드 통합'/'스킬 키워드 마스터' 시트를 확인·수정한 뒤 runApplyPipeline.
        '''
        reasons = self.findUnappliedWorkbookEdits()
        if reasons and not force:
            print("[ERROR] 엑셀 생성 중단: 기존 엑셀에 db에 반영하지 않은 수작업이 있음 — 덮어쓰면 사라진다")
            for reason in reasons:
                print(f"  - {reason}")
            print("  먼저 RUN_APPLY로 db에 반영하거나, 수작업을 버려도 되면 FORCE_WORKBOOK = True로 실행")
            return False
        if reasons:
            print(f"[WARN] FORCE_WORKBOOK : 반영하지 않은 수작업 {len(reasons)}건을 버리고 엑셀을 새로 만듦")
        # 덮어쓰기 전에 기존 엑셀의 '반영 기준'을 읽어 둔다(새 엑셀로 넘기고 변경 표시의 비교 기준으로 씀)
        self.loadPreviousBaseline()
        return self.runSteps(self.WORKBOOK_PIPELINE_STEPS)

    def runApplyPipeline(self):
        '''
        [진입점: db 반영] 사람이 고친 엑셀로 12~14단계를 실행한다 — 링크 재생성 → db 적재(마스터·링크·통합)
        → '반영 기준' 시트 저장(다음 RUN_WORKBOOK의 변경 표시 기준) → 검증.
        엑셀을 새로 만들지 않으므로 수작업이 그대로 db에 들어간다. 검증 결과는 엑셀의 '검증 결과' 시트에 남는다.
        '''
        return self.runSteps(self.APPLY_PIPELINE_STEPS)

# ==============================================================================
# 메인 실행 파이프라인
# ==============================================================================
# 사용법 : 아래 실행 설정에서 실행할 작업만 True로 바꾼 뒤 이 파일을 실행한다(VSCode 실행 버튼).
#   True인 작업이 여럿이면 crawl → workbook → apply 순서로 실행하고, 하나라도 실패하면 거기서 멈춘다.
# 일반적인 흐름 : RUN_CRAWLING → RUN_WORKBOOK → (엑셀에서 '통합 점검' 확인, '키워드 통합'/'스킬 키워드 마스터' 수정) → RUN_APPLY
if __name__ == '__main__':
    # ===== 실행 설정 =====
    RUN_CRAWLING = True     # 1단계 : 크롤링으로 nikke_info_table 최신화 + 변경 보고서 + 프로필 이미지 다운로드
    RUN_WORKBOOK = False    # 2~12단계 : db 기준으로 스킬 키워드 엑셀(nikke_skill_keyword.xlsx)을 새로 만듦
    FORCE_WORKBOOK = False  # RUN_WORKBOOK에서 db에 반영하지 않은 엑셀 수작업이 있어도 엑셀을 새로 만듦(수작업은 사라짐)
    RUN_APPLY = False       # 12~14단계 : 사람이 고친 엑셀로 링크 재생성 → db 반영 → 검증

    # python 실제 실행 경로 기준에서 db 파일 지정
    db_file_path = Path(__file__).resolve().parent / "nikke_info.db"
    is_success = True

    if RUN_CRAWLING:
        manager = DBManager(str(db_file_path))

        # DB 테이블 생성
        manager.createDBTable()

        # 보고서 작성용 - 임시로 DB 백업
        manager.backupDB()

        # 1단계: 크롤링 페이지 접속 및 DB 최신화
        manager.accessPage()
        manager.crawlingAllUpdateDB(only_new=False)
        # 변경 사항 보고서 생성. backupDB와 한쌍 함수. backupDB로 백업한 DB를 이 함수에서 삭제.
        manager.generateChangeReport()

        # 이미지 다운로드
        manager.DownloadResource()

    if is_success and (RUN_WORKBOOK or RUN_APPLY):
        # 스킬 키워드 테이블(마스터/링크/통합)이 없는 db에서도 동작하도록 테이블부터 보장한다(이미 있으면 그대로)
        DBManager(str(db_file_path)).createDBTable()
        analyzer = KoreanAnalyzer(str(db_file_path))

        if RUN_WORKBOOK:
            # 2~12단계: db 기준으로 엑셀 전체 재생성
            is_success = analyzer.runWorkbookPipeline(force=FORCE_WORKBOOK)
        if is_success and RUN_APPLY:
            # 12~14단계: 사람이 고친 엑셀로 링크 재생성 → db 반영 → 검증
            is_success = analyzer.runApplyPipeline()

    print("[완료]" if is_success else "[실패] 위 로그를 확인")
